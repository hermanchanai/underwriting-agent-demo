"""Parse one uploaded file in a child process that cannot take the server down with it.

Every parser in here is handed a stranger's file. ``pypdf``, ``openpyxl``, ``xlrd`` and
``python-docx`` are all perfectly ordinary libraries that will happily allocate whatever a
malformed structure asks them to, and none of them has a "stop after N seconds" argument.
So the isolation is the process, not the library: address space and CPU time are capped
with ``setrlimit`` inside the child before a single byte is read, and the parent adds a
wall-clock timeout for the case where the child is blocked rather than busy.

Protocol: a JSON job on stdin, a JSON result on stdout, nothing else on stdout ever.

    {"kind": "xlsx", "bytes": 12345, "run": "<run id>", "filename": "original name.xlsx"}
    → {"ok": true, "document": {...}}  |  {"ok": false, "code": "...", "message": "..."}

``filename`` is METADATA. It is what the visitor called the file and what the page shows
back to them; it never names anything on disk. The child writes the document under a name
built from ``run`` and ``kind`` — see :func:`scratch_name`.

The normalized document is the only thing that survives the child:

    {filename, kind, text|None, sheets: [{name, rows}], stats: {...}}

Run directly for a one-off:  python -m server.parse_worker  < job.json
"""

from __future__ import annotations

import datetime as dt
import io
import json
import logging
import math
import os
import re
try:
    import resource  # POSIX only — the rlimits below
except ImportError:  # Windows: no rlimits; the wall-clock timeout is the backstop
    resource = None
import shutil
import signal
import subprocess
import sys
import tempfile
import threading
import time
from dataclasses import dataclass
from pathlib import Path

log = logging.getLogger("upload_demo.parse")

# --------------------------------------------------------------------------------------
# Limits. Env-overridable because the tests need to prove the CPU limit actually bites,
# and a limit you cannot demonstrate killing something is a limit you have not tested.
# --------------------------------------------------------------------------------------

DEFAULT_MEM_MB = int(os.environ.get("DEMO_PARSE_MEM_MB", "1024"))
DEFAULT_CPU_SECONDS = int(os.environ.get("DEMO_PARSE_CPU_SECONDS", "30"))
DEFAULT_WALL_SECONDS = float(os.environ.get("DEMO_PARSE_WALL_SECONDS", "45"))

MAX_TEXT_BYTES = 2 * 1024 * 1024
MAX_SHEETS = 30
#: Rows across the WHOLE workbook, not per sheet. Per sheet, thirty tabs of twenty-five
#: thousand rows was 750,000 rows and up to 150 million cells handed back to the single
#: API process as JSON — the child's memory limit does not protect the parent while it
#: decodes that.
MAX_ROWS = 25_000
MAX_COLS = 200
#: And a cell budget, because 25,000 rows of 200 columns is still five million values.
MAX_CELLS = int(os.environ.get("DEMO_MAX_CELLS", 400_000))
#: The parent refuses a result larger than this before decoding it.
MAX_RESULT_BYTES = int(os.environ.get("DEMO_MAX_RESULT_BYTES", 64 * 1024 * 1024))

#: How much of the child's stderr the parent keeps. Small: it is a log line, not data.
MAX_STDERR_BYTES = 64 * 1024


@dataclass(frozen=True)
class ParseFailure(Exception):
    code: str
    message: str


#: Filled by :func:`child_main` before parsing and folded into the document's stats, so
#: the parent can see which limits the child actually got rather than which it asked for.
APPLIED_LIMITS: dict = {}


# --------------------------------------------------------------------------------------
# Child side
# --------------------------------------------------------------------------------------


def apply_limits(mem_mb: int = DEFAULT_MEM_MB, cpu_seconds: int = DEFAULT_CPU_SECONDS) -> dict:
    """Cap this process before it touches the file, and report what actually stuck.

    ``RLIMIT_CPU`` turns an infinite loop into a ``SIGXCPU`` and works everywhere.
    ``RLIMIT_AS`` would turn a runaway allocation into a ``MemoryError``, and **macOS
    refuses to set it**: both ``RLIMIT_AS`` and ``RLIMIT_DATA`` raise ``ValueError:
    current limit exceeds maximum limit`` from an unprivileged process, so the memory cap
    is a Linux control and inert on the build machine.

    That distinction is returned rather than logged away. A limit that silently failed to
    apply is worse than no limit, because everything downstream is written as though it
    held — so the child reports it, the stats carry it, and the test asserts the truth for
    the platform it is running on instead of a hope.
    """
    applied: dict = {}
    if resource is None:
        # No rlimit facility on this platform (Windows). Said in the stats, not assumed
        # away: everything downstream must know only the wall clock is holding the line.
        print(f"parse: no rlimit facility on {sys.platform}; the wall-clock timeout is the only backstop",
              file=sys.stderr)
        return {"cpu_limit_applied": False, "memory_limit_applied": False, "rlimits_unavailable": sys.platform}
    soft_cpu = max(1, cpu_seconds)
    try:
        resource.setrlimit(resource.RLIMIT_CPU, (soft_cpu, soft_cpu + 5))
        applied["cpu_limit_applied"] = True
    except (ValueError, OSError) as exc:  # pragma: no cover - would be a broken platform
        applied["cpu_limit_applied"] = False
        print(f"parse: CPU limit could not be set: {exc}", file=sys.stderr)

    applied["memory_limit_applied"] = False
    if mem_mb > 0:
        limit = mem_mb * 1024 * 1024
        for name in ("RLIMIT_AS", "RLIMIT_DATA"):
            which = getattr(resource, name, None)
            if which is None:
                continue
            try:
                resource.setrlimit(which, (limit, limit))
                applied["memory_limit_applied"] = True
                applied["memory_limit_via"] = name
                break
            except (ValueError, OSError) as exc:
                print(
                    f"parse: {name} not settable on {sys.platform} ({exc}); "
                    f"the wall-clock timeout is the only backstop here",
                    file=sys.stderr,
                )
    return applied


def current_limits() -> dict:
    """What the limits actually are, read back rather than assumed."""
    def read(name: str) -> list:
        which = getattr(resource, name, None) if resource is not None else None
        if which is None:
            return [-1, -1]
        try:
            return list(resource.getrlimit(which))
        except (ValueError, OSError):  # pragma: no cover
            return [-1, -1]

    return {
        "rlimit_cpu": read("RLIMIT_CPU"),
        "rlimit_as": read("RLIMIT_AS"),
        "platform": sys.platform,
    }


def parse_pdf(data: bytes) -> dict:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:  # noqa: BLE001
            raise ParseFailure("encrypted_file", "That PDF is password-protected.") from exc

    chunks: list[str] = []
    size = 0
    truncated = False
    pages = len(reader.pages)
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:  # noqa: BLE001 - one bad page must not lose the other 36
            text = ""
        # The page marker is what lets an extracted quote cite a page number later.
        block = f"[[page {index}]]\n{text}\n"
        size += len(block.encode("utf-8"))
        if size > MAX_TEXT_BYTES:
            truncated = True
            break
        chunks.append(block)

    body = "".join(chunks)
    return {
        "text": body,
        "sheets": [],
        "stats": {
            "pages": pages,
            "pages_read": len(chunks),
            "characters": len(body),
            "words": len(body.split()),
            "truncated": truncated,
        },
    }


def parse_docx(data: bytes) -> dict:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts: list[str] = [p.text for p in document.paragraphs if p.text.strip()]
    tables = 0
    for table in document.tables:
        tables += 1
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append("\t".join(cells))
    body = "\n".join(parts)
    truncated = False
    if len(body.encode("utf-8")) > MAX_TEXT_BYTES:
        body = body.encode("utf-8")[:MAX_TEXT_BYTES].decode("utf-8", "ignore")
        truncated = True
    return {
        "text": body,
        "sheets": [],
        "stats": {
            "paragraphs": len(document.paragraphs),
            "tables": tables,
            "characters": len(body),
            "words": len(body.split()),
            "truncated": truncated,
        },
    }


def parse_doc(path: Path) -> dict:
    """Legacy ``.doc`` via ``antiword``.

    There is no antiword on the build Mac and no Homebrew to install one, so this path is
    exercised in the container and reports an honest refusal here rather than pretending.
    """
    binary = shutil.which("antiword")
    if not binary:
        raise ParseFailure(
            "doc_not_supported",
            "This demo cannot read the old .doc format right now. Save it as .docx or "
            "PDF and try again.",
        )
    try:
        completed = subprocess.run(
            [binary, "-w", "0", str(path)],
            capture_output=True,
            timeout=30,
            check=False,
            # antiword is a native binary being handed a stranger's file. Even inside
            # this already-sandboxed child it gets the allowlist and nothing more, so
            # there is no environment here for it to read either.
            env=child_environment(),
        )
    except subprocess.TimeoutExpired as exc:
        raise ParseFailure("unreadable_file", "That document took too long to read.") from exc
    if completed.returncode != 0:
        raise ParseFailure("unreadable_file", "That .doc file could not be read.")
    body = completed.stdout.decode("utf-8", "replace")[:MAX_TEXT_BYTES]
    return {
        "text": body,
        "sheets": [],
        "stats": {"characters": len(body), "words": len(body.split()), "tool": "antiword"},
    }


def _truncation_reason(
    hit_budget: bool, over_sheets: bool, over_width: bool = False
) -> str | None:
    """Which cap stopped the read, in words a visitor can act on.

    ``truncated`` on its own is a boolean nobody can do anything with. Assembly refuses a
    truncated tabular document outright — a rent roll read to row 25,000 and then
    analysed as if that were the whole building understates unit count, rents and
    occupancy, all in the flattering direction — so the refusal has to name the limit.

    Width counts too. Columns past the cap used to be sliced away in silence, and a month
    column, a rent column or a lease-status column sitting past 200 turns into a complete,
    confident, WRONG analysis rather than a refusal.
    """
    if over_sheets:
        return f"more than {MAX_SHEETS} sheets"
    if over_width:
        return f"more than {MAX_COLS} columns"
    if hit_budget:
        return f"more than {MAX_ROWS:,} rows or {MAX_CELLS:,} cells"
    return None


def _clean_cell(value: object) -> object:
    """One cell, normalised for JSON.

    Dates are the case worth spelling out. openpyxl hands back a real ``datetime`` for any
    date-formatted cell, and ``str()`` renders that as ``2026-01-15 00:00:00`` — which
    every date parser downstream refuses, so lease starts and move-ins silently became
    None on exactly the workbooks Excel itself wrote. ISO, and a bare date when there is
    no time of day, because a midnight is a date typed by a spreadsheet, not a time.
    """
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, dt.datetime):
        if (value.hour, value.minute, value.second, value.microsecond) == (0, 0, 0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, dt.time):
        return value.isoformat()
    if isinstance(value, (int, float)):
        return value
    return str(value).strip()


#: How many formula cells with no cached result a workbook may hold before it is refused.
#:
#: None. `data_only=True` answers None for a formula the file never stored a result for,
#: which is indistinguishable from an empty cell — so a T-12 whose amounts are `=SUM(...)`
#: read as a statement of zeros, and a rent roll's calculated rents vanished. Both produce
#: a complete, confident, wrong output set. A workbook openpyxl wrote —
#: including this demo's OWN outputs — is exactly that file, because openpyxl does not
#: evaluate; opening it in a spreadsheet application and saving is what stores the values.
UNCACHED_FORMULA_MESSAGE = (
    "That spreadsheet stores formulas but not their results, so every calculated cell "
    "reads as blank and the numbers cannot be trusted. Open it in a spreadsheet "
    "application and save it, or export the values, then upload it again."
)


def uncached_formula_cells(data: bytes) -> int:
    """How many formula cells in ``data`` carry no cached result.

    Two views of the same package: the formula view says which cells ARE formulas, and
    the value view says which of those the file actually stored an answer for. Bounded by
    the same sheet and row budget the value pass uses, and it stops at the first offender
    — the question is whether the workbook has any, not how many.
    """
    from openpyxl import load_workbook

    formulas = load_workbook(io.BytesIO(data), data_only=False, read_only=True)
    values = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    try:
        for formula_sheet, value_sheet in zip(
            formulas.worksheets[:MAX_SHEETS], values.worksheets[:MAX_SHEETS]
        ):
            rows = zip(
                formula_sheet.iter_rows(values_only=True, max_row=MAX_ROWS),
                value_sheet.iter_rows(values_only=True, max_row=MAX_ROWS),
            )
            for formula_row, value_row in rows:
                for formula, value in zip(formula_row, value_row):
                    # `ArrayFormula` and `DataTableFormula` are objects rather than
                    # strings, and both are formulas; a plain string starting with "="
                    # is the ordinary case.
                    is_formula = (
                        isinstance(formula, str) and formula.startswith("=")
                    ) or type(formula).__name__ in ("ArrayFormula", "DataTableFormula")
                    if is_formula and value is None:
                        return 1
        return 0
    finally:
        for book in (formulas, values):
            try:
                book.close()
            except Exception:  # noqa: BLE001 - closing a read-only book is best effort
                pass


def parse_xlsx(data: bytes) -> dict:
    from openpyxl import load_workbook

    # BEFORE the values are read, because the refusal is about what the values ARE. A
    # sheet of Nones is a valid-looking parse; the only place the difference is visible is
    # here, against the formula view.
    if uncached_formula_cells(data):
        raise ParseFailure("formula_values_missing", UNCACHED_FORMULA_MESSAGE)

    # A STREAM, not a path — which also retires the extension check `load_workbook`
    # applies to a filename and nothing else. See `KIND_SUFFIX`.
    workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    try:
        sheets = []
        truncated = False
        over_width = False
        budget_rows = MAX_ROWS
        budget_cells = MAX_CELLS
        for worksheet in workbook.worksheets[:MAX_SHEETS]:
            rows: list[list] = []
            if budget_rows <= 0 or budget_cells <= 0:
                truncated = True
                break
            for row in worksheet.iter_rows(values_only=True):
                if len(row) > MAX_COLS:
                    over_width = True
                cells = [_clean_cell(v) for v in row[:MAX_COLS]]
                # Refuse the row that does not FIT, rather than the one that exactly
                # fills the budget. Decrementing first and testing `<= 0` marked a file of
                # exactly MAX_ROWS truncated because its last permitted row took the
                # budget to zero — the documented limit, reported as a partial read.
                if budget_rows <= 0 or budget_cells < len(cells):
                    truncated = True
                    break
                rows.append(cells)
                budget_rows -= 1
                budget_cells -= len(cells)
            sheets.append({"name": worksheet.title, "rows": rows})
        sheet_total = len(workbook.worksheets)
    finally:
        workbook.close()

    over_sheets = sheet_total > MAX_SHEETS
    return {
        "text": None,
        "sheets": sheets,
        "stats": {
            "sheets": sheet_total,
            "sheets_read": len(sheets),
            "rows": sum(len(s["rows"]) for s in sheets),
            "truncated": truncated or over_sheets or over_width,
            "truncated_reason": _truncation_reason(
                truncated, over_sheets, over_width
            ),
        },
    }


def _xls_cell(value: object, cell_type: int, datemode: int, xlrd) -> object:
    """One legacy cell, with the date type actually honoured.

    ``.xls`` has no date type on the wire: a date is a float, and the only thing that says
    so is the cell's type code. Returning the float means every date on the sheet arrives
    downstream as ``45306.0`` — refused by the lease-date parser and by month-header
    detection, so a legacy statement reads as "not a T-12" and a legacy roll loses every
    lease date, on a format this demo advertises as supported.
    """
    if cell_type != xlrd.XL_CELL_DATE:
        return _clean_cell(value)
    try:
        return _clean_cell(xlrd.xldate_as_datetime(value, datemode))
    except (ValueError, OverflowError, xlrd.XLDateError):
        # A date cell holding a value the epoch cannot represent. The raw number is more
        # honest than a wrong date, and it will simply fail to parse as a date later.
        return _clean_cell(value)


def parse_xls(data: bytes) -> dict:
    import xlrd

    book = xlrd.open_workbook(file_contents=data)
    sheets = []
    truncated = False
    over_width = False
    budget_rows = MAX_ROWS
    # The aggregate CELL budget, which this parser never had. A compact 25,000 x 200 file
    # is inside every size limit on the way in and five million Python objects once the
    # parent decodes it, which is where the container dies.
    budget_cells = MAX_CELLS
    for sheet in book.sheets()[:MAX_SHEETS]:
        rows: list[list] = []
        for i in range(sheet.nrows):
            if sheet.ncols > MAX_COLS:
                over_width = True
            values = sheet.row_values(i)[:MAX_COLS]
            types = sheet.row_types(i)[:MAX_COLS]
            if budget_rows <= 0 or budget_cells < len(values):
                truncated = True
                break
            rows.append([
                _xls_cell(v, t, book.datemode, xlrd) for v, t in zip(values, types)
            ])
            budget_rows -= 1
            budget_cells -= len(values)
        if sheet.nrows > len(rows):
            truncated = True
        sheets.append({"name": sheet.name, "rows": rows})
    over_sheets = book.nsheets > MAX_SHEETS
    return {
        "text": None,
        "sheets": sheets,
        "stats": {
            "sheets": book.nsheets,
            "sheets_read": len(sheets),
            "rows": sum(len(s["rows"]) for s in sheets),
            "truncated": truncated or over_sheets or over_width,
            "truncated_reason": _truncation_reason(
                truncated, over_sheets, over_width
            ),
        },
    }


def parse_csv(data: bytes, display_name: str | None = None) -> dict:
    """A delimited file, with its single sheet named after the UPLOAD.

    ``display_name`` rather than a path stem: this parser is handed BYTES and there is no
    path to read a name off — and when there was one it named the run, not the upload.
    """
    import csv as csv_mod

    raw = data
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")

    sample = "\n".join(text.splitlines()[:20])
    try:
        dialect = csv_mod.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv_mod.Error:
        dialect = csv_mod.excel

    rows = []
    truncated = False
    over_width = False
    budget_cells = MAX_CELLS
    # A newline-preserving STREAM, not `text.splitlines()`. Splitting first removed the
    # line terminators inside quoted fields before `csv.reader` could see them, so a
    # perfectly valid record whose account label spans two lines arrived as two rows —
    # or, once the reader rejoined them, as "RentalIncome". `newline=""`
    # is what the csv module's own documentation asks for, for exactly this reason.
    for row in csv_mod.reader(io.StringIO(text, newline=""), dialect):
        if len(row) > MAX_COLS:
            over_width = True
        cells = [_clean_cell(v) for v in row[:MAX_COLS]]
        # Refuse the row that does not fit rather than the one that exactly fills the
        # budget: a file of exactly MAX_ROWS is the documented limit, not a partial read.
        if len(rows) >= MAX_ROWS or budget_cells < len(cells):
            truncated = True
            break
        rows.append(cells)
        budget_cells -= len(cells)

    return {
        "text": None,
        "sheets": [{"name": Path(display_name or "sheet").stem, "rows": rows}],
        "stats": {
            "sheets": 1,
            "sheets_read": 1,
            "rows": len(rows),
            "delimiter": getattr(dialect, "delimiter", ","),
            "truncated": truncated or over_width,
            "truncated_reason": _truncation_reason(truncated, False, over_width),
        },
    }


def probe_pdf(data: bytes) -> dict:
    """Page count and encryption only — no text extraction.

    Split out so the intake gate can learn whether a PDF is encrypted or over the page cap
    WITHOUT running pypdf in the API process. Opening an attacker-controlled PDF is
    exactly the work the sandbox exists for, and doing it in the parent to answer a cap
    question handed a malformed file the one process that has no CPU, memory or wall-clock
    limit around it.
    """
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    encrypted = bool(reader.is_encrypted)
    if encrypted:
        try:
            encrypted = not reader.decrypt("")
        except Exception:  # noqa: BLE001
            encrypted = True
    pages = 0 if encrypted else len(reader.pages)
    return {"text": None, "sheets": [], "stats": {"pages": pages, "encrypted": encrypted}}


def probe_ole(data: bytes) -> dict:
    """Which legacy Office format this compound file is — stream names only.

    Split out for exactly the reason ``probe_pdf`` was: the question has to be answered
    before a parser is chosen, and answering it in the API process meant the root process,
    on the event loop, running a general-purpose parser over a stranger's crafted file.
    ``olefile`` is a perfectly ordinary library that will allocate whatever a malformed
    compound file asks it to; here it does that under the child's CPU, memory and
    wall-clock limits, where a bad file costs one child rather than the service.
    """
    import olefile

    ole = olefile.OleFileIO(io.BytesIO(data))
    try:
        streams = sorted({"/".join(entry).lower() for entry in ole.listdir()})
    finally:
        ole.close()
    if "workbook" in streams or "book" in streams:
        kind = "xls"
    elif "worddocument" in streams:
        kind = "doc"
    else:
        kind = None
    return {
        "text": None,
        "sheets": [],
        # The stream NAMES, capped: they are structure, not content, and they are what a
        # maintainer reads when a legacy file is refused.
        "stats": {"ole_kind": kind, "streams": streams[:40]},
    }


PARSERS = {
    "pdf": parse_pdf,
    "probe_pdf": probe_pdf,
    "probe_ole": probe_ole,
    "docx": parse_docx,
    "doc": parse_doc,
    "xlsx": parse_xlsx,
    "xls": parse_xls,
    "csv": parse_csv,
}

#: Parsers that also need the name the VISITOR gave the file. Only ``csv``: a delimited
#: file has no sheet name of its own, and its parser is handed bytes with no path to read
#: a name off, so the display name is passed in.
PARSERS_WANTING_NAME: frozenset = frozenset({"csv"})

#: The ONLY kinds whose parser needs a file on disk.
#:
#: Every other parser reads the document out of memory: pypdf, python-docx and olefile
#: take a stream, openpyxl takes one too, and xlrd takes ``file_contents``. So an ordinary
#: parse now writes NOTHING — the child is handed the bytes on stdin and never opens a
#: path. Scratch directories were mode 0700 but chowned to the single
#: `demo-parser` account, so two overlapping parsers shared an owner and either could read
#: the other visitor's raw upload; the fix is to stop having the file at all.
#:
#: ``antiword`` is a native binary that takes a path, so ``.doc`` still needs one — and
#: those parses are serialised under `DOC_LOCK`, so at most one such directory exists at
#: any moment and it never coexists with another visitor's.
KINDS_NEEDING_FILE: frozenset = frozenset({"doc"})

#: Serialises the one parse that still writes the upload to disk. Process-wide, held
#: across the whole child launch: the scratch is made inside it and removed inside it.
DOC_LOCK = threading.Lock()

#: The suffix a parser is allowed to see, chosen from the MAGIC-BYTE decision and never
#: from the upload. Only the file-backed kinds reach it now.
#:
#: ``openpyxl.load_workbook`` refuses outright any path whose extension is not one of its
#: own four, so a perfectly good workbook uploaded as ``quote.csv`` came back
#: ``unreadable_file`` — the uploader's extension overriding the gate's own reading of the
#: bytes. The other parsers took the path as given, which was the same problem pointing
#: the other way. Reading from a stream retires the whole question for every kind but one.
KIND_SUFFIX: dict[str, str] = {
    "pdf": ".pdf",
    "probe_pdf": ".pdf",
    "probe_ole": ".bin",
    "docx": ".docx",
    "doc": ".doc",
    "xlsx": ".xlsx",
    "xls": ".xls",
    "csv": ".csv",
}

#: Anything outside this is dropped from a run id before it names a file. Run ids are
#: ``secrets.token_urlsafe``, so nothing legitimate is ever removed; the filter is here so
#: that a job assembled by hand cannot reintroduce what this whole function exists to stop.
_UNSAFE_IN_NAME = re.compile(r"[^A-Za-z0-9_-]")


def scratch_name(job: dict) -> str:
    """The name the document is written under inside the child's own directory.

    **Never the visitor's filename.** ``Path(scratch) / "/tmp/payload.xlsx"`` is
    ``/tmp/payload.xlsx`` — an absolute component discards everything to its left — and
    ``../../evil.xlsx`` walks out of the directory the cleanup deletes. Either one let a
    public upload leave up to 25 MB at any path the parser account can write, outliving
    the run that produced it. The name is built from the run id and the kind the
    magic-byte gate decided; the uploaded name travels as metadata and nothing else.

    A job with no run id still gets a safe name: the scratch DIRECTORY is already unique
    per child, so the file inside it only has to be predictable and inert.
    """
    run = _UNSAFE_IN_NAME.sub("", str(job.get("run") or ""))[:48] or "document"
    return f"{run}{KIND_SUFFIX.get(str(job.get('kind') or ''), '.bin')}"


def parse_job(job: dict, data: bytes, path: "Path | None" = None) -> dict:
    """One document, from the bytes the parent sent — or from a file, for ``.doc`` alone.

    ``path`` is passed only for the kinds in `KINDS_NEEDING_FILE`; everything else is
    parsed out of memory and never touches the filesystem.
    """
    kind = job["kind"]
    parser = PARSERS.get(kind)
    if parser is None:
        raise ParseFailure("unsupported_type", "That file type is not one this demo reads.")
    display = job.get("filename") or "document"
    if kind in KINDS_NEEDING_FILE:
        if path is None:
            raise ParseFailure("unreadable_file", "That file could not be read.")
        result = parser(path)
    else:
        result = parser(data, display) if kind in PARSERS_WANTING_NAME else parser(data)
    result["filename"] = display
    result["kind"] = kind
    result["stats"] = {**result.get("stats", {}), **current_limits(), **APPLIED_LIMITS}
    return result


def child_main() -> int:
    applied = apply_limits()
    APPLIED_LIMITS.update(applied)
    stream = sys.stdin.buffer
    try:
        header = stream.readline()
        job = json.loads(header.decode("utf-8") or "{}")
        wanted = int(job.get("bytes") or 0)
    except (ValueError, UnicodeDecodeError):
        sys.stdout.write(json.dumps({"ok": False, "code": "bad_job", "message": "unreadable job"}))
        return 2

    # The document arrives on stdin and, for every kind but one, STAYS IN MEMORY. Writing
    # it to a scratch directory gave a compromised parser a path to read — every scratch
    # was 0700 but owned by the single `demo-parser` account, so two overlapping children
    # shared an owner and either could open the other visitor's raw upload. pypdf, python-docx, olefile and openpyxl all take a stream and xlrd takes
    # `file_contents`, so the file was never needed.
    #
    # `.doc` is the exception: antiword is a native binary that takes a path. Those parses
    # are serialised by the parent (`DOC_LOCK`), and the directory is still the PARENT's —
    # a SIGKILL here runs no `finally`, so this process must never be the one that has to
    # tidy up. The NAME inside it is generated — see `scratch_name`.
    data = b""
    while len(data) < wanted:
        chunk = stream.read(wanted - len(data))
        if not chunk:
            break
        data += chunk
    path = None
    if str(job.get("kind") or "") in KINDS_NEEDING_FILE:
        scratch = str(job.get("scratch") or "")
        if not scratch or not Path(scratch).is_dir():
            sys.stdout.write(json.dumps({
                "ok": False, "code": "bad_job",
                "message": "unreadable job",
            }))
            return 2
        path = Path(scratch) / scratch_name(job)

    try:
        # Inside the try, so a write that fails answers with a code and still deletes its
        # directory rather than dying with a traceback and leaving the scratch behind.
        if path is not None:
            path.write_bytes(data)
        document = parse_job(job, data, path)
        # Serialise, MEASURE, and refuse here rather than writing it and letting the
        # parent discover the size. The child is the only process that has to hold this
        # string at all, and it is the sandboxed one.
        body = json.dumps({"ok": True, "document": document})
        if len(body) > MAX_RESULT_BYTES:
            sys.stdout.write(
                json.dumps(
                    {
                        "ok": False,
                        "code": "file_too_complex",
                        "message": (
                            f"{job.get('filename') or 'That file'} holds more data than "
                            f"this demo will read."
                        ),
                    }
                )
            )
            return 0
        sys.stdout.write(body)
        return 0
    except ParseFailure as failure:
        sys.stdout.write(
            json.dumps({"ok": False, "code": failure.code, "message": failure.message})
        )
        return 0
    except MemoryError:
        sys.stdout.write(
            json.dumps(
                {
                    "ok": False,
                    "code": "file_too_complex",
                    "message": "That file needed more memory than this demo allows.",
                }
            )
        )
        return 0
    except Exception as exc:  # noqa: BLE001 - the parent gets a code, the log gets a class
        # The exception CLASS only. `repr(exc)` on a parser failure routinely contains the
        # cell, the row or the byte range it failed on, and this stream is drained by the
        # parent — belt and braces with the parent's own sanitising, since the child is
        # the process actually holding the document.
        print(f"parse worker failed: {type(exc).__name__}", file=sys.stderr)
        sys.stdout.write(
            json.dumps(
                {
                    "ok": False,
                    "code": "unreadable_file",
                    "message": "That file could not be read.",
                }
            )
        )
        return 0
    finally:
        # The FILE, not the directory: the directory belongs to the parent, which removes
        # it after reaping this process however it ended. Best-effort here so a normal
        # exit does not leave the document sitting there for the microseconds until the
        # parent gets to it. Nothing to do at all on the in-memory path.
        if path is not None:
            try:
                path.unlink()
            except OSError:
                pass


# --------------------------------------------------------------------------------------
# Parent side
# --------------------------------------------------------------------------------------


#: The only variables a parser child is given.
#:
#: A child here opens a stranger's file with pypdf, openpyxl, xlrd, python-docx or the
#: native `antiword` binary. If any of them can be made to run something, the process it
#: runs in used to hold the whole server environment: the Anthropic key, the download
#: service's own secrets. rlimits bound what a child
#: can ALLOCATE; they say nothing about what it can read and send. So the child is handed
#: an allowlist, and everything else stays in the parent.
#:
#: PYTHONPATH is added by :func:`child_environment` rather than listed here, because it is
#: computed rather than inherited.
CHILD_ENV_ALLOW: tuple[str, ...] = (
    "PATH", "HOME", "TMPDIR", "TEMP", "TMP",
    "LANG", "LC_ALL", "LC_CTYPE",
    "PYTHONHASHSEED", "PYTHONIOENCODING", "SYSTEMROOT",
    # The parser limits, so a deployment that tunes them tunes the child too. All five,
    # including the two rlimits: those are read at IMPORT in the child, so a
    # DEMO_PARSE_MEM_MB or DEMO_PARSE_CPU_SECONDS left off this list is an override that
    # is documented, accepted, and silently ignored exactly where it is applied — the
    # child imported the hardcoded 1024 MB and 30 seconds instead.
    "DEMO_MAX_RESULT_BYTES", "DEMO_MAX_CELLS", "DEMO_PARSE_WALL_SECONDS",
    "DEMO_PARSE_MEM_MB", "DEMO_PARSE_CPU_SECONDS",
)


def child_environment(extra: dict | None = None) -> dict:
    """The minimal environment a parser child runs in."""
    env = {
        name: os.environ[name] for name in CHILD_ENV_ALLOW if name in os.environ
    }
    env.setdefault("PATH", "/usr/local/bin:/usr/bin:/bin")
    # openpyxl switches its serialiser to lxml the moment lxml is importable, and lxml
    # writes the same worksheet to different bytes. Pinned here as well as in the parent
    # because the child does its own reading.
    env["OPENPYXL_LXML"] = "False"
    env["PYTHONDONTWRITEBYTECODE"] = "1"
    env["PYTHONPATH"] = str(Path(__file__).resolve().parent.parent)
    if extra:
        env.update(extra)
    return env


#: The unprivileged account parser children run as, created in the container image.
PARSER_USER = os.environ.get("DEMO_PARSER_USER", "demo-parser")

#: What the isolation probe found, filled once at import and logged at startup.
_ISOLATION: dict = {}


def _parser_identity() -> "tuple[int, int] | None":
    """(uid, gid) of the parser account, or None when it does not exist here."""
    try:
        import pwd

        entry = pwd.getpwnam(PARSER_USER)
    except (ImportError, KeyError):
        return None
    return entry.pw_uid, entry.pw_gid


#: The binary that gives a child its own network namespace, when the host permits it.
#:
#: A BINARY rather than a call, because the drop has to happen between fork and exec
#: without running Python in the forked child — see :func:`child_launch`. ``unshare -n --``
#: is exec-safe by construction: it is the thing being exec'd.
UNSHARE_BINARY = "unshare"

#: The binary that drops the uid AFTER the namespace has been applied.
#:
#: ``unshare -n`` needs CAP_SYS_ADMIN, which the parser account does not have — so the two
#: cannot be done in that order. ``Popen(user=...)`` drops between fork and exec, i.e.
#: BEFORE the ``unshare`` it was prepending, and every launch on a host that granted the
#: namespace would have failed. ``setpriv`` is the exec-safe way round:
#: root unshares, then setpriv drops, then the parser runs — one exec chain, no fork hook.
SETPRIV_BINARY = "setpriv"

#: The four ways a child can be launched here, in descending order of isolation. Stored on
#: the isolation report so the startup line says which one this host actually gets.
LAUNCH_NAMESPACED_DROP = "unshare -n, then setpriv drops to the parser account"
LAUNCH_NAMESPACED = "unshare -n, with no parser account to drop to"
LAUNCH_DROP = "Popen user/group drop to the parser account, no namespace"
LAUNCH_PLAIN = "no isolation available on this host"


def _launch_plan(state: dict) -> str:
    """Which launch path a child gets, from what the probe found. One decision.

    Read by :func:`isolation_report` so the startup line can print it and by
    :func:`child_launch` so the launch does it — the two cannot disagree, which is the
    whole finding: the probe said "namespaced" and the launch then dropped the privileges
    the namespace needed.
    """
    have_unshare = bool(shutil.which(UNSHARE_BINARY))
    have_setpriv = bool(shutil.which(SETPRIV_BINARY))
    if state.get("network_isolated") and have_unshare:
        if not state.get("uid_dropped"):
            return LAUNCH_NAMESPACED
        if have_setpriv:
            return LAUNCH_NAMESPACED_DROP
        # Both were wanted and only one is reachable. The uid drop is the one that has
        # been proven to hold; the namespace has never been available on any host this
        # has run on. Dropping is strictly better than a launch that fails.
        return LAUNCH_DROP
    return LAUNCH_DROP if state.get("uid_dropped") else LAUNCH_PLAIN


def _probe_network_namespace() -> bool:
    """Can a child unshare its network? Asked ONCE, in a throwaway process.

    Two conditions, both probed rather than assumed: the ``unshare`` binary has to exist
    (it is what applies the namespace now that no fork hook does), and it has to actually
    succeed here. ``CLONE_NEWNET`` needs CAP_SYS_ADMIN, which most containers do not grant.
    """
    if not sys.platform.startswith("linux"):
        return False
    binary = shutil.which(UNSHARE_BINARY)
    if not binary:
        return False
    probe = subprocess.run(
        [binary, "-n", "--", sys.executable, "-c", "raise SystemExit(0)"],
        capture_output=True, timeout=20, check=False,
    )
    return probe.returncode == 0


def isolation_report() -> dict:
    """What isolation actually applied, in words a startup line can print.

    Probed rather than assumed, and never fatal. The honest answer on the build Mac is
    "none of it": there is no root and no parser account, and saying so is worth more than
    a comment claiming a sandbox that is not there.
    """
    if _ISOLATION:
        return dict(_ISOLATION)
    identity = _parser_identity()
    is_root = hasattr(os, "geteuid") and os.geteuid() == 0
    if identity is None:
        uid_reason = f"no {PARSER_USER!r} account on this host"
    elif not is_root:
        uid_reason = f"not running as root, so cannot drop to {PARSER_USER!r}"
    else:
        uid_reason = f"children run as {PARSER_USER!r} (uid {identity[0]})"
    network = _probe_network_namespace() if is_root else False
    _ISOLATION.update({
        "uid_dropped": bool(identity and is_root),
        "uid_reason": uid_reason,
        "network_isolated": bool(network),
        "network_reason": (
            f"children are launched through {UNSHARE_BINARY!r} in their own network "
            f"namespace" if network
            else "CLONE_NEWNET is not permitted here, so children share the network"
        ),
        "parser_user": PARSER_USER,
    })
    _ISOLATION["launch"] = _launch_plan(_ISOLATION)
    return dict(_ISOLATION)


#: Every prefix a RAW uploaded document is ever written under, keyed by what makes it.
#:
#: ONE table, because the sweep iterates it. Three of these four had no cleanup at all:
#: a process killed mid-run left a visitor's document under `upload-parse-`,
#: `upload-probe-` or `upload-ole-` for the life of the filesystem, and the startup sweep
#: only ever looked for `parse-scratch-`. A prefix typed at a call site
#: is a prefix this sweep cannot know about, so :func:`raw_scratch` is the only way to
#: make one and it refuses a kind that is not in here.
#:
#: The names differ on purpose: an operator who finds one on a running host can tell
#: which side of the service made it.
RAW_SCRATCH_PREFIXES: dict[str, str] = {
    # antiword's one file, written BY the child — the only parse that needs a path.
    "doc": "parse-scratch-",
    # The upload worker's copy of a visitor's batch, one file at a time.
    "upload": "upload-parse-",
    # The intake page-count / encryption probe.
    "probe": "upload-probe-",
    # The intake doc-vs-xls probe.
    "ole": "upload-ole-",
}

#: The `.doc` scratch prefix under its old name. Same string, one definition.
SCRATCH_PREFIX = RAW_SCRATCH_PREFIXES["doc"]

#: The file a scratch directory carries saying when it was created.
#:
#: The same marker, and the same name, that a run directory carries — a retention clock
#: read off mtime is a clock the directory's own contents reset, and every one of these
#: directories has a document written into it after it is made.
CREATED_MARKER = ".created"


def write_creation_stamp(path: Path, now: float | None = None) -> None:
    """Stamp ``path`` with the moment it was created. A failure is logged, never raised.

    The fallback is the pre-marker behaviour — age from mtime — which is worse but not
    broken, and a directory that cannot take a marker file is a bigger problem than
    retention.
    """
    try:
        (path / CREATED_MARKER).write_text(
            f"{time.time() if now is None else now:.6f}\n", encoding="utf-8"
        )
    except OSError:  # pragma: no cover - a temp root we cannot write is a bigger problem
        log.error("retention: could not stamp the creation time of a scratch directory")


def _stamp_in(marker: Path) -> "float | None":
    """The timestamp one marker file holds, or None when there is not one to read.

    None means "ask the filesystem instead" rather than "brand new": a directory written
    by an older build carries no marker, and reading that as freshly created would keep
    it forever.
    """
    try:
        raw = marker.read_text(encoding="utf-8").strip()
    except (OSError, UnicodeDecodeError):
        return None
    try:
        stamp = float(raw)
    except ValueError:
        return None
    return stamp if math.isfinite(stamp) else None


def creation_stamp(path: Path) -> "float | None":
    """When the RUN directory ``path`` was created, from the marker inside it.

    Run directories are the server's own and nothing else writes into them, so their
    marker stays where round 17 put it. A raw scratch directory's marker is a SIBLING —
    see :func:`scratch_creation_stamp` — because that one can be handed to another uid.
    """
    return _stamp_in(path / CREATED_MARKER)


def scratch_marker(path: Path) -> Path:
    """The creation marker for a raw scratch directory — BESIDE it, never inside it.

    A marker inside the directory it dates is writable by whoever owns that directory,
    and the `.doc` scratch is chowned to the parser account: the one process whose
    lifetime the clock bounds could delete or rewrite it. A sibling in the temp root is
    owned by the server, mode 0600, and the root's sticky bit stops another uid unlinking
    it — so the timestamp is one the subject cannot touch (round 19, ruling 3).

    It also keeps the directory EMPTY until the document arrives, which is what round 6's
    "no raw upload of this run is on disk by then" assertion is really about.
    """
    return Path(str(path) + CREATED_MARKER)


def write_scratch_stamp(path: Path) -> None:
    """Stamp a raw scratch directory's sibling marker. A failure is logged, never raised."""
    marker = scratch_marker(path)
    try:
        marker.write_text(f"{time.time():.6f}\n", encoding="utf-8")
        marker.chmod(0o600)
    except OSError:  # pragma: no cover - a temp root we cannot write is a bigger problem
        log.error("retention: could not stamp the creation time of a scratch directory")


def scratch_creation_stamp(path: Path) -> "float | None":
    """When a raw scratch directory was created, from its sibling. None if unreadable."""
    return _stamp_in(scratch_marker(path))


def remove_raw_scratch(path: Path) -> None:
    """Remove a raw scratch directory AND its sibling marker. The only remover.

    One function because there are five call sites — two probes, the upload worker's two
    exits and the `.doc` parse — and a marker whose directory is gone is litter the sweep
    then has to collect. Everything that makes one through :func:`raw_scratch` unmakes it
    through here.
    """
    shutil.rmtree(path, ignore_errors=True)
    try:
        scratch_marker(path).unlink(missing_ok=True)
    except OSError:  # pragma: no cover - the sweep collects it on the next pass
        pass


def raw_scratch(kind: str, *, child_writable: bool = False) -> Path:
    """One private directory for one raw document, under its registered prefix.

    ``0700`` and, where ``child_writable`` and children drop to the parser account, owned
    BY that account: the child has to write the document into it, and after a uid drop the
    only paths it can write are ones granted to it. Nobody else on the host can read it
    either way.

    ``kind`` must be a key of :data:`RAW_SCRATCH_PREFIXES`. A kind that is not in the
    table is a directory the sweep would never look for, so it is refused here rather
    than created and forgotten.
    """
    try:
        prefix = RAW_SCRATCH_PREFIXES[kind]
    except KeyError:
        raise ValueError(
            f"{kind!r} is not a registered raw-upload scratch kind; add it to "
            f"RAW_SCRATCH_PREFIXES so the sweep can find what it leaves behind"
        ) from None
    scratch = Path(tempfile.mkdtemp(prefix=prefix))
    scratch.chmod(0o700)
    # BESIDE the directory, so the chown below cannot hand the clock to the thing it times.
    write_scratch_stamp(scratch)
    if child_writable:
        state = isolation_report()
        identity = _parser_identity()
        if state["uid_dropped"] and identity is not None:
            try:
                os.chown(scratch, identity[0], identity[1])
            except OSError as exc:  # pragma: no cover - would mean we are not root
                log.error("parse: could not hand %s to the parser account (%s); the "
                          "child will not be able to write its document", scratch, exc)
    return scratch


def make_scratch() -> Path:
    """The ``.doc`` parse's own directory — the one a child writes into."""
    return raw_scratch("doc", child_writable=True)


#: How old a scratch directory has to be before a startup sweep will remove it.
#:
#: Comfortably past the longest a child may live (`DEFAULT_WALL_SECONDS`), so the sweep
#: cannot take a directory out from under a parse that is running — which matters because
#: this is called from `create_app`, and a test process builds many apps while parses are
#: in flight.
SCRATCH_MAX_AGE_SECONDS = 10 * 60


def sweep_scratch(max_age_seconds: float = SCRATCH_MAX_AGE_SECONDS) -> int:
    """Remove raw-upload scratch directories left by a process that is gone.

    Returns how many. **Every prefix in :data:`RAW_SCRATCH_PREFIXES`**, not only the
    `.doc` one: each writer removes its own in a `finally`, so the way one survives is the
    server process itself being killed — a redeploy, an OOM, a `docker kill` — and every
    one of them holds somebody's uploaded document.

    Called at startup AND from the retention sweeper, because a restart is not the only
    moment a leftover exists: the process that made one can be killed at any point in the
    service's life, and a sweep that runs once leaves that directory forever.

    Age comes from the directory's own creation marker where there is one, and from mtime
    where there is not. mtime is refreshed by the document being written into it, so it
    measures the last write rather than the lifetime — the timestamp that bounds a
    lifetime has to be one the thing itself does not touch.

    **The count is all that is logged.** A path here is `<prefix><random>` and the files
    inside it are the visitor's; an operator needs to know that leftovers existed and how
    many, which is exactly what a count says.
    """
    root = Path(tempfile.gettempdir())
    cutoff = time.time() - max_age_seconds
    removed = 0
    candidates: list[Path] = []
    for prefix in RAW_SCRATCH_PREFIXES.values():
        try:
            candidates.extend(root.glob(prefix + "*"))
        except OSError:  # pragma: no cover - an unreadable temp root is not our outage
            return removed
    for path in candidates:
        try:
            if path.is_dir():
                born = scratch_creation_stamp(path)
                when = born if born is not None else path.stat().st_mtime
                if when > cutoff:
                    continue
                remove_raw_scratch(path)
                removed += not path.exists()
                continue
            # A marker whose directory is already gone. Since round 19's ruling 3 the
            # marker is a SIBLING, so a removal that got half way — or a `rmtree` by
            # something that did not know about it — leaves the file behind. It carries
            # no upload, but litter in the temp root is how the next reader loses track
            # of which directories are real.
            if path.name.endswith(CREATED_MARKER):
                directory = Path(str(path)[: -len(CREATED_MARKER)])
                if directory.exists():
                    continue
                if (_stamp_in(path) or path.stat().st_mtime) > cutoff:
                    continue
                path.unlink(missing_ok=True)
        except OSError:
            continue
    if removed:
        log.warning(
            "parse: removed %d raw-upload scratch director%s left by an earlier process",
            removed, "y" if removed == 1 else "ies",
        )
    return removed


def child_launch(argv: list[str]) -> "tuple[list[str], dict]":
    """The argv to exec and the ``Popen`` keywords that drop this child's privileges.

    **No ``preexec_fn``, anywhere.** Python documents it as unsafe in a program that uses
    threads, and this one launches every parser from a worker or intake thread: the hook
    runs in the forked child, between fork and exec, in a process that inherited the
    parent's locks without the threads that would release them. A hook that blocks there
    deadlocks the child *before exec*, ``Popen`` never returns, and the wall-clock loop
    that exists to kill a stuck parser is code that has not started yet — so the request
    and its worker slot hang for as long as the process lives.

    Both jobs it used to do are done exec-safely instead:

    * the uid drop is ``Popen``'s own ``user`` / ``group`` / ``extra_groups``, applied by
      CPython in C after the fork with no interpreter state involved (and ``umask`` while
      we are here, so a compromised parser cannot leave group- or world-readable files);
    * the network namespace is the ``unshare`` binary, which is simply the thing exec'd.

    Neither is attempted where the probe said it does not apply, because a launch that
    fails is a parse that fails.
    """
    state = isolation_report()
    plan = _launch_plan(state)
    # Recorded so `isolation_report()` — and therefore the startup line — names the path
    # actually taken rather than the one the probe hoped for.
    _ISOLATION["launch"] = plan
    kwargs: dict = {}
    identity = _parser_identity()
    if plan == LAUNCH_NAMESPACED_DROP and identity is not None:
        uid, gid = identity
        # One exec chain, root the whole way to setpriv: unshare needs CAP_SYS_ADMIN, so
        # it has to run BEFORE the drop, not after it. `umask` is still Popen's, because
        # it is applied in the forked child before any of this execs and setpriv has no
        # equivalent.
        argv = [
            shutil.which(UNSHARE_BINARY), "-n", "--",
            shutil.which(SETPRIV_BINARY),
            f"--reuid={uid}", f"--regid={gid}", "--clear-groups", "--",
            *argv,
        ]
        kwargs["umask"] = 0o077
    elif plan == LAUNCH_NAMESPACED:
        argv = [shutil.which(UNSHARE_BINARY), "-n", "--", *argv]
    elif plan == LAUNCH_DROP and identity is not None:
        uid, gid = identity
        kwargs.update({"user": uid, "group": gid, "extra_groups": [], "umask": 0o077})
    return argv, kwargs


#: Every child gets its OWN session, and therefore its own process group.
#:
#: `subprocess.Popen(start_new_session=True)` calls `setsid` in the forked child before
#: exec — CPython's own C code, not a `preexec_fn`, so the rule at the top of
#: `child_launch` still holds. It is what makes `kill_process_group` able to reach a
#: DESCENDANT: a `.doc` child launches antiword, and `process.kill()` reached only the
#: Python process, leaving antiword running with the visitor's document open after the
#: run was cancelled.
NEW_SESSION = True


def kill_process_group(process) -> bool:
    """SIGKILL the child's whole process group. True when the GROUP was signalled.

    False means the group kill was not safe or not possible and the caller's single-
    process kill is what happened.

    **The guard is the important half.** A child that is somehow NOT in a session of its
    own shares this server's process group, and `killpg` on that group kills the server —
    turning a parser timeout into an outage. So the group is signalled only when it is
    demonstrably the child's own, which `start_new_session` makes true (pgid == pid) and
    which is checked rather than assumed.
    """
    if not hasattr(os, "getpgid") or not hasattr(os, "killpg"):
        return False  # Windows: no process groups to signal; the caller's process.kill() stands
    try:
        group = os.getpgid(process.pid)
    except (ProcessLookupError, PermissionError, OSError):
        return False
    if group != process.pid or group == os.getpgid(0):
        # Not a group this child owns — or, impossibly, ours. Never signal it.
        return False
    try:
        os.killpg(group, signal.SIGKILL)
    except (ProcessLookupError, PermissionError, OSError):
        return False
    return True


#: Return code meaning "the run this parse belonged to was cancelled".
CANCELLED_CODE = -15
#: Return code meaning "the child produced more than the parent will hold".
OVERFLOW_CODE = -8


def _drain(stream, cap: int, keep: list, counter: list, over: threading.Event) -> None:
    """Read a pipe to the end, keeping at most ``cap`` bytes of it.

    The counting continues past the cap so the caller can tell that it overflowed, and the
    keeping does not — which is the whole point. ``communicate()`` accumulates everything
    the child writes and hands it over in one string, so the cap that came after it was a
    check on an allocation the parent had already made.
    """
    try:
        while True:
            chunk = stream.read(65536)
            if not chunk:
                return
            counter[0] += len(chunk)
            if counter[0] > cap:
                if not over.is_set():
                    over.set()
                return
            keep.append(chunk)
    except (ValueError, OSError):
        return


def run_child(
    argv: list[str],
    payload: str,
    timeout: float = DEFAULT_WALL_SECONDS,
    env: dict | None = None,
    cancel: "threading.Event | None" = None,
    stdin_bytes: bytes | None = None,
) -> tuple[int, str, str]:
    """Run a child to completion or kill it, and say which happened.

    Split out from :func:`parse_file` so the kill path is the same code in production and
    in the test that proves it works. ``kill()`` rather than ``terminate()``: a process
    that is ignoring the clock is not one to ask politely. A return code of ``-9`` means
    the wall clock did it, ``-15`` that the run was cancelled underneath it.

    ``cancel`` is what makes the run deadline real rather than advisory. Without it a
    timed-out run still had a subprocess of its own chewing through a stranger's
    spreadsheet for another 45 seconds, on a container the next visitor is waiting for.
    The pump thread exists because ``communicate()`` cannot be interrupted; polling a
    short join is the only way to watch two conditions at once.
    """
    # Binary stdin, because the document itself travels this way — the child is given no
    # path into any directory the server owns, so a compromised parser has nothing to
    # read. Text is decoded on the way out instead.
    launch_argv, launch_kwargs = child_launch(argv)
    process = subprocess.Popen(
        launch_argv,
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        env=env,
        # Its own session, so `kill_process_group` can take the whole tree — see
        # `NEW_SESSION`.
        start_new_session=NEW_SESSION,
        **launch_kwargs,
    )
    out_keep: list = []
    err_keep: list = []
    out_count = [0]
    err_count = [0]
    overflow = threading.Event()

    try:
        if payload:
            process.stdin.write(payload.encode("utf-8") if isinstance(payload, str)
                                else payload)
        if stdin_bytes:
            process.stdin.write(stdin_bytes)
        process.stdin.close()
    except (BrokenPipeError, ValueError, OSError):
        pass

    readers = [
        threading.Thread(
            target=_drain,
            args=(process.stdout, MAX_RESULT_BYTES, out_keep, out_count, overflow),
            name="parse-stdout", daemon=True,
        ),
        threading.Thread(
            target=_drain,
            args=(process.stderr, MAX_STDERR_BYTES, err_keep, err_count,
                  threading.Event()),
            name="parse-stderr", daemon=True,
        ),
    ]
    for reader in readers:
        reader.start()

    def collected() -> tuple[str, str]:
        def text(parts: list) -> str:
            joined = b"".join(p if isinstance(p, bytes) else p.encode("utf-8", "replace")
                              for p in parts)
            return joined.decode("utf-8", "replace")

        return text(out_keep), text(err_keep)

    def stop(code: int, why: str) -> tuple[int, str, str]:
        # The GROUP, not just the process. A `.doc` child launches antiword; killing only
        # the Python process left antiword running with the visitor's document open, and
        # it also held this pipe's write end open — so the drain below blocked to its full
        # join and the parent collected nothing.
        grouped = kill_process_group(process)
        process.kill()
        try:
            # Reap the direct child. Anything else in the group is dead and is reaped by
            # init; this is the one process whose exit status is ours to collect.
            process.wait(timeout=10)
        except subprocess.TimeoutExpired:  # pragma: no cover - a SIGKILL that did not land
            log.error("parse: the killed child did not exit")
        for reader in readers:
            reader.join(5)
        log.warning("parse: child killed (%s) — %s",
                    "process group" if grouped else "single process", why)
        out, err = collected()
        return code, out, err

    deadline = time.monotonic() + timeout
    while any(r.is_alive() for r in readers) or process.poll() is None:
        if overflow.is_set():
            return stop(
                OVERFLOW_CODE,
                f"it wrote more than the {MAX_RESULT_BYTES:,}-byte result ceiling",
            )
        if cancel is not None and cancel.is_set():
            return stop(CANCELLED_CODE, "its run was cancelled")
        if time.monotonic() >= deadline:
            return stop(-9, f"it passed its {timeout:.0f}s wall clock")
        time.sleep(0.05)

    for reader in readers:
        reader.join(5)
    if overflow.is_set():
        return stop(
            OVERFLOW_CODE, f"it wrote more than the {MAX_RESULT_BYTES:,}-byte ceiling"
        )
    out, err = collected()
    return process.returncode, out, err


def parse_file(
    path: Path,
    kind: str,
    filename: str | None = None,
    timeout: float = DEFAULT_WALL_SECONDS,
    env: dict | None = None,
    cancel: "threading.Event | None" = None,
    run_id: str | None = None,
) -> dict:
    """Parse one file in a subprocess. Raises :class:`ParseFailure` on any refusal.

    The document travels ON STDIN — a job line, then its bytes — rather than as a path the
    child opens, and for every kind but ``.doc`` it is never written down at all: the
    parsers read it out of memory, so a compromised parser has no path to any other run's
    document because none exists.

    ``.doc`` is the exception, because antiword takes a path. Those parses hold `DOC_LOCK`
    for their whole duration, so at most one scratch directory exists process-wide.

    ``run_id`` names the file the child writes on that one path, alongside ``kind``.
    ``filename`` is what the visitor called it and is metadata only: it reaches the
    ``documents`` list and the refusal messages, never a path.
    """
    data = Path(path).read_bytes()
    child_env = child_environment() if env is None else dict(env)
    process_argv = [sys.executable, "-m", "server.parse_worker"]
    job_fields = {
        "kind": kind,
        "filename": filename or Path(path).name,
        "bytes": len(data),
        "run": run_id or "",
    }

    if kind not in KINDS_NEEDING_FILE:
        # NOTHING on disk. The child gets the bytes and no path into any directory this
        # server owns, so a compromised parser has nothing of anyone else's to open — not
        # even another concurrent parse's, which one shared `demo-parser` uid used to make
        # readable across visitors.
        code, out, err = run_child(
            process_argv, json.dumps(job_fields) + "\n", timeout, child_env, cancel,
            stdin_bytes=data,
        )
    else:
        # antiword needs a real file. One at a time, process-wide, with the directory made
        # and removed inside the lock — so two visitors' scratch directories never exist
        # at the same moment and there is nothing for a compromised parser to enumerate.
        #
        # The parent makes it and the parent removes it: the child is SIGKILLed on three
        # ordinary paths — the wall clock, the result ceiling, a cancelled run — and a
        # killed process runs no `finally`.
        with DOC_LOCK:
            scratch = make_scratch()
            try:
                code, out, err = run_child(
                    process_argv,
                    json.dumps({**job_fields, "scratch": str(scratch)}) + "\n",
                    timeout, child_env, cancel, stdin_bytes=data,
                )
            finally:
                # After the reap, on every path this can leave by. `run_child` does not
                # return until the child is dead and its readers are joined, so nothing is
                # writing in here by the time it goes. Through the one remover, so the
                # sibling marker goes with it (round 19, ruling 3).
                remove_raw_scratch(scratch)
    if code == CANCELLED_CODE:
        raise ParseFailure(
            "run_cancelled", "This run was stopped before it finished reading."
        )
    if code == -9:
        raise ParseFailure(
            "parse_timeout",
            f"{filename or path.name} took too long to read and was stopped.",
        )
    if code == OVERFLOW_CODE or len(out) > MAX_RESULT_BYTES:
        log.error("parse: child output passed the %d-byte ceiling", MAX_RESULT_BYTES)
        raise ParseFailure(
            "file_too_complex",
            f"{filename or path.name} holds more data than this demo will read.",
        )
    if not out.strip():
        # Counts and the exit code, never the streams. A parser's diagnostics quote the
        # document — openpyxl names the cell it choked on, antiword echoes what it read —
        # so `err[-500:]` was up to half a kilobyte of a stranger's rent roll written to a
        # log that outlives the one-hour retention the page promises.
        log.error(
            "parse: child produced no output (exit %s, %d stdout byte(s), "
            "%d stderr byte(s))", code, len(out.encode("utf-8", "replace")),
            len(err.encode("utf-8", "replace")),
        )
        raise ParseFailure("unreadable_file", f"{filename or path.name} could not be read.")
    try:
        result = json.loads(out)
    except ValueError as exc:
        # A partial success starts with normalised DOCUMENT TEXT, so the first 300
        # characters of it are the worst possible thing to keep. The category of the JSON
        # error and how much arrived are what an operator can act on.
        log.error(
            "parse: child output was not json (%s at position %s, %d byte(s) received)",
            type(exc).__name__, getattr(exc, "pos", "?"),
            len(out.encode("utf-8", "replace")),
        )
        raise ParseFailure(
            "unreadable_file", f"{filename or path.name} could not be read."
        ) from exc
    if not result.get("ok"):
        raise ParseFailure(
            result.get("code", "unreadable_file"),
            result.get("message", "That file could not be read."),
        )
    return result["document"]


if __name__ == "__main__":
    raise SystemExit(child_main())
