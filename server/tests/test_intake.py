"""The intake gate, at the unit level.

Each refusal is checked against the gate function and the code it raises with. (The hosted
demo also checks every refusal through its HTTP route; that service layer is not part of
this repository.)

The zip bomb gets the strictest check in the file. It is not enough that it is refused —
it has to be refused **without the payload being decompressed and without openpyxl being
invoked**, so the test replaces both ``zipfile.ZipFile.read`` and ``openpyxl.load_workbook``
with recording sentinels and asserts on what they saw. Refusing a bomb after unpacking it
is losing politely.

Usage: python server/tests/test_intake.py
"""

from __future__ import annotations

import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from _harness import Checker, guard

FIXTURES = Path(__file__).resolve().parent / "fixtures"


def load(name: str) -> bytes:
    return (FIXTURES / name).read_bytes()


def refusal(c: Checker, label: str, fn, status: int, code: str) -> None:
    """Call something that must refuse, and check both halves of the refusal."""
    from server.errors import ApiError

    try:
        fn()
    except ApiError as exc:
        c.equal(f"{label} -> {status}", exc.status, status)
        c.equal(f"{label} code", exc.code, code)
        c.ok(
            f"{label} message is a plain sentence",
            exc.message.strip().endswith((".", "!", "?")) and "Traceback" not in exc.message,
            repr(exc.message),
        )
        return
    c.ok(f"{label} -> {status}", False, "nothing was raised")


def main() -> int:
    c = Checker("test_intake")
    from server import intake
    from server.errors import ApiError

    # ---- batch caps, before any content ------------------------------------------
    c.section("batch caps")
    refusal(c, "six files", lambda: intake.check_batch(6, ""), 400, "too_many_files")
    refusal(c, "zero files", lambda: intake.check_batch(0, ""), 400, "no_files")
    refusal(
        c,
        "a 2,001-character note",
        lambda: intake.check_batch(1, "x" * 2001),
        400,
        "notes_too_long",
    )
    try:
        intake.check_batch(5, "x" * 2000)
        c.ok("five files and a 2,000-character note are fine", True)
    except ApiError as exc:
        c.ok("five files and a 2,000-character note are fine", False, exc.message)

    refusal(
        c,
        "a 26 MB file",
        lambda: intake.check_file_size("big.pdf", 26 * 1024 * 1024),
        413,
        "file_too_large",
    )
    refusal(
        c,
        "41 MB in total",
        lambda: intake.check_total_size(41 * 1024 * 1024),
        413,
        "batch_too_large",
    )

    # ---- what the bytes say ------------------------------------------------------
    c.section("type is decided by the bytes, never the name")
    for name, kind in (
        ("om.pdf", "pdf"),
        ("om.docx", "docx"),
        ("t12.xlsx", "xlsx"),
        # A compound file sniffs as the CONTAINER since an earlier review: which of the two
        # legacy Office formats it is takes a parser, and that runs in the sandboxed
        # child (`parse_worker.probe_ole`), never in the API process. Same split the PDF
        # page count already had.
        ("rent-roll.xls", "ole"),
        ("t12.csv", "csv"),
    ):
        c.equal(f"{name} sniffs as {kind}", intake.sniff(load(name), name), kind)

    c.equal(
        "a docx renamed .xlsx is still a docx",
        intake.sniff(load("om.docx"), "actually-a-spreadsheet.xlsx"),
        "docx",
    )
    refusal(
        c,
        "a bare .zip",
        lambda: intake.sniff(load("plain.zip"), "plain.zip"),
        415,
        "unsupported_type",
    )
    refusal(
        c,
        "a png",
        lambda: intake.sniff(load("not-a-document.png"), "not-a-document.png"),
        415,
        "unsupported_type",
    )
    refusal(
        c,
        "an .xlsm renamed .xlsx",
        lambda: intake.sniff(
            load("macro-workbook-named-xlsx.xlsx"), "macro-workbook-named-xlsx.xlsx"
        ),
        415,
        "macros_not_accepted",
    )

    # ---- pdf ---------------------------------------------------------------------
    # The page and encryption facts come from the sandboxed child (parse_worker's
    # probe_pdf job). What is checked here is the decision made from that report — the
    # parsing itself is exercised in test_parsers.py.
    c.section("pdf, from the child's report")
    import tempfile as _tf

    from server.parse_worker import parse_file

    with _tf.TemporaryDirectory() as tmpdir:
        probe = Path(tmpdir) / "om.pdf"
        probe.write_bytes(load("om.pdf"))
        report = parse_file(probe, "probe_pdf", "om.pdf")["stats"]
    c.equal("the child reports the sample OM's page count", report["pages"], 3)
    c.equal("and that it is not encrypted", report["encrypted"], False)
    c.equal("which the gate accepts", intake.check_pdf_report(report, "om.pdf"), 3)
    refusal(
        c,
        "an encrypted pdf",
        lambda: intake.check_pdf_report({"encrypted": True, "pages": 0}, "encrypted.pdf"),
        422,
        "encrypted_file",
    )
    refusal(
        c,
        "a 151-page pdf",
        lambda: intake.check_pdf_report({"encrypted": False, "pages": 151}, "big.pdf"),
        413,
        "too_many_pages",
    )

    # ---- the zip bomb, and what must NOT have happened ---------------------------
    c.section("the zip bomb is refused from the central directory alone")
    bomb = load("zip-bomb.xlsx")
    with zipfile.ZipFile(Path(FIXTURES / "zip-bomb.xlsx")) as archive:
        declared = sum(i.file_size for i in archive.infolist())
    c.ok(
        "the fixture really does declare a huge expansion",
        declared > 250 * 1024 * 1024,
        f"declares {declared / 1024 / 1024:.0f} MB",
    )

    import openpyxl

    read_names: list[str] = []
    openpyxl_calls: list[tuple] = []
    real_read = zipfile.ZipFile.read
    real_load = openpyxl.load_workbook

    def recording_read(self, name, pwd=None):
        read_names.append(getattr(name, "filename", name))
        return real_read(self, name, pwd)

    def recording_load(*args, **kwargs):
        openpyxl_calls.append((args, kwargs))
        return real_load(*args, **kwargs)

    zipfile.ZipFile.read = recording_read
    openpyxl.load_workbook = recording_load
    try:
        refusal(
            c,
            "the zip bomb",
            lambda: intake.sniff(bomb, "zip-bomb.xlsx"),
            413,
            "archive_expands_too_far",
        )
    finally:
        zipfile.ZipFile.read = real_read
        openpyxl.load_workbook = real_load

    c.equal("openpyxl was never invoked", len(openpyxl_calls), 0)
    c.ok(
        "the bomb's payload was never decompressed",
        "xl/worksheets/sheet1.xml" not in read_names,
        f"entries read: {read_names}",
    )
    c.ok(
        "in fact nothing at all was read from it",
        read_names == [],
        f"entries read: {read_names}",
    )

    return c.finish()


if __name__ == "__main__":
    raise SystemExit(guard(main))
