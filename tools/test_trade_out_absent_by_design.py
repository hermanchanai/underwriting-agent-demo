"""Trade-out is absent by design: no dead writer or preview path stays behind.

The demo never ingests a lease trade-out report —
``server/assemble.py`` always states ``NO_TRADE_OUT``, and ``policy.py`` /
``restage_fixture.py`` empty the block for the public samples — so the trade-out sheet
writer and its HTML fragment were reachable only from tests. Delete them and keep the
stated absence. ``compute.trade_out`` stays: it is the parity-pinned math
(``tools/test_rentroll_parity.py``) an uploaded report would be wired through.

Usage: python tools/test_trade_out_absent_by_design.py
"""

from __future__ import annotations

import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
for p in (str(REPO), str(HERE)):
    if p not in sys.path:
        sys.path.insert(0, p)

from _testkit import Checker, guard  # noqa: E402

DEAD = ("_write_trade_out_tab", "trade_out_fragment", "SHEET_TRADE")


def main() -> int:
    import json

    import render_previews
    import schema
    import write_rent_roll
    from openpyxl import load_workbook

    c = Checker("trade-out: absent by design, no dead paths")

    c.section("the dead writer and preview paths are gone")
    c.ok("write_rent_roll has no trade-out sheet writer",
         not hasattr(write_rent_roll, "_write_trade_out_tab"))
    c.ok("write_rent_roll has no trade-out sheet name",
         not hasattr(write_rent_roll, "SHEET_TRADE"))
    c.ok("render_previews has no trade-out fragment",
         not hasattr(render_previews, "trade_out_fragment"))
    shipped = [*sorted((REPO / "tools").glob("*.py")), *sorted((REPO / "server").glob("*.py"))]
    hits = [f"{p.relative_to(REPO)}: {name}" for p in shipped
            if not p.name.startswith("test_") and p.name != Path(__file__).name
            for name in DEAD if name in p.read_text(encoding="utf-8")]
    c.equal("no shipped module names a dead path", hits, [])
    c.ok("no preview file is a trade-out tab", "trade-out.html" not in schema.PREVIEW_FILES)

    c.section("the formats doc no longer promises a trade-out tab or preview")
    formats = (REPO / "formats" / "FORMATS.md").read_text(encoding="utf-8")
    c.ok("FORMATS.md has no `Lease Trade-Out` tab section", "Tab `Lease Trade-Out`" not in formats)
    c.ok("FORMATS.md lists no trade-out.html preview", "trade-out.html" not in formats)

    c.section("the absence is still stated, by every live path")
    from server import assemble

    src = (REPO / "server" / "assemble.py").read_text(encoding="utf-8")
    c.ok("assemble still emits rows=[] with NO_TRADE_OUT",
         '"trade_out": {"rows": [], "absent_reason": NO_TRADE_OUT}' in src)
    c.ok("and NO_TRADE_OUT says why", bool(str(assemble.NO_TRADE_OUT).strip()))

    c.section("a deal that DOES carry trade-out rows still writes no trade-out sheet")
    deal = json.loads((HERE / "fixtures" / "synthetic_full.json").read_text())
    deal["trade_out"] = {"rows": [{"cohort": "New", "leases": 3, "prior_rent": 1500.0,
                                   "new_rent_gross": 1520.0, "new_rent_effective": 1440.0}],
                         "absent_reason": None}
    with tempfile.TemporaryDirectory() as tmp:
        out = write_rent_roll.build(deal, Path(tmp) / "rr.xlsx")
        names = load_workbook(out).sheetnames
    c.ok("no sheet is a trade-out sheet", not any("trade" in n.lower() for n in names), names)
    return c.finish()


if __name__ == "__main__":
    sys.exit(guard(main))
