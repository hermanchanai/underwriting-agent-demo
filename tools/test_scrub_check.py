"""Tests for the publish gate — including the red leg, which is the point.

A gate nobody has watched fail is not a gate. So this file seeds a copy of a real fixture
output with contamination in the three places a leak actually hides:

1. a visible cell value,
2. a sheet name,
3. the workbook's ``creator`` property — invisible on screen, and the one a human
   reviewer cannot catch by looking,

then asserts the gate exits 1 and that its report names all three. Only after that does it
check the clean file exits 0.

Also checks a leaked absolute path, a contaminated filename, a contaminated HTML fragment,
the extra-terms file, and that ordinary English containing a token is not flagged.

Usage: python tools/test_scrub_check.py
"""

from __future__ import annotations

import shutil
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from openpyxl import load_workbook

import json as _json
import scrub_check
from _testkit import Checker, guard

PROJECT = HERE.parent


def _fixture_set() -> tuple[Path, str]:
    """The synthetic fixture deal written through every writer and the preview renderer."""
    import render_previews
    import write_deal_summary
    import write_manifest
    import write_rent_roll
    import write_t12_analysis

    built = Path(tempfile.mkdtemp(prefix="scrub_fixture_"))
    source = HERE / "fixtures" / "synthetic_full.json"
    deal = _json.loads(source.read_text(encoding="utf-8"))
    write_deal_summary.build(deal, built / "sample-deal-summary.xlsx")
    write_t12_analysis.build(deal, built / "sample-t12-analysis.xlsx")
    write_rent_roll.build(deal, built / "sample-rent-roll-analysis.xlsx")
    write_manifest.build(built / "sample-model-ready-outputs.xlsx", deal)
    render_previews.render_all(deal, built / "previews")
    shutil.copy2(source, built / "deal_outputs.json")
    return built, "built from tools/fixtures/synthetic_full.json"


FIXTURE_DIR, FIXTURE_ORIGIN = _fixture_set()
FIXTURE_XLSX = FIXTURE_DIR / "sample-deal-summary.xlsx"

SEEDED_TERM = "Acme"
SEEDED_PATH = "/Users/jdoe/Documents/secret-deal"


def _seed_contaminated(src: Path, dest_dir: Path) -> Path:
    """Copy a clean output and plant the term in a cell, a sheet name, and the creator."""
    dest = dest_dir / "seeded-deal-summary.xlsx"
    shutil.copy2(src, dest)
    wb = load_workbook(dest)
    ws = wb.worksheets[0]
    ws["E2"] = f"Prepared for {SEEDED_TERM} Capital"
    ws.title = f"{SEEDED_TERM} Summary"
    wb.properties.creator = f"{SEEDED_TERM} Analyst"
    wb.save(dest)
    return dest


def main() -> int:
    c = Checker("test_scrub_check")

    if not FIXTURE_XLSX.exists():
        c.ok(
            "fixture output present",
            False,
            f"{FIXTURE_XLSX} missing — run the writers first",
        )
        return c.finish()
    c.ok(f"fixture output present ({FIXTURE_ORIGIN})", True)

    terms = scrub_check.load_terms([])

    with tempfile.TemporaryDirectory() as tmp:
        tmpdir = Path(tmp)

        # ---- RED LEG ---------------------------------------------------------------
        c.section("red leg: seeded contamination must be caught")
        seeded = _seed_contaminated(FIXTURE_XLSX, tmpdir)
        findings = scrub_check.run([seeded], terms)
        rendered = "\n".join(f.render() for f in findings)

        c.ok("red leg produced findings", bool(findings), f"{len(findings)} findings")
        c.ok(
            "names the contaminated CELL",
            any("cell" in f.location for f in findings),
            rendered,
        )
        c.ok(
            "names the contaminated SHEET NAME",
            any("sheet name" in f.location for f in findings),
            rendered,
        )
        c.ok(
            "names the contaminated CREATOR property",
            any("document property 'creator'" in f.location for f in findings),
            rendered,
        )
        exit_code = scrub_check.main([str(seeded), "--quiet"])
        c.equal("red leg exit code is 1", exit_code, 1)

        # ---- a leaked absolute path -------------------------------------------------
        c.section("red leg: leaked local path")
        path_file = tmpdir / "fragment.html"
        path_file.write_text(
            f"<p>Built from <code>{SEEDED_PATH}</code></p>", encoding="utf-8"
        )
        path_findings = scrub_check.run([path_file], terms)
        c.ok(
            "flags a leaked /Users path",
            any(f.term == "/users/" for f in path_findings),
            "\n".join(f.render() for f in path_findings),
        )
        c.ok(
            "reports the line number for a text hit",
            any(f.location.startswith("line ") for f in path_findings),
        )

        # ---- a contaminated filename ------------------------------------------------
        c.section("red leg: contaminated filename")
        named = tmpdir / f"{SEEDED_TERM.lower()}-export.json"
        named.write_text("{}", encoding="utf-8")
        name_findings = scrub_check.run([named], terms)
        c.ok(
            "flags the filename itself",
            any(f.location == "file path" for f in name_findings),
            "\n".join(f.render() for f in name_findings),
        )

        # ---- extra terms file --------------------------------------------------------
        c.section("extra terms from a deal's scrub_terms.txt")
        terms_file = tmpdir / "scrub_terms.txt"
        terms_file.write_text(
            "# broker identifiers for this deal\nOakmont Partners\n\n", encoding="utf-8"
        )
        extra = scrub_check.load_terms([terms_file])
        c.ok("extra term loaded", "Oakmont Partners" in extra)
        c.ok("comment line ignored", not any(t.startswith("#") for t in extra))
        broker_file = tmpdir / "broker.html"
        broker_file.write_text("<p>Listed by Oakmont Partners.</p>", encoding="utf-8")
        c.ok(
            "extra term is enforced",
            bool(scrub_check.run([broker_file], extra)),
        )
        c.ok(
            "extra term is not enforced when not supplied",
            not scrub_check.run([broker_file], terms),
        )

        # ---- benign English ----------------------------------------------------------
        c.section("benign English is not flagged")
        benign = tmpdir / "prose.html"
        benign.write_text(
            "<p>This layout is the acme of plain tables.</p>",
            encoding="utf-8",
        )
        benign_findings = scrub_check.run([benign], terms)
        c.ok(
            "\"the acme of\" does not trip the gate",
            not benign_findings,
            "\n".join(f.render() for f in benign_findings),
        )
        still = tmpdir / "prose2.html"
        still.write_text("<p>Prepared by Acme.</p>", encoding="utf-8")
        c.ok(
            "the bare token still trips the gate",
            bool(scrub_check.run([still], terms)),
        )

        # ---- RED LEG: a term hidden inside a provenance quote -----------------------
        c.section("red leg: contamination inside an audit-trail quote")
        # The trail carries verbatim text lifted out of source documents, and an offering
        # memorandum's prose is exactly where a brokerage or an owner entity is named. A
        # gate that scans the numbers but not the evidence behind them is a gate with a
        # hole the size of the trail.
        import json as _json

        import write_deal_summary as _wds

        deal = _json.loads(
            (HERE / "fixtures" / "synthetic_full.json").read_text(encoding="utf-8")
        )
        deal["provenance"]["summary.deal_name"] = {
            "source_doc": "Synthetic offering memorandum",
            "locator": "p.1",
            "method": "read",
            "quote": f"Offered exclusively by {SEEDED_TERM} Capital Markets",
        }
        trail_path = _wds.build(deal, tmpdir / "trail-seeded.xlsx")
        trail_findings = scrub_check.run([trail_path], terms)
        c.ok(
            "the gate reads the Audit Trail tab",
            any("Audit Trail" in f.location for f in trail_findings),
            "\n".join(f.render() for f in trail_findings) or "no findings at all",
        )
        c.ok(
            "and names the seeded term",
            any(f.term.lower() == SEEDED_TERM.lower() for f in trail_findings),
            "\n".join(f.render() for f in trail_findings),
        )
        c.equal(
            "exit code is 1 for a contaminated trail",
            scrub_check.main([str(trail_path), "--quiet"]),
            1,
        )
        clean_trail = _wds.build(
            _json.loads((HERE / "fixtures" / "synthetic_full.json").read_text(encoding="utf-8")),
            tmpdir / "trail-clean.xlsx",
        )
        c.equal(
            "and 0 once the quote is clean again",
            scrub_check.main([str(clean_trail), "--quiet"]),
            0,
        )

    # ---- GREEN LEG -------------------------------------------------------------------
    # Two forms, on purpose. The directory walk is what a person types; the explicit file
    # list is what a publishing step gates on, and it does not benefit from any
    # directory exclusion.
    c.section("green leg: the real outputs are clean")
    fixture_dir = FIXTURE_DIR
    clean_findings = scrub_check.run([fixture_dir], terms)
    c.ok(
        "fixture output set is clean (directory walk)",
        not clean_findings,
        "\n".join(f.render() for f in clean_findings),
    )
    c.equal(
        "green leg exit code is 0",
        scrub_check.main([str(fixture_dir), "--quiet"]),
        0,
    )

    artifacts = sorted(fixture_dir.glob("*.xlsx")) + sorted(
        (fixture_dir / "previews").glob("*.html")
    )
    c.ok("publishable artifacts found", bool(artifacts), f"{len(artifacts)} files")
    explicit_findings = scrub_check.run(artifacts, terms)
    c.ok(
        "every publishable artifact is clean (explicit file list)",
        not explicit_findings,
        "\n".join(f.render() for f in explicit_findings),
    )

    c.section("the evidence directory is excluded from a directory walk, not from a file list")
    c.ok("qa is excluded", "qa" in scrub_check.EXCLUDED_DIR_NAMES)
    qa_dir = fixture_dir / "qa"
    if qa_dir.exists():
        walked = scrub_check.iter_files([fixture_dir])
        c.ok(
            "directory walk skips qa/",
            not any("qa" in p.relative_to(fixture_dir).parts[:-1] for p in walked),
        )
        # Plant a known term rather than relying on whichever captured red run happens
        # to quote one: the invariant under test is that naming a file explicitly
        # bypasses the directory exclusion, not what today's evidence files contain.
        planted = qa_dir / "_scrub_gate_selftest.txt"
        try:
            planted.write_text(f"seeded {SEEDED_TERM} for the gate self-test\n", encoding="utf-8")
            c.ok(
                "a directory walk still skips it",
                planted not in scrub_check.iter_files([fixture_dir]),
            )
            c.ok(
                "but naming it explicitly scans it",
                bool(scrub_check.run([planted], terms)),
            )
        finally:
            planted.unlink(missing_ok=True)

    c.section("gate coverage")
    c.ok("checks xlsx", ".xlsx" in scrub_check.XLSX_SUFFIXES)
    c.ok("checks html", ".html" in scrub_check.TEXT_SUFFIXES)
    c.ok("checks json", ".json" in scrub_check.TEXT_SUFFIXES)
    c.ok(
        "denylist covers brand and path families",
        len(scrub_check.BRAND_TERMS) >= 3 and len(scrub_check.PATH_TERMS) >= 3,
    )

    return c.finish()


if __name__ == "__main__":
    raise SystemExit(guard(main))
