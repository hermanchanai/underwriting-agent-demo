"""Tests for the four xlsx writers, read back by LABEL rather than by cell address.

Every assertion locates its block by searching column A for the label, so the layout can
move without the suite going red for the wrong reason — and so a block that silently
stopped being written fails loudly instead of passing because row 17 still holds a number.

The formula assertions are the substance. Three of them exist because the arithmetic they
pin is where a T-12 rollup actually goes wrong:

* Total Revenue must sum from Gross Scheduled Rent down, never from Gross Potential Rent
  (which double-counts loss-to-lease);
* Gross Scheduled Rent must carry the printed-else-derived fallback;
* the per-unit column must be blank-guarded, so a zero unit count cannot fill it with
  errors.

Also covers: no denylisted term anywhere, no tenant-name column, a text date surviving as
text, contra rows staying negative, and the manifest containing every inventory field
exactly once.

Usage: python tools/test_writers.py
"""

from __future__ import annotations

import ast
import copy
import datetime as dt
import hashlib
import json
import re
import sys
import tempfile
import time
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from openpyxl import load_workbook

import compute
import schema
import scrub_check
import write_deal_summary
import write_manifest
import write_rent_roll
import write_t12_analysis
from _testkit import Checker, guard
from _xl import FMT_CURRENCY_CENTS, save_workbook

PROJECT = HERE.parent
#: The pre-2026-09-10 sample (fully synthetic, TX, every bed count covered, four supply
#: stages, trade-out rows) — the deal the writer checks below were written against. The
#: PUBLISHED sample is fixture_deal.json, pinned by test_published_sample.py.
FIXTURE = HERE / "fixtures" / "synthetic_full.json"
OUT_DIR = PROJECT / "outputs" / "fixture"


#: A nonzero-total test. `<>` compared against 0 is unambiguous: the correct existence
#: gate is spelled COUNTIF(...)>0, so it never produces this.
_NONZERO_GATE_RE = re.compile(r"<>\s*0")

#: A SUMIF written inline (not through the helper) and compared to 0 in the same literal.
_INLINE_SUMIF_GATE_RE = re.compile(r"sumif\s*\(.*?(?:<>|>|<|=)\s*0", re.IGNORECASE | re.DOTALL)


def _code_string_literals(source: str) -> list[tuple[int, str]]:
    """Every string literal that is *code* — the text that becomes a formula.

    Comments are not literals at all, and docstrings are excluded explicitly, so prose
    quoting a pattern cannot register as an offender.
    """
    tree = ast.parse(source)
    docstrings: set[int] = set()
    for node in ast.walk(tree):
        if isinstance(
            node, (ast.Module, ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)
        ):
            body = getattr(node, "body", [])
            if (
                body
                and isinstance(body[0], ast.Expr)
                and isinstance(body[0].value, ast.Constant)
                and isinstance(body[0].value.value, str)
            ):
                docstrings.add(id(body[0].value))
    out: list[tuple[int, str]] = []
    for node in ast.walk(tree):
        if (
            isinstance(node, ast.Constant)
            and isinstance(node.value, str)
            and id(node) not in docstrings
        ):
            out.append((node.lineno, node.value))
    return out


def aggregate_gate_offenders(paths: list[Path]) -> list[str]:
    """Source-sweep for a presence decision spelled as "the total is nonzero".

    Matches the **semantic shape** inside code string literals, not a literal token
    sequence. That distinction is the whole point: the real defect built its gate through
    a local ``sumif`` helper, so the source read
    ``f'IF({sumif(c, "Gross Scheduled Rent")}<>0,'`` — the ``SUMIF(`` token is generated at
    runtime and never appears next to the comparison. An uppercase ``SUMIF\\(...\\)<>0``
    regex therefore reported *nothing* against the faulty source, and its apparent hit came
    from a docstring that spelled the pattern out in prose. This version cannot be fooled
    by interpolation, by case, or by prose.
    """
    offenders: list[str] = []
    for path in paths:
        text = path.read_text(encoding="utf-8")
        try:
            literals = _code_string_literals(text)
        except SyntaxError as exc:  # a file we cannot parse is not a file we can clear
            offenders.append(f"{path.name}: unparseable ({exc})")
            continue
        for lineno, value in literals:
            if _NONZERO_GATE_RE.search(value):
                offenders.append(
                    f"{path.name}:{lineno}: nonzero-total gate in {value.strip()[:48]!r}"
                )
            elif _INLINE_SUMIF_GATE_RE.search(value):
                offenders.append(
                    f"{path.name}:{lineno}: inline SUMIF compared to 0 in "
                    f"{value.strip()[:48]!r}"
                )
    return offenders


def find_row(ws, label: str, col: int = 1, start: int = 1) -> int | None:
    """First row whose column-``col`` value equals ``label``, ignoring case and spacing."""
    target = " ".join(str(label).split()).lower()
    for r in range(start, ws.max_row + 1):
        value = ws.cell(row=r, column=col).value
        if value is None:
            continue
        if " ".join(str(value).split()).lower() == target:
            return r
    return None


def find_row_prefix(ws, prefix: str, col: int = 1) -> int | None:
    target = " ".join(str(prefix).split()).lower()
    for r in range(1, ws.max_row + 1):
        value = ws.cell(row=r, column=col).value
        if value is None:
            continue
        if " ".join(str(value).split()).lower().startswith(target):
            return r
    return None


def main() -> int:
    c = Checker("test_writers")
    deal = json.loads(FIXTURE.read_text(encoding="utf-8"))

    with tempfile.TemporaryDirectory() as tmp:
        tmpdir = Path(tmp)
        paths = {
            "deal_summary": write_deal_summary.build(
                deal, tmpdir / schema.OUTPUT_FILES["deal_summary"]
            ),
            "t12_analysis": write_t12_analysis.build(
                deal, tmpdir / schema.OUTPUT_FILES["t12_analysis"]
            ),
            "rent_roll": write_rent_roll.build(
                deal, tmpdir / schema.OUTPUT_FILES["rent_roll"]
            ),
            "manifest": write_manifest.build(
                tmpdir / schema.OUTPUT_FILES["manifest"], deal
            ),
        }

        c.section("all four workbooks build and reopen")
        for key, path in paths.items():
            c.ok(f"{key} written", path.exists() and path.stat().st_size > 4000)
        books = {k: load_workbook(p, data_only=False) for k, p in paths.items()}

        # ---- sheet inventory ---------------------------------------------------------
        c.section("sheet inventory")
        c.equal(
            "deal summary tabs",
            books["deal_summary"].sheetnames,
            ["Deal Summary", "Rent Comps", "Property Tax", "Demand", "Supply",
             "Notes & Disclosures", "Audit Trail"],
        )
        c.equal(
            "t12 tabs",
            books["t12_analysis"].sheetnames,
            ["T-12 Data", "P&L Summary", "Checks & Pricing", "Audit Trail"],
        )
        c.equal(
            "rent roll tabs",
            books["rent_roll"].sheetnames,
            ["Rent Roll", "Unit Mix Summary", "Audit Trail"],
        )
        c.equal("manifest tabs", books["manifest"].sheetnames, ["Model-Ready Outputs", "Audit Trail"])

        # ---- T-12: the formulas that matter ------------------------------------------
        c.section("P&L formulas — the arithmetic traps")
        pnl = books["t12_analysis"]["P&L Summary"]
        rows = {
            label: find_row(pnl, label)
            for label in (
                "Gross Potential Rent",
                "Gain / (Loss) to Lease",
                "Gross Scheduled Rent",
                "Revenue Adjustments",
                "Total Revenue",
                "Total Operating Expenses",
                "Net Operating Income",
                "Utilities",
            )
        }
        c.ok("every P&L label found", all(v for v in rows.values()), str(rows))

        total_rev_formula = pnl.cell(row=rows["Total Revenue"], column=2).value
        expected_span = f"=SUM(B{rows['Gross Scheduled Rent']}:B{rows['Revenue Adjustments']})"
        c.equal(
            "Total Revenue sums GSR..Revenue Adjustments", total_rev_formula, expected_span
        )
        c.ok(
            "Total Revenue does NOT start at Gross Potential Rent",
            f"B{rows['Gross Potential Rent']}:" not in str(total_rev_formula),
            f"formula was {total_rev_formula!r}",
        )

        gsr_formula = str(pnl.cell(row=rows["Gross Scheduled Rent"], column=2).value)
        c.ok("GSR uses an IF fallback", gsr_formula.startswith("=IF("), gsr_formula)
        c.ok(
            "gated on the row's EXISTENCE, not on a nonzero total",
            "COUNTIF(" in gsr_formula and "<>0" not in gsr_formula,
            gsr_formula,
        )
        g, l = rows['Gross Potential Rent'], rows['Gain / (Loss) to Lease']
        c.ok(
            # an absent input is "" (never 0), so each is read as 0 explicitly — "" + a
            # number would be #VALUE! — and the fallback is blank when neither is mapped.
            "GSR fallback adds GPR + GLL, each absent input read as 0, blank when neither is mapped",
            f'IF(B{g}="",0,B{g})+IF(B{l}="",0,B{l})' in gsr_formula
            and "COUNTIF('T-12 Data'!" in gsr_formula and gsr_formula.endswith(',""))'),
            gsr_formula,
        )

        noi_formula = pnl.cell(row=rows["Net Operating Income"], column=2).value
        c.equal(
            "NOI is revenue less opex",
            noi_formula,
            f"=B{rows['Total Revenue']}-B{rows['Total Operating Expenses']}",
        )
        opex_formula = str(pnl.cell(row=rows["Total Operating Expenses"], column=2).value)
        c.ok("opex total is a SUM over the expense block", opex_formula.startswith("=SUM(B"), opex_formula)

        per_unit = str(pnl.cell(row=rows["Payroll"] if rows.get("Payroll") else find_row(pnl, "Payroll"), column=4).value)
        c.ok("per-unit column is blank-guarded", per_unit.startswith("=IFERROR("), per_unit)
        c.ok(
            "per-unit divides by the unit count",
            f"/{deal['t12']['units']}" in per_unit,
            per_unit,
        )
        pct = str(pnl.cell(row=find_row(pnl, "Payroll"), column=5).value)
        c.ok(
            "% of revenue references the Total Revenue row",
            f"B{rows['Total Revenue']}" in pct and pct.startswith("=IFERROR("),
            pct,
        )
        util_sumif = str(pnl.cell(row=rows["Utilities"], column=2).value)
        u = rows['Utilities']
        c.ok(
            # keyed on its own label cell, and BLANK when no line carries the category.
            "a category row is a SUMIF keyed on its own label cell, blank when no line is mapped to it",
            util_sumif.startswith("=IF(COUNTIF(") and f",$A{u})>0,SUMIF(" in util_sumif
            and f"$A{u}" in util_sumif and util_sumif.endswith(',"")'),
            util_sumif,
        )

        # ---- T-12 data tab -----------------------------------------------------------
        c.section("T-12 data tab")
        data = books["t12_analysis"]["T-12 Data"]
        header_row_idx = find_row(data, "Line item")
        c.ok("data header found", header_row_idx is not None)
        headers = [data.cell(row=header_row_idx, column=i).value for i in range(1, 18)]
        c.equal("month columns start at C", headers[2], deal["t12"]["month_labels"][0])
        c.equal("T-12 total column present", headers[14], "T-12 total")
        c.equal("T-3 total column present", headers[15], "T-3 total")

        first_data = header_row_idx + 1
        c.equal(
            "T-12 total is a 12-month SUM",
            data.cell(row=first_data, column=15).value,
            f"=SUM(C{first_data}:N{first_data})",
        )
        t3_first, t3_last = compute.recent_span(deal["t12"]["month_labels"], 3)
        c.equal(
            "T-3 total sums the three most recent months",
            data.cell(row=first_data, column=16).value,
            f"=SUM({chr(ord('C') + t3_first)}{first_data}:{chr(ord('C') + t3_last)}{first_data})",
        )

        subtotal_row = find_row(data, "Total Revenue (per statement)")
        c.ok("the uncategorized source subtotal is still written", subtotal_row is not None)
        # openpyxl does not persist an empty string, so a blank cell reads back as None.
        # Either way the SUMIF has nothing to match, which is the actual requirement.
        c.ok(
            "and its category cell is blank, so nothing sums it twice",
            data.cell(row=subtotal_row, column=2).value in (None, ""),
            repr(data.cell(row=subtotal_row, column=2).value),
        )

        vac_row = find_row(data, "Vacancy Loss")
        c.ok(
            "a contra-revenue row stays negative",
            all(
                data.cell(row=vac_row, column=col).value < 0 for col in range(3, 15)
            ),
        )
        helper = find_row(data, write_t12_analysis.REVENUE_HELPER_LABEL)
        c.ok("the computed monthly-revenue row exists", helper is not None)
        c.ok(
            "and it is a formula, not a pasted number",
            str(data.cell(row=helper, column=3).value).startswith("=IF(COUNTIF("),
            str(data.cell(row=helper, column=3).value)[:80],
        )

        # ---- checks & pricing --------------------------------------------------------
        c.section("checks and pricing")
        checks = books["t12_analysis"]["Checks & Pricing"]
        for label in ("Total Revenue", "Total Operating Expenses", "Net Operating Income"):
            r = find_row(checks, label)
            c.ok(f"error check row for {label}", r is not None)
            if r:
                status = str(checks.cell(row=r, column=5).value)
                c.ok(
                    f"{label} status distinguishes NOT CHECKED from OK",
                    "NOT CHECKED" in status and "REVIEW" in status,
                    status,
                )
        ur_row = find_row(checks, "Utility Reimbursements")
        c.ok(
            "the absent source subtotal leaves its cell empty",
            ur_row is not None and checks.cell(row=ur_row, column=3).value is None,
        )

        price_row = find_row(checks, "Total purchase price")
        c.ok(
            "total price is a formula off price-per-unit and unit count",
            str(checks.cell(row=price_row, column=2).value).startswith("=B"),
            str(checks.cell(row=price_row, column=2).value),
        )
        for label in ("T-12 cap rate", "T-3 annualized cap rate"):
            r = find_row_prefix(checks, label)
            c.ok(f"{label} row present", r is not None)
            if r:
                f = str(checks.cell(row=r, column=2).value)
                c.ok(f"{label} is IFERROR-guarded", f.startswith("=IFERROR("), f)
                c.ok(
                    f"{label} carries the illustrative marker",
                    "illustrative" in str(checks.cell(row=r, column=1).value),
                    str(checks.cell(row=r, column=1).value),
                )
        t6 = find_row(checks, "T-6 revenue (annualized)")
        c.ok("T-6 trend row present", t6 is not None)
        c.ok(
            "T-6 is a formula over the data tab, not a pasted value",
            "'T-12 Data'!" in str(checks.cell(row=t6, column=2).value),
            str(checks.cell(row=t6, column=2).value),
        )

        # ---- rent roll ---------------------------------------------------------------
        c.section("rent roll")
        roll = books["rent_roll"]["Rent Roll"]
        hdr = find_row(roll, "Unit ID")
        header_texts = [
            str(roll.cell(row=hdr, column=i).value or "")
            for i in range(1, len(write_rent_roll.ROLL_HEADERS) + 1)
        ]
        c.equal("roll headers exact", header_texts, write_rent_roll.ROLL_HEADERS)
        c.ok(
            "NO name-like column anywhere in the roll (I4)",
            not any(
                tok in h.lower()
                for h in header_texts
                for tok in ("tenant", "resident", "name")
            ),
            str(header_texts),
        )
        c.equal("roll has one row per unit", roll.max_row - hdr, len(deal["rent_roll"]["units"]))

        # the text-date unit must survive as text
        text_unit = next(
            u for u in deal["rent_roll"]["units"] if u.get("lease_start") == "n/a"
        )
        tr = find_row(roll, text_unit["unit_id"])
        c.ok("text-date unit found in the roll", tr is not None)
        c.equal(
            "a text date stays text — not coerced to a date",
            roll.cell(row=tr, column=7).value,
            "n/a",
        )
        c.ok(
            "and its lease type is blank rather than guessed",
            roll.cell(row=tr, column=9).value in ("", None),
            repr(roll.cell(row=tr, column=9).value),
        )
        real_date_row = hdr + 1
        c.ok(
            "a parseable date is written as a real date",
            isinstance(roll.cell(row=real_date_row, column=7).value, (dt.date, dt.datetime)),
            repr(roll.cell(row=real_date_row, column=7).value),
        )
        c.ok(
            "rent per SF is blank-guarded",
            str(roll.cell(row=real_date_row, column=11).value).startswith("=IFERROR("),
        )

        # ---- unit mix ----------------------------------------------------------------
        c.section("unit mix summary")
        mix = books["rent_roll"]["Unit Mix Summary"]
        for knob in ("Months lookback", "As-of date", "New-tenant tol (days)"):
            c.ok(
                f"knob '{knob}' is stated on the output",
                find_row(mix, knob) is not None
                or find_row(mix, knob, col=4) is not None
                or find_row(mix, knob, col=7) is not None,
            )
        plans = compute.plan_summary(deal)
        for plan in plans:
            r = find_row(mix, plan["floor_plan"], col=2)
            c.ok(f"plan row for {plan['floor_plan']}", r is not None)
        c.section("unit mix follows the Summary column contract")
        hdr_row = find_row(mix, "Bedroom")
        c.equal(
            "column order matches the contract",
            [mix.cell(row=hdr_row, column=i).value for i in range(1, 12)],
            # the internal rent-roll tool's Summary vocabulary (rentroll_rules
            # SUMMARY_HEADERS) — signed (F/G), new move-in (H/I), then the model feed.
            ["Bedroom", "Floor plan", "Avg SF", "Unit count", "In-place rent",
             "LEASES SIGNED rent", "Signed #", "NEW MOVE-IN rent", "Move-In #",
             "Market rent for model", "Basis"],
        )
        c.ok(
            "plans are sorted by bedroom then plan code",
            [p_["beds"] for p_ in plans] == sorted(p_["beds"] for p_ in plans),
        )
        fallback_plan = next(p for p in plans if p["movein_count"] == 0)
        fr = find_row(mix, fallback_plan["floor_plan"], col=2)
        c.ok(
            "a plan with no new leases states the in-place basis",
            "in-place" in str(mix.cell(row=fr, column=11).value).lower(),
            str(mix.cell(row=fr, column=11).value),
        )
        c.equal(
            "and its market rent equals the in-place average",
            round(float(mix.cell(row=fr, column=10).value), 2),
            round(fallback_plan["avg_in_place_rent"], 2),
        )
        newlease_plan = next(p for p in plans if p["movein_count"] > 0)
        nr = find_row(mix, newlease_plan["floor_plan"], col=2)
        c.ok(
            "a plan with new move-ins states the move-in basis and its sample size",
            "new move-ins" in str(mix.cell(row=nr, column=11).value).lower()
            and f"n={newlease_plan['movein_count']}" in str(mix.cell(row=nr, column=11).value),
            str(mix.cell(row=nr, column=11).value),
        )
        c.equal("the move-in rent column is the number the model reads on that plan",
                round(float(mix.cell(row=nr, column=8).value), 2),
                round(float(mix.cell(row=nr, column=10).value), 2))
        c.equal("and its count is the plan's move-in count",
                mix.cell(row=nr, column=9).value, newlease_plan["movein_count"])
        c.equal("the signed count is the plan's signed count",
                mix.cell(row=nr, column=7).value, newlease_plan["signed_count"])
        # Review of this change: the TOTAL row carries the property-level move-in and
        # signed averages (count-weighted), and propagates n/a rather than summing it to 0.
        tot_r = find_row(mix, "TOTAL", col=2)
        for col_, cnt_ in ((6, "G"), (8, "I")):
            f_ = str(mix.cell(row=tot_r, column=col_).value)
            c.ok(f"TOTAL col {col_} is the {cnt_}-weighted average, not blank",
                 f_.startswith("=IFERROR(SUMPRODUCT(") and f"${cnt_}$" in f_, f_)
        nmi = copy.deepcopy(deal)
        for u_ in nmi["rent_roll"]["units"]:
            u_["move_in"] = None
        with tempfile.TemporaryDirectory() as td_:
            out_ = write_rent_roll.build(nmi, Path(td_) / "nmi.xlsx")
            mx_ = load_workbook(out_)[write_rent_roll.SHEET_MIX]
            tr_ = find_row(mx_, "TOTAL", col=2)
            c.ok("Basis (K) is the wide column, not the move-in count (I)",
                 (mix.column_dimensions["K"].width or 0) >= 30
                 and (mix.column_dimensions["I"].width or 0) < 20,
                 str((mix.column_dimensions["I"].width, mix.column_dimensions["K"].width)))
            c.ok("the signed definition says CLASSIFIED lease starts",
                 "classified" in next(f_.definition for f_ in schema.FIELD_INVENTORY
                                      if f_.name == "plan_signed_rent").lower())
            c.ok("the market-rent definition names its lease-start prerequisite",
                 "lease start" in next(f_.definition for f_ in schema.FIELD_INVENTORY
                                       if f_.name == "market_rent_for_model").lower())
            c.equal("a roll with no move-in column reads n/a on the TOTAL row, never 0",
                    [mx_.cell(row=tr_, column=k).value for k in (6, 7, 8, 9)],
                    ["n/a", "n/a", "n/a", "n/a"])
        units_formula = str(mix.cell(row=nr, column=4).value)
        c.ok("plan unit count is a COUNTIF over the roll", units_formula.startswith("=COUNTIF("), units_formula)
        inplace_formula = str(mix.cell(row=nr, column=5).value)
        c.ok("in-place rent is an AVERAGEIFS on status", inplace_formula.startswith("=IFERROR(AVERAGEIFS("), inplace_formula)
        c.section("the mix tab carries its charts")
        c.equal("two native charts on the unit mix tab", len(mix._charts), 2)
        c.section("move-ins are counted, by move-in date")
        c.ok("move-ins block present", find_row(mix, "Move-ins in window") is not None)
        for label in ("Occupancy", "Vacancy", "Total units", "Occupied units", "Non-revenue units"):
            c.ok(f"occupancy block has '{label}'", find_row(mix, label) is not None)
        occ_row = find_row(mix, "Occupancy")
        c.ok(
            "occupancy is blank-guarded",
            str(mix.cell(row=occ_row, column=2).value).startswith("=IFERROR("),
        )

        # ---- trade-out: not a deliverable (2026-09-10) ---------------------------------
        c.section("no lease trade-out tab")
        c.ok(
            "the rent-roll workbook carries no trade-out sheet — an internal-pipeline input "
            "the demo does not ask for",
            "Lease Trade-Out" not in books["rent_roll"].sheetnames,
            books["rent_roll"].sheetnames,
        )
        absent = copy.deepcopy(deal)
        absent["trade_out"] = {
            "rows": [],
            "absent_reason": "The materials contained no lease trade-out report.",
        }
        absent_path = write_rent_roll.build(absent, tmpdir / "absent.xlsx")
        c.equal(
            "an absent trade-out section changes nothing about the sheets written",
            load_workbook(absent_path).sheetnames,
            books["rent_roll"].sheetnames,
        )

        # ---- deal summary ------------------------------------------------------------
        c.section("deal summary card")
        summary = books["deal_summary"]["Deal Summary"]
        for label in (
            "Deal name", "Street address", "County", "Unit count", "Product type",
            "Listing brokerage", "Pricing guidance", "Market rent for model",
            "Tax rate as billed", "Elementary school", "Notes",
        ):
            c.ok(f"summary has '{label}'", find_row(summary, label) is not None)
        tax_row = find_row(summary, "Computed annual tax")
        tax_formula = str(summary.cell(row=tax_row, column=2).value)
        c.ok("computed tax is a formula", tax_formula.startswith("=IFERROR("), tax_formula)
        c.ok(
            "and it degrades to MISSING, not to a plausible number",
            schema.MISSING in tax_formula,
            tax_formula,
        )
        cap_row = find_row_prefix(summary, "T-12 cap rate")
        c.ok(
            "cap rate is marked illustrative",
            "illustrative" in str(summary.cell(row=cap_row, column=1).value),
        )
        ppu_row = find_row_prefix(summary, "Price per unit")
        c.ok(
            "price per unit is marked illustrative",
            "illustrative" in str(summary.cell(row=ppu_row, column=1).value),
        )

        c.section("R#3: the workbook carries the same tax evidence the preview does")
        for label in (
            "Parcel / APN", "Tax area code", "Billed by the county",
            "Check: computed vs billed", "Transfer tax", "Mortgage / recording tax",
        ):
            c.ok(
                f"summary workbook has '{label}'",
                find_row_prefix(summary, label) is not None,
                "present in the HTML preview but not in the downloadable workbook",
            )
        chk = find_row_prefix(summary, "Check: computed vs billed")
        if chk:
            formula = str(summary.cell(row=chk, column=2).value)
            c.ok("the cross-check is a formula, not a pasted verdict",
                 formula.startswith("=IF"), formula[:80])
            c.ok(
                "RR#1: it guards on ISNUMBER so a MISSING input cannot make it #VALUE!",
                "ISNUMBER(" in formula,
                formula[:160],
            )

        c.section("RR#1b: an incomplete tax block still ships a readable check")
        incomplete = copy.deepcopy(deal)
        incomplete["summary"]["tax_rate_as_billed"] = schema.MISSING
        inc_path = write_deal_summary.build(incomplete, tmpdir / "missing-tax.xlsx")
        inc = load_workbook(inc_path)["Deal Summary"]
        inc_chk = find_row_prefix(inc, "Check: computed vs billed")
        c.ok("the check row is still written", inc_chk is not None)
        if inc_chk:
            f2 = str(inc.cell(row=inc_chk, column=2).value)
            c.ok("and still ISNUMBER-guarded", "ISNUMBER(" in f2, f2[:160])
            c.equal(
                "compute's mirror says NOT CHECKED for the same deal (I6)",
                compute.tax_check(incomplete)["status"],
                "NOT CHECKED",
            )

        c.section("R#2: the TOTAL in-place rent is the occupied-unit average")
        total_row = find_row(mix, "TOTAL", col=2)
        c.ok("TOTAL row found", total_row is not None)
        expected = compute.roll_totals(deal)["wavg_in_place_rent"]
        inplace_formula = str(mix.cell(row=total_row, column=5).value)
        c.ok(
            "it is computed off the roll, not by weighting plan averages by unit count",
            "AVERAGEIFS(" in inplace_formula,
            inplace_formula[:110],
        )
        c.ok(
            "and it does not SUMPRODUCT the plan column by unit count",
            "SUMPRODUCT" not in inplace_formula,
            inplace_formula[:110],
        )

        c.section("R#10: every provenance locator resolves to the cell it claims")
        import provenance as _pv
        import write_rent_roll as _wrr

        sheet_of = {
            "Unit Mix Summary": books["rent_roll"]["Unit Mix Summary"],
            "Rent Roll": books["rent_roll"]["Rent Roll"],
        }
        metric_col = {
            "summary.in_place_rent_avg": "In-place rent",
            "summary.market_rent_for_model": "Market rent for model",
        }
        bad = []
        for path, st in (deal.get("provenance") or {}).items():
            loc = str(st.get("locator", ""))
            if "!" not in loc:
                continue
            sheet, ref = loc.split("!", 1)
            ws_ = sheet_of.get(sheet)
            if ws_ is None or not ref.isalpha():
                continue
            hdr = find_row(ws_, "Bedroom") or find_row(ws_, "Unit ID")
            if hdr is None:
                continue
            col_idx = ord(ref.upper()) - 64
            header = ws_.cell(row=hdr, column=col_idx).value
            want = metric_col.get(path)
            if want and str(header) != want:
                bad.append(f"{path}: locator {loc} -> header {header!r}, expected {want!r}")
            elif header is None:
                bad.append(f"{path}: locator {loc} points at a column with no header")
        c.equal("no stamp points at the wrong or a nonexistent column", bad, [])
        c.equal(
            "the mix layout map is the single source for those letters",
            _wrr.MIX_COLUMNS["in_place_rent"],
            "E",
        )
        c.equal("and market rent", _wrr.MIX_COLUMNS["market_rent_for_model"], "J")

        c.section("R#3b: the audit trail carries values, not just paths")
        trail = books["deal_summary"]["Audit Trail"]
        thdr = find_row(trail, "Value")
        c.ok("trail header found", thdr is not None)
        headers = [trail.cell(row=thdr, column=i).value for i in range(1, 8)]
        c.ok(
            "the trail shows the value itself",
            "Value as published" in headers or "Value" in headers,
            str(headers),
        )
        c.ok(
            "a stamped scalar's value appears on its trail row",
            any(
                str(trail.cell(row=r, column=1).value) == "summary.units_total"
                and str(trail.cell(row=r, column=2).value) == str(deal["summary"]["units_total"])
                for r in range(thdr + 1, trail.max_row + 1)
            ),
            "the trail lists paths and sources but not the value each one describes",
        )

        c.section("comps tab")
        comps_ws = books["deal_summary"]["Rent Comps"]
        chdr = find_row(comps_ws, "Community")
        c.ok("comps header found", chdr is not None)
        comp_headers = [comps_ws.cell(row=chdr, column=i).value for i in range(1, 15)]
        for needed in ("Year built", "Asking rent", "Concession", "Net effective rent",
                       "Source"):
            c.ok(f"comps column '{needed}'", needed in comp_headers, str(comp_headers))
        src_col = comp_headers.index("Source") + 1
        tier1_rows = [r for r in deal["comps"]["rows"] if r["tier"] == 1]
        sources = {
            comps_ws.cell(row=r, column=src_col).value
            for r in range(chdr + 1, chdr + 1 + len(tier1_rows))
        }
        c.equal("every comp row says own weekly scrape", sources, {schema.COMP_SOURCE})
        c.ok(
            "subject-versus-comp-set block present",
            find_row_prefix(comps_ws, "Subject vs comp set") is not None,
        )

        c.section("two-tier comp display")
        # Tier headers, in order, so the workbook and the preview cannot diverge on which
        # communities are being called the market.
        t1 = find_row_prefix(comps_ws, "Tier 1 — Top comps")
        t2 = find_row_prefix(comps_ws, "Tier 2 — Additional comps")
        t3 = find_row_prefix(comps_ws, "Tier 3 — Also nearby")
        c.ok("tier 1 section present", t1 is not None)
        c.ok("tier 2 section present", t2 is not None)
        c.ok("tier 3 section is GONE", t3 is None, f"still at row {t3}")
        c.ok("tiers appear in order", None not in (t1, t2) and t1 < t2, f"{t1}, {t2}")
        c.ok("the nearby data is still on the deal the workbook was built from",
             bool(deal["comps"]["nearby"]), "the writer must not be dropping the data")
        sheet_text = " ".join(
            str(comps_ws.cell(row=r, column=col).value or "")
            for r in range(1, comps_ws.max_row + 1)
            for col in range(1, 15)
        )
        c.ok("and no nearby reason is written to the sheet",
             not any(str(n.get("reason") or "") in sheet_text for n in deal["comps"]["nearby"]))
        c.ok(
            "the vintage footnote sits under the tier-1 header",
            str(comps_ws.cell(row=t1 + 1, column=1).value or "").startswith("Vintage filter"),
            repr(comps_ws.cell(row=t1 + 1, column=1).value),
        )
        c.ok(
            "the footnote discloses that some rows carry no year",
            "where year known" in str(comps_ws.cell(row=t1 + 1, column=1).value or "").lower(),
            repr(comps_ws.cell(row=t1 + 1, column=1).value),
        )

        c.ok("no 'Not a comp because' column survives anywhere on the sheet",
             "Not a comp because" not in sheet_text)

        # The averages must span BOTH detail tiers. One range would quietly drop tier 2.
        avg_row = find_row(comps_ws, "Comp-set average asking")
        c.ok("comp-set average present", avg_row is not None)
        if avg_row:
            formula = str(comps_ws.cell(row=avg_row, column=2).value)
            c.ok(
                "the average spans both detail tiers, as separate ranges",
                formula.count(":") == 2 and "," in formula,
                formula,
            )
            c.ok(
                "and the note says which population it is over",
                "detailed tiers" in str(comps_ws.cell(row=avg_row, column=3).value or ""),
                repr(comps_ws.cell(row=avg_row, column=3).value),
            )

        c.section("D#2: the comps tab shares compute's withholding decision (I6)")
        comps_ws = books["deal_summary"]["Rent Comps"]
        cm = compute.comps(deal)
        for label in ("Comp set, mix-matched", "Subject rent, same bed counts", "Mix coverage"):
            c.ok(f"workbook shows '{label}'", find_row(comps_ws, label) is not None)
        sp_row = find_row_prefix(comps_ws, "Subject vs comp set")
        c.ok("workbook has the spread row", sp_row is not None)
        if sp_row:
            cell = comps_ws.cell(row=sp_row, column=2).value
            if cm["subject_vs_compset_spread"] is None:
                c.ok("published nothing when compute withheld", str(cell) == "WITHHELD",
                     repr(cell))
            else:
                c.ok(
                    "the published spread is a formula over the MATCHED cells",
                    str(cell).startswith("=IFERROR(") and "-1" in str(cell),
                    str(cell)[:90],
                )
                c.ok(
                    "and not over the flat average",
                    "AVERAGE(" not in str(cell),
                    str(cell)[:90],
                )

        c.section("D#2b: a withheld spread reaches the workbook as WITHHELD, with the reason")
        thin_deal = copy.deepcopy(deal)
        thin_deal["comps"]["rows"] = [
            r for r in thin_deal["comps"]["rows"] if r["beds"] != 3
        ]
        tc = compute.comps(thin_deal)
        c.ok("compute withholds on this deal", tc["subject_vs_compset_spread"] is None,
             str(tc["subject_vs_compset_spread"]))
        thin_path = write_deal_summary.build(thin_deal, tmpdir / "thin-comps.xlsx")
        tw = load_workbook(thin_path)["Rent Comps"]
        t_row = find_row_prefix(tw, "Subject vs comp set")
        c.ok("the workbook row exists", t_row is not None)
        if t_row:
            c.equal("and reads WITHHELD", str(tw.cell(row=t_row, column=2).value), "WITHHELD")
            note = str(tw.cell(row=t_row, column=3).value or "")
            # Compare against compute's own string rather than guessing a keyword: the
            # point of I6 is that one decision produces one wording on both surfaces.
            c.equal("carrying compute's own reason verbatim", note, tc["spread_withheld_reason"])

        c.section("demand tab — the COMPLETE report, not a summary of it")
        dem = books["deal_summary"]["Demand"]
        cols = deal["demand"]["columns"]
        c.ok("demand has seven columns", len(cols) == 7, str(cols))
        # Every section the interchange declares must appear, including the two awkward
        # shapes: the ZIP home-value block (one value, not seven) and the labor section's
        # year-by-year metro table.
        for sec in deal["demand"]["sections"]:
            c.ok(f"demand section '{sec['title'][:34]}'",
                 find_row(dem, sec["title"]) is not None)
        c.ok(
            "the ring column headers are written",
            find_row(dem, cols[0], col=2) is not None,
            f"expected {cols[0]!r} as a column header",
        )
        c.ok(
            "the metro history table is rendered",
            find_row(dem, "Year") is not None
            and find_row(dem, "Nonfarm jobs (000s)", col=2) is not None,
        )
        # 🔴 The 1-mile caveat must be ON the tab. It is a known overcount in the engine,
        # and a published artifact cannot leave it to the reader to know.
        caveat_row = None
        for r in range(1, 12):
            v = str(dem.cell(row=r, column=1).value or "")
            if "1-mile" in v and "overstate" in v:
                caveat_row = r
        c.ok("the 1-mile overcount caveat is on the tab", caveat_row is not None)
        c.ok("and it sits above the data", (caveat_row or 99) < 10, str(caveat_row))
        c.ok(
            "a ZIP-geography row carries one value, not seven",
            any(
                len(r["values"]) == 1
                for s in deal["demand"]["sections"] for r in s["rows"]
                if r.get("kind") == "zip"
            ),
        )

        c.section("supply tab — unchanged figures, with a coverage note")
        sup = books["deal_summary"]["Supply"]
        c.ok("supply ring totals block", find_row(sup, "Ring totals") is not None)
        c.ok(
            "supply stages are named",
            find_row(sup, "Entitlement", col=2) is not None
            and find_row(sup, "Under construction", col=2) is not None,
        )
        note_found = any(
            "ongoing work" in str(sup.cell(row=r, column=1).value or "")
            for r in range(1, 10)
        )
        c.ok("the coverage note says coverage is ongoing work", note_found)
        c.ok(
            "and that a zero is a floor, not a census",
            any("floor" in str(sup.cell(row=r, column=1).value or "").lower()
                for r in range(1, 10)),
        )

        c.section("property tax tab")
        tax = books["deal_summary"]["Property Tax"]
        import compute as _cmp
        tu = _cmp.tax_underwriting(deal)
        for label in ("Assessed-value basis", "Computed annual tax", "County billed amount",
                      "Does the basis reset on sale?", "Year-one full value"):
            c.ok(f"tax tab shows '{label}'", find_row(tax, label) is not None)
        # The illustrative rows carry the chip in their label, so match by prefix.
        for label in ("Year-one tax", "Illustrative purchase price"):
            c.ok(f"tax tab shows '{label}'", find_row_prefix(tax, label) is not None)
        gen_row = find_row(tax, "Assessed value if it reset to the price")
        c.ok("the generic shortcut is shown too", gen_row is not None)
        c.ok(
            "the computed-vs-billed cell is ISNUMBER-guarded",
            "NOT CHECKED" in str(tax.cell(row=find_row(tax, "Computed less billed"),
                                          column=2).value),
        )
        reset_row = find_row(tax, "Does the basis reset on sale?")
        if reset_row:
            c.equal(
                "a no-reset state says No",
                tax.cell(row=reset_row, column=2).value,
                {True: "Yes", False: "No", None: "not vendored"}[tu["resets_on_sale"]],
            )

        c.section("R17 F4: a MISSING exemption reaches the formula as a NUMBER")
        import copy as _copy
        miss = _copy.deepcopy(deal)
        miss["summary"]["tax_exemption"] = schema.MISSING
        miss_path = write_deal_summary.build(miss, tmpdir / "exempt-missing.xlsx")
        mws = load_workbook(miss_path)["Property Tax"]
        ex_row = find_row(mws, "Exemption")
        c.ok("the exemption row exists", ex_row is not None)
        if ex_row:
            val = mws.cell(row=ex_row, column=2).value
            c.ok(
                "it holds a number, not the MISSING text",
                isinstance(val, (int, float)),
                f"a text cell here makes Excel return #VALUE! while Python reports a "
                f"number — got {val!r}",
            )
            c.equal("and it is the normalized zero compute used", val, 0)

        c.section("R17 F5: the tax tab labels full value and assessed value separately")
        tw = books["deal_summary"]["Property Tax"]
        c.ok("year-one FULL value row", find_row(tw, "Year-one full value") is not None)
        c.ok("year-one ASSESSED value row",
             find_row(tw, "Year-one assessed value") is not None)
        fv_row = find_row(tw, "Year-one full value")
        if fv_row:
            c.ok(
                "the full-value note does not describe the assessed value",
                "x assessment ratio" not in str(tw.cell(row=fv_row, column=3).value or ""),
                repr(tw.cell(row=fv_row, column=3).value),
            )

        c.section("notes and disclosures tab")
        nt = books["deal_summary"]["Notes & Disclosures"]
        c.ok("notes header", find_row(nt, "Topic") is not None)
        notes = deal["meta"]["notes"]
        c.ok("the deal carries notes", len(notes) >= 5, str(len(notes)))
        for n in notes[:3]:
            c.ok(f"note '{n['topic'][:34]}' is written",
                 find_row(nt, n["topic"]) is not None)

        c.section("provenance renders BESIDE the values, not only in the trail tab")
        summ = books["deal_summary"]["Deal Summary"]
        addr_row = find_row(summ, schema.summary_label("street_address"))
        c.ok("a summary value row was found", addr_row is not None)
        if addr_row:
            c.equal(
                "its source sits in column D",
                summ.cell(row=addr_row, column=4).value,
                _pv.compact_for(deal, "summary.street_address"),
            )
            c.ok(
                "and that source names a document",
                bool(str(summ.cell(row=addr_row, column=4).value or "").strip()),
            )
        c.ok(
            "the dedicated Audit Trail tab still carries the full registry",
            "Audit Trail" in books["deal_summary"].sheetnames,
        )

        # ---- manifest -----------------------------------------------------------
        c.section("model-ready manifest")
        man = books["manifest"][write_manifest.SHEET]
        hdr_row = find_row(man, "Group")
        c.ok("manifest header row found", hdr_row is not None)
        c.equal(
            "manifest columns exact",
            [man.cell(row=hdr_row, column=i).value for i in range(1, len(schema.MANIFEST_COLUMNS) + 1)],
            list(schema.MANIFEST_COLUMNS),
        )
        pairs: list[tuple[str, str]] = []
        for r in range(hdr_row + 1, man.max_row + 1):
            group = man.cell(row=r, column=1).value
            name = man.cell(row=r, column=2).value
            if group and name:
                pairs.append((str(group), str(name)))
        expected_pairs = [(f.group, f.name) for f in schema.MANIFEST_FIELDS]
        print(f"  ...  MANIFEST ROWS = {len(pairs)}  (FIELD_INVENTORY = {len(expected_pairs)})")
        c.equal(
            "row count equals len(MANIFEST_FIELDS)", len(pairs), len(expected_pairs)
        )
        c.equal("every field appears exactly once, in order", pairs, expected_pairs)
        c.ok("cover states the pitch", _sheet_contains(man, "hands to"), "")
        c.ok(
            "cover states the licensed-data boundary",
            _sheet_contains(man, "excluded by design"),
        )
        c.ok(
            "cover states the field count",
            find_row(man, "Fields listed") is not None,
        )

        c.section("manifest maps-to text (spot checks)")
        by_name = {f.name: f for f in schema.MANIFEST_FIELDS}
        spot = {
            "plan_market_rent_for_model": "Unit-mix market rent",
            "vacancy": "Vacancy assumption",
            "tax_rate_as_billed": "Property tax rate",
        }
        for name, expected in spot.items():
            c.equal(f"maps-to for {name}", by_name[name].fills, expected)

        c.section("manifest without a deal")
        no_deal = write_manifest.build(tmpdir / "manifest-nodeal.xlsx", None)
        nw = load_workbook(no_deal)[write_manifest.SHEET]
        nhdr = find_row(nw, "Group")
        examples = {
            nw.cell(row=r, column=8).value   # H: Example (demo deal); G is Public demo
            for r in range(nhdr + 1, nw.max_row + 1)
            if nw.cell(row=r, column=2).value
        }
        c.equal("every example reads an em dash", examples, {"—"})

        # ---- cross-cutting -----------------------------------------------------------
        c.section("cross-cutting: gate, properties, validation")
        terms = scrub_check.load_terms([])
        findings = scrub_check.run(list(paths.values()), terms)
        c.ok(
            "no denylisted term in any of the four workbooks",
            not findings,
            "\n".join(f.render() for f in findings),
        )
        for key, wb in books.items():
            c.ok(
                f"{key} has no personal authorship metadata",
                wb.properties.creator == "AI Underwriting Pipeline",
                repr(wb.properties.creator),
            )

        c.section("N1: the workbook decides GSR presence the same way compute does (I6)")
        gsr_row = find_row(pnl, "Gross Scheduled Rent")
        gsr_b = str(pnl.cell(row=gsr_row, column=2).value)
        gsr_c = str(pnl.cell(row=gsr_row, column=3).value)
        c.ok(
            "B column gates on COUNTIF (existence), not SUMIF<>0 (aggregate)",
            "COUNTIF(" in gsr_b and "<>0" not in gsr_b,
            gsr_b[:120],
        )
        c.ok(
            "C column gates on the same COUNTIF",
            "COUNTIF(" in gsr_c and "<>0" not in gsr_c,
            gsr_c[:120],
        )
        zero_net = copy.deepcopy(deal)
        zero_net["t12"]["rows"].append(
            {
                "label": "Gross Scheduled Rent",
                "category": "Gross Scheduled Rent",
                "monthly": [-1000.0] * 9 + [3000.0] * 3,
            }
        )
        zn_path = write_t12_analysis.build(zero_net, tmpdir / "zeronet.xlsx")
        zn_pnl = load_workbook(zn_path)["P&L Summary"]
        zn_row = find_row(zn_pnl, "Gross Scheduled Rent")
        c.ok(
            "a zero-net GSR row still produces an existence-gated formula",
            "COUNTIF(" in str(zn_pnl.cell(row=zn_row, column=2).value),
            str(zn_pnl.cell(row=zn_row, column=2).value)[:120],
        )

        c.section("N2: an out-of-range generation date never half-writes a file")
        for bad in ("1900-01-01", "2200-06-30"):
            bad_deal = copy.deepcopy(deal)
            bad_deal["meta"]["generated_on"] = bad
            target = tmpdir / f"range-{bad}.xlsx"
            c.raises(
                f"generated_on={bad} is refused by the writer",
                lambda bd=bad_deal, t=target: write_deal_summary.build(bd, t),
                ValueError,
            )
            c.ok(
                f"and leaves no file behind for {bad}",
                not target.exists(),
                f"{target} exists — saved, then failed while normalizing",
            )

        c.section("R1: the monthly revenue helper gates on existence too")
        # Same defect class as N1, one call site further along: the helper row on the data
        # tab decided GSR presence PER MONTH from a nonzero SUMIF, so a printed row with a
        # zero month substituted GPR+GLL for that month alone -- and the workbook's T-6
        # check then disagreed with the preview, which treats the printed row as
        # authoritative for every month.
        zero_month = copy.deepcopy(deal)
        gsr_monthly = [50000.0] * 12
        gsr_monthly[8] = 0.0  # 2026-03, inside the trailing-six window
        zero_month["t12"]["rows"].append(
            {
                "label": "Gross Scheduled Rent",
                "category": "Gross Scheduled Rent",
                "monthly": gsr_monthly,
            }
        )
        c.equal("the zero-month fixture is schema-valid", schema.validate(zero_month), [])
        t6_idx = compute.recent_indices(zero_month["t12"]["month_labels"], 6)
        c.ok("and month 8 really is inside the T-6 window", 8 in t6_idx, str(t6_idx))

        zm_path = write_t12_analysis.build(zero_month, tmpdir / "zeromonth.xlsx")
        zm_data = load_workbook(zm_path)["T-12 Data"]
        helper_row = find_row(zm_data, write_t12_analysis.REVENUE_HELPER_LABEL)
        c.ok("helper row present", helper_row is not None)
        zero_col = write_t12_analysis.MONTH_FIRST_COL + 8
        zero_formula = str(zm_data.cell(row=helper_row, column=zero_col).value)
        c.ok(
            "the zero month's helper formula gates on COUNTIF existence",
            "COUNTIF(" in zero_formula,
            zero_formula[:140],
        )
        c.ok(
            "and carries no aggregate <>0 presence gate",
            "<>0" not in zero_formula,
            zero_formula[:140],
        )
        for offset in range(12):
            f = str(zm_data.cell(row=helper_row, column=write_t12_analysis.MONTH_FIRST_COL + offset).value)
            c.ok(
                f"month {offset} helper formula is existence-gated",
                "COUNTIF(" in f and "<>0" not in f,
                f[:100],
            )

        c.section("R1b: no aggregate presence gate survives anywhere in the writers")
        # The durable form of the grep: a presence decision must never be spelled as
        # "the total is nonzero" in any emitted formula.
        sweep_targets = sorted(HERE.glob("write_*.py")) + [
            HERE / "render_previews.py",
            HERE / "compute.py",
        ]
        c.equal(
            "no aggregate presence gate in any writer, renderer, or compute",
            aggregate_gate_offenders(sweep_targets),
            [],
        )

        c.section("R3: the sweep itself can see the known offender")
        # A sweep that cannot fail is not a sweep. These two samples are the REAL pre-fix
        # source, verbatim from git (289e26f^ and e0b1a6f^). Both build the gate through
        # the local `sumif` helper, so the SUMIF( token never sits adjacent to the
        # comparison — the operator and the 0 land in a string segment of their own. An
        # uppercase-literal regex sees nothing here; its only "hit" was docstring prose.
        historical = HERE / "fixtures" / "historical_gate_sample.py.txt"
        c.ok("historical sample present", historical.exists(), str(historical))
        caught = aggregate_gate_offenders([historical])
        c.ok(
            "the sweep flags the pre-R1 monthly-helper shape",
            any(o.split(":")[1] == "18" for o in caught if len(o.split(":")) > 1),
            f"caught={caught}",
        )
        c.ok(
            "the sweep flags both pre-N1 rollup shapes",
            len([o for o in caught if o.split(":")[1] in ("27", "28")]) == 2,
            f"caught={caught}",
        )
        c.equal("three gates in the sample, three offenders", len(caught), 3)
        c.ok("and each offender names its line", all(":" in o for o in caught), str(caught))

        c.section("R3b: prose alone is not an offender")
        # The old detector's only hit was a docstring, so it reported the right answer for
        # the wrong reason and went silent as soon as the prose was reworded.
        prose_only = HERE / "fixtures" / "prose_only_gate_sample.py.txt"
        c.ok("prose-only sample present", prose_only.exists(), str(prose_only))
        c.equal(
            "a file whose only match is prose reports no offender",
            aggregate_gate_offenders([prose_only]),
            [],
        )

        # And the same sweep over what actually ships: every formula in every generated
        # workbook. Source prose cannot fool this one.
        emitted = []
        for key, wb in books.items():
            for ws_ in wb.worksheets:
                for row_cells in ws_.iter_rows():
                    for cell in row_cells:
                        v = cell.value
                        if isinstance(v, str) and v.startswith("=") and "<>0" in v:
                            emitted.append(f"{key}/{ws_.title}!{cell.coordinate}")
        c.equal("no emitted formula gates presence on a nonzero total", emitted, [])
        zm_emitted = []
        for ws_ in load_workbook(zm_path).worksheets:
            for row_cells in ws_.iter_rows():
                for cell in row_cells:
                    v = cell.value
                    if isinstance(v, str) and v.startswith("=") and "<>0" in v:
                        zm_emitted.append(f"{ws_.title}!{cell.coordinate}")
        c.equal("nor in the printed-GSR-with-a-zero-month workbook", zm_emitted, [])

        c.section("R2: the archive backstop survives an out-of-range leap day")
        # validate() rejects such a deal, so this exercises the backstop directly: a
        # caller bypassing validation must still get a valid file, never a crash after
        # the workbook has already been written.
        from openpyxl import Workbook as _WB

        for bad_date in ("2400-02-29", "1900-02-29", "0001-01-01", "9999-12-31"):
            probe = tmpdir / f"leap-{bad_date}.xlsx"
            wb_probe = _WB()
            wb_probe.active["A1"] = "x"
            try:
                save_workbook(wb_probe, probe, bad_date)
                ok = probe.exists() and probe.stat().st_size > 0
                c.ok(f"backstop handles generated_on={bad_date}", ok, "no file produced")
                if ok:
                    load_workbook(probe)  # must be a readable archive
            except Exception as exc:  # noqa: BLE001
                c.ok(
                    f"backstop handles generated_on={bad_date}",
                    False,
                    f"{type(exc).__name__}: {exc} (file left: {probe.exists()})",
                )

        c.section("F3: a formula cell keeps its numeric format")
        # A formula is a Python string, so a naive isinstance(str) check formats every
        # formula cell as text -- and Excel then shows a recalculated cap rate as 0.0549
        # instead of 5.49%.
        checks_ws = books["t12_analysis"]["Checks & Pricing"]
        summary_ws = books["deal_summary"]["Deal Summary"]
        roll_ws = books["rent_roll"]["Rent Roll"]
        fmt_cases = [
            ("deal summary total price", summary_ws, find_row_prefix(summary_ws, "Total price"), 2, "#,##0"),
            ("deal summary T-12 cap rate", summary_ws, find_row_prefix(summary_ws, "T-12 cap rate"), 2, "0.00%"),
            ("deal summary computed tax", summary_ws, find_row(summary_ws, "Computed annual tax"), 2, "#,##0"),
            ("checks T-12 cap rate", checks_ws, find_row_prefix(checks_ws, "T-12 cap rate"), 2, "0.00%"),
            ("checks total purchase price", checks_ws, find_row(checks_ws, "Total purchase price"), 2, "#,##0"),
        ]
        for label, ws_, r, col_idx, expected_fmt in fmt_cases:
            if r is None:
                c.ok(f"{label}: row found", False, "label not located")
                continue
            cell = ws_.cell(row=r, column=col_idx)
            c.ok(
                f"{label} is stored as a formula",
                cell.data_type == "f",
                f"data_type={cell.data_type!r} value={cell.value!r}",
            )
            c.equal(f"{label} number_format", cell.number_format, expected_fmt)
        rent_psf = roll_ws.cell(row=find_row(roll_ws, "Unit ID") + 1, column=11)
        c.ok("rent per SF is a formula", rent_psf.data_type == "f", repr(rent_psf.value))
        c.equal("rent per SF number_format", rent_psf.number_format, FMT_CURRENCY_CENTS)

        c.section("F3b: every formula cell in every workbook is typed as a formula")
        mistyped = []
        for key, wb in books.items():
            for ws_ in wb.worksheets:
                for row_cells in ws_.iter_rows():
                    for cell in row_cells:
                        if isinstance(cell.value, str) and cell.value.startswith("="):
                            if cell.data_type != "f":
                                mistyped.append(f"{key}/{ws_.title}!{cell.coordinate}")
        c.equal("no formula is stored as text", mistyped, [])

        c.section("F4: source text is never coerced into a formula")
        injected = copy.deepcopy(deal)
        injected["rent_roll"]["units"][0]["unit_id"] = '=HYPERLINK("http://evil.example","x")'
        c.raises(
            "a deal whose source text starts with '=' is refused",
            lambda: write_rent_roll.build(injected, tmpdir / "inj.xlsx"),
            ValueError,
        )
        text_cell = roll_ws.cell(row=find_row(roll_ws, "Unit ID") + 1, column=1)
        c.equal("an ordinary text cell is typed as a string", text_cell.data_type, "s")

        c.section("builds are byte-reproducible")
        # A workbook rebuilt from the same JSON must be the same bytes. Otherwise every
        # regeneration shows up as a modified file and a real change becomes
        # indistinguishable from a rebuild — and the site's "downloaded file matches the local
        # copy" check loses its meaning.
        # Advance the wall clock deliberately. Without this the two builds can land in
        # the same second and the check passes by coincidence -- which is exactly how an
        # incomplete determinism fix (docProps pinned, zip entry timestamps not) went
        # green the first time.
        time.sleep(1.2)
        second = Path(tmp) / "again"
        rebuilds = {
            "deal_summary": write_deal_summary.build(
                deal, second / schema.OUTPUT_FILES["deal_summary"]
            ),
            "t12_analysis": write_t12_analysis.build(
                deal, second / schema.OUTPUT_FILES["t12_analysis"]
            ),
            "rent_roll": write_rent_roll.build(
                deal, second / schema.OUTPUT_FILES["rent_roll"]
            ),
            "manifest": write_manifest.build(
                second / schema.OUTPUT_FILES["manifest"], deal
            ),
        }
        for key, again in rebuilds.items():
            first_digest = hashlib.md5(paths[key].read_bytes()).hexdigest()
            second_digest = hashlib.md5(again.read_bytes()).hexdigest()
            c.equal(f"{key} rebuild is byte-identical", second_digest, first_digest)

        c.section("writers refuse an invalid deal")
        bad = copy.deepcopy(deal)
        bad["comps"]["source"] = "Licensed comp feed"
        for name, fn in (
            ("deal summary", lambda: write_deal_summary.build(bad, tmpdir / "bad1.xlsx")),
            ("t12", lambda: write_t12_analysis.build(bad, tmpdir / "bad2.xlsx")),
            ("rent roll", lambda: write_rent_roll.build(bad, tmpdir / "bad3.xlsx")),
            ("manifest", lambda: write_manifest.build(tmpdir / "bad4.xlsx", bad)),
        ):
            c.raises(f"{name} refuses an invalid deal", fn, ValueError)

        # ---- newest-first months (direction detection) --------------------------------
        c.section("newest-first statement writes a different T-3 range")
        reversed_deal = copy.deepcopy(deal)
        reversed_deal["t12"]["month_labels"] = list(
            reversed(reversed_deal["t12"]["month_labels"])
        )
        for row in reversed_deal["t12"]["rows"]:
            row["monthly"] = list(reversed(row["monthly"]))
        rp = write_t12_analysis.build(reversed_deal, tmpdir / "reversed.xlsx")
        rdata = load_workbook(rp)["T-12 Data"]
        rhdr = find_row(rdata, "Line item")
        c.equal(
            "newest-first T-3 sums the FIRST three month columns",
            rdata.cell(row=rhdr + 1, column=16).value,
            f"=SUM(C{rhdr + 1}:E{rhdr + 1})",
        )
        c.equal(
            "oldest-first T-3 sums the LAST three month columns",
            data.cell(row=first_data, column=16).value,
            f"=SUM(L{first_data}:N{first_data})",
        )

    return c.finish()


def _sheet_contains(ws, needle: str) -> bool:
    needle = needle.lower()
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and needle in str(cell.value).lower():
                return True
    return False


if __name__ == "__main__":
    raise SystemExit(guard(main))
