"""Shared openpyxl helpers for the writers.

Two jobs. First, one consistent sheet chrome so a reader always finds the same offsets.
Second, and more important: a :class:`Block` that records the row every label was written
on, so formulas are emitted against recorded rows rather than hardcoded ones. The tests
read files back by searching for the label, which is what makes a layout change safe.

Pure openpyxl. No Excel, no COM, no macros -- everything here runs headless.
"""

from __future__ import annotations

import datetime as _dt
import re
import zipfile
from dataclasses import dataclass, field as dc_field
from pathlib import Path
from typing import Any

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from schema import ZIP_MAX_YEAR, ZIP_MIN_YEAR
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# ---- neutral formatting vocabulary ---------------------------------------------------

BODY_FONT = Font(name="Calibri", size=11)
BOLD = Font(name="Calibri", size=11, bold=True)
TITLE = Font(name="Calibri", size=14, bold=True)
SUBTLE = Font(name="Calibri", size=10, italic=True, color="595959")
SECTION_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
SUBTOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")
SECTION_FILL = PatternFill("solid", fgColor="595959")

THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
UNDERLINE = Border(bottom=Side(style="medium", color="808080"))

FMT_CURRENCY = "#,##0"
FMT_CURRENCY_CENTS = "#,##0.00"
FMT_RENT = "#,##0"
FMT_PCT = "0.0%"
FMT_PCT2 = "0.00%"
FMT_COUNT = "#,##0"
FMT_SF = "#,##0"
FMT_DATE = "mm/dd/yyyy"
FMT_NUM1 = "#,##0.0"
FMT_TEXT = "@"

WRAP = Alignment(wrap_text=True, vertical="top")
RIGHT = Alignment(horizontal="right")
LEFT = Alignment(horizontal="left")


def assign(cell, value: Any, fmt: str | None = None) -> None:
    """Put ``value`` in ``cell`` with the right type and number format.

    Three cases, and the distinction between the first two is the point:

    * **a formula** — a string starting with ``=``. openpyxl types it as a formula, and
      the requested number format is applied, because a formula IS a Python string and a
      naive ``isinstance(value, str)`` check would format every cap-rate and per-unit cell
      as text. Excel would then recalculate 0.0549 and display it as ``0.0549`` rather
      than ``5.49%``.
    * **source text** — any other string. Written with an explicit string data type so
      openpyxl cannot guess it into a formula. Reading a leading ``=`` as "formula" is
      only safe because :func:`schema.validate` rejects any source string that starts with
      one; these two halves are a single invariant and must move together.
    * **anything else** — number, date, ``None`` — assigned with the requested format.
    """
    if isinstance(value, str):
        if value.startswith("="):
            cell.value = value
            if fmt:
                cell.number_format = fmt
            return
        if value == "":
            cell.value = None
            return
        # Assign, then force the string type. openpyxl guesses "formula" from a leading
        # '=' when the value is bound; overriding data_type makes the writer emit
        # t="inlineStr", so Excel shows the literal text and never evaluates it.
        cell.value = value
        cell.data_type = "s"
        cell.number_format = FMT_TEXT
        return
    cell.value = value
    if fmt and value is not None:
        cell.number_format = fmt


@dataclass
class Block:
    """Records where each label landed, so formulas reference rows rather than guesses."""

    rows: dict[str, int] = dc_field(default_factory=dict)

    def mark(self, label: str, row: int) -> int:
        # First write wins: a duplicated label would otherwise silently repoint every
        # formula built against the earlier one.
        self.rows.setdefault(label, row)
        return row

    def row(self, label: str) -> int:
        try:
            return self.rows[label]
        except KeyError as exc:  # pragma: no cover - a writer bug, not a data problem
            raise KeyError(
                f"no row recorded for label {label!r}; recorded: {sorted(self.rows)}"
            ) from exc

    def has(self, label: str) -> bool:
        return label in self.rows


def sheet_header(
    ws: Worksheet,
    title: str,
    deal_name: str,
    as_of: str,
    *,
    note: str | None = None,
) -> int:
    """Write the standard three-row chrome. Returns the first data row (5)."""
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws["A2"] = f"{deal_name} — as of {as_of}"
    ws["A2"].font = SUBTLE
    if note:
        ws["A3"] = note
        ws["A3"].font = SUBTLE
    for col in range(1, 9):
        ws.cell(row=4, column=col).border = UNDERLINE
    return 5


def section_bar(ws: Worksheet, row: int, text: str, width: int = 3) -> int:
    """A filled section bar spanning ``width`` columns. Returns the next row."""
    for col in range(1, width + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = SECTION_FILL
        cell.font = SECTION_FONT
    ws.cell(row=row, column=1).value = text
    return row + 1


def header_row(ws: Worksheet, row: int, headers: list[str]) -> int:
    """Write a bold, filled, boxed header row. Returns the next row."""
    for i, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=i)
        cell.value = text
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.border = BOX
        cell.alignment = WRAP
    return row + 1


def label_value(
    ws: Worksheet,
    row: int,
    label: str,
    value: Any,
    *,
    fmt: str | None = None,
    note: str | None = None,
    source: str | None = None,
    block: Block | None = None,
    bold: bool = False,
) -> int:
    """Label in A, value in B, note in C, source in D. Returns the next row."""
    a = ws.cell(row=row, column=1)
    assign(a, label)
    a.font = BOLD if bold else BODY_FONT
    b = ws.cell(row=row, column=2)
    assign(b, value, fmt)
    b.font = BODY_FONT
    if note is not None:
        c = ws.cell(row=row, column=3, value=note)
        c.font = SUBTLE
        c.alignment = WRAP
    if source:
        # Provenance BESIDE the value, not only in a trail tab. A reader deciding whether
        # to trust a number is looking at the number, and a source they have to go to
        # another sheet for is a source most readers never see.
        d = ws.cell(row=row, column=4, value=source)
        d.font = SUBTLE
        d.alignment = WRAP
    if block is not None:
        block.mark(label, row)
    return row + 1


def write_cells(
    ws: Worksheet,
    row: int,
    values: list[Any],
    *,
    formats: list[str | None] | None = None,
    bold: bool = False,
    fill: PatternFill | None = None,
) -> int:
    """Write one data row across columns A.. Returns the next row."""
    for i, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=i)
        fmt = None
        if formats is not None and i - 1 < len(formats):
            fmt = formats[i - 1]
        # `assign` decides type and format together: a formula keeps its numeric format,
        # source text stays text (the source said "n/a", and coercing that to a date or a
        # zero would invent data) and can never be guessed into a formula.
        assign(cell, value, fmt)
        cell.font = BOLD if bold else BODY_FONT
        if fill is not None:
            cell.fill = fill
    return row + 1


def set_widths(ws: Worksheet, widths: dict[str, float]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def autosize(ws: Worksheet, first_col: int, last_col: int, *, cap: float = 42.0) -> None:
    """Width from the longest cell in each column, capped."""
    for col in range(first_col, last_col + 1):
        longest = 0
        for cell in ws[get_column_letter(col)]:
            if cell.value is None:
                continue
            longest = max(longest, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col)].width = min(
            max(10.0, longest + 2), cap
        )


def freeze(ws: Worksheet, cell: str) -> None:
    ws.freeze_panes = cell


def col(idx: int) -> str:
    return get_column_letter(idx)


#: Fallback timestamp when a deal supplies no generation date. Any fixed value works; what
#: matters is that it is not the wall clock.
_EPOCH = _dt.datetime(2026, 1, 1, 0, 0, 0)


def _stamp(generated_on: str | None) -> _dt.datetime:
    if generated_on:
        try:
            return _dt.datetime.fromisoformat(generated_on)
        except ValueError:
            pass
    return _EPOCH


def _clamp_to_zip_range(stamp: _dt.datetime) -> _dt.datetime:
    """A whole valid datetime inside the ZIP-representable range.

    Deterministic and total: any input maps to exactly one output, and the output is
    always a real date. Clamping the year in isolation is what fails on a leap day.
    """
    if stamp.year > ZIP_MAX_YEAR:
        return _dt.datetime(ZIP_MAX_YEAR, 12, 31, 23, 59, 59)
    if stamp.year < ZIP_MIN_YEAR:
        return _dt.datetime(ZIP_MIN_YEAR, 1, 1, 0, 0, 0)
    return stamp


_MODIFIED_RE = re.compile(
    rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)"
)


def _normalize_archive(path: Path, stamp: _dt.datetime) -> None:
    """Make the saved file a function of its content alone, not of when it was written.

    An xlsx is a zip, and there are **two** independent clock sources in it:

    1. every zip entry carries its own modification time, taken from the wall clock at
       save, so two builds of identical content differ as soon as they land in different
       seconds;
    2. ``docProps/core.xml`` carries ``dcterms:modified``, which openpyxl overwrites with
       ``datetime.now()`` inside its own save path (``writer/excel.py``) — so setting the
       property *before* saving cannot work, and this is the only place to pin it.

    Missing the second one is how the first version of this fix appeared to work: the two
    builds in the test happened to land in the same second. Entry order and compression
    are preserved, so only the timestamps change.
    """
    # Backstop, not the primary control: schema.validate rejects a generation date
    # outside the ZIP range, so a deal cannot reach here with one. Clamping anyway means
    # a caller that bypasses validation gets a valid file rather than a ValueError raised
    # *after* the workbook has already been written to disk.
    #
    # Clamp to a whole valid date, never by replacing the year alone: 2400-02-29 is a real
    # date whose year clamps to 2107, which is not a leap year, so replace(year=...) would
    # raise "day is out of range for month" — from inside the very backstop meant to stop
    # a half-written artifact.
    clamped = _clamp_to_zip_range(stamp)
    date_time = (
        clamped.year,
        clamped.month,
        clamped.day,
        clamped.hour,
        clamped.minute,
        clamped.second,
    )
    iso = clamped.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    with zipfile.ZipFile(path) as src:
        entries = [(info, src.read(info.filename)) for info in src.infolist()]
    tmp = path.with_name(path.name + ".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as out:
        for info, data in entries:
            if info.filename == "docProps/core.xml":
                data = _MODIFIED_RE.sub(rb"\g<1>" + iso + rb"\g<2>", data)
            pinned = zipfile.ZipInfo(info.filename, date_time=date_time)
            pinned.compress_type = info.compress_type
            pinned.external_attr = info.external_attr
            pinned.internal_attr = info.internal_attr
            pinned.create_system = info.create_system
            out.writestr(pinned, data)
    tmp.replace(path)


def save_workbook(wb, out_path: Path, generated_on: str | None = None) -> Path:
    """The single save path: strip metadata, pin both clocks, write.

    Every writer goes through here so a workbook's bytes are a function of the deal alone.
    """
    clean_properties(wb, generated_on)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    _normalize_archive(out_path, _stamp(generated_on))
    return out_path


def clean_properties(wb, generated_on: str | None = None) -> None:
    """Strip authorship metadata and pin the timestamps.

    Two jobs. First, authorship: openpyxl stamps a creator on every workbook it writes, and
    document properties are a real leak channel precisely because they are invisible on
    screen — the scrub gate checks them for that reason.

    Second, reproducibility. openpyxl also stamps ``created``/``modified`` from the wall
    clock, which makes every rebuild a different file. That costs two things worth keeping:
    a regenerated workbook shows up as modified in git even when nothing changed, so a real
    change becomes indistinguishable from a rebuild; and it leaks when the file was built.
    Pinning both to the deal's own generation date fixes it, and is the honest value anyway.
    """
    props = wb.properties
    props.creator = "AI Underwriting Pipeline"
    props.lastModifiedBy = "AI Underwriting Pipeline"
    props.title = None
    props.subject = None
    props.description = None
    props.keywords = None
    props.category = None
    props.company = None
    props.manager = None

    stamp = _stamp(generated_on)
    props.created = stamp
    props.modified = stamp


def write_audit_trail(wb, deal: dict, workbook_key: str, *, title: str) -> None:
    """Add an "Audit Trail" tab: every value in this workbook and where it came from.

    The tab is generated from the deal's provenance stamps, so it cannot drift from the
    numbers — if a value is not stamped, the schema refused the deal long before here.
    """
    import provenance as _prov

    ws = wb.create_sheet("Audit Trail")
    stamps = _prov.for_workbook(deal.get("provenance") or {}, workbook_key)
    stamped, required, pct = _prov.coverage(deal.get("provenance") or {})

    row = sheet_header(
        ws,
        "Audit trail",
        deal["meta"]["deal_name"],
        deal["meta"]["as_of_date"],
        note=(
            f"{title}. Every value below names the document it came from and where in it. "
            f"A derived value names the values it was built from — one that could not "
            f"would not have passed validation."
        ),
    )
    ws.cell(row=row, column=1, value="Stamped values in this workbook").font = BOLD
    ws.cell(row=row, column=2, value=len(stamps)).number_format = FMT_COUNT
    row += 1
    ws.cell(row=row, column=1, value="Coverage of the required set").font = BOLD
    ws.cell(row=row, column=2, value=f"{stamped}/{required} ({pct:.0f}%)")
    row += 2

    import schema as _schema

    row = header_row(ws, row, list(_prov.AUDIT_COLUMNS))
    for path, st in stamps:
        # What the deal ASSERTS at the path, not merely what it holds: a not_asserted
        # closing-cost amount publishes as MISSING here, never 0.
        cells = _prov.as_row(path, st, _schema.asserted_value(deal, path))
        write_cells(ws, row, cells, formats=[None] * len(cells))
        ws.cell(row=row, column=7).alignment = WRAP
        ws.cell(row=row, column=1).font = Font(name="Calibri", size=10)
        row += 1

    set_widths(ws, {"A": 34, "B": 22, "C": 34, "D": 26, "E": 11, "F": 34, "G": 62})
    freeze(ws, "A9")


ILLUSTRATIVE_MARK = "(illustrative)"
