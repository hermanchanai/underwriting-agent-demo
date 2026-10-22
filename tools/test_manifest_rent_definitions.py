"""The field manifest states the rent definitions the engine computes.

In-place rent is the mean contract rent over OCCUPIED units
(``compute.roll_totals``), and effective trade-out is measured against the prior EFFECTIVE
rent (``compute.trade_out``: the prior gross when no prior concession is stated). The
manifest still called in-place rent "unit-weighted across the property", measured the
effective change "over prior rent", and had no field for ``prior_rent_effective``. The text
is printed into the published model-ready workbook, the pack's manifest sheet and the manifest
preview, so this builds those from the published interchange file and checks them.

Usage: python tools/test_manifest_rent_definitions.py
"""

from __future__ import annotations

import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
for p in (str(REPO), str(HERE)):
    if p not in sys.path:
        sys.path.insert(0, p)

from _testkit import Checker, guard  # noqa: E402


def main() -> int:
    import schema

    c = Checker("manifest rent definitions match the engine")
    by_name = {f.name: f for f in schema.MANIFEST_FIELDS}
    by_path = {f.path: f for f in schema.MANIFEST_FIELDS if f.path}

    c.section("in-place rent is the mean over occupied units")
    d = by_name["wavg_in_place_rent"].definition
    c.ok("wavg_in_place_rent says OCCUPIED units", "occupied" in d.lower(), d)
    c.ok("and no longer 'across the property'", "across the property" not in d.lower(), d)
    # compute.mean skips an occupied unit whose contract rent is not stated.
    c.ok("and only those with a stated contract rent", "stated contract rent" in d.lower(), d)

    c.section("effective trade-out is over the prior EFFECTIVE rent")
    d = by_name["cohort_effective_change_pct"].definition
    c.ok("cohort_effective_change_pct says prior EFFECTIVE", "prior effective" in d.lower(), d)
    c.ok("and says the prior gross stands in when no concession is stated",
         "prior gross" in d.lower(), d)
    d = by_name["cohort_gross_change_pct"].definition
    c.ok("gross change stays over prior (gross) rent", "prior" in d.lower(), d)

    c.section("prior_rent_effective has a field")
    f = by_path.get("trade_out.rows[0].prior_rent_effective")
    c.ok("a Field points at trade_out.rows[0].prior_rent_effective", f is not None)
    if f is not None:
        c.equal("it is a currency", f.type, "currency")
        c.equal("it sits in the Trade-out group", f.group, "Trade-out")
        c.ok("it says it is optional", "optional" in f.definition.lower(), f.definition)

    c.section("every trade-out field reads internal-only, the derived ones too")
    # Public status keys off Field.path; a derived field with no path read "yes" while every
    # other trade-out field reads internal-only.
    for f in (x for x in schema.MANIFEST_FIELDS if x.group == "Trade-out"):
        status = schema.public_demo_status(f.path)
        c.ok(f"{f.name} is internal-only in the manifest", status.startswith("Internal pipeline only"), status)

    c.section("the formats doc agrees")
    formats = (REPO / "formats" / "FORMATS.md").read_text(encoding="utf-8")
    c.ok("FORMATS.md does not call in-place rent a weighted average",
         "weighted average in-place rent" not in formats)
    c.ok("FORMATS.md says the in-place mean is over units with a stated contract rent",
         "occupied units with a stated contract rent" in formats)

    c.section("the published writers print the corrected text")
    # Built here from the published interchange file, never read from outputs/ (gitignored,
    # absent on a clean checkout). test_sample_run pins the served set == outputs/fixture.
    import json
    import tempfile

    import render_previews
    import write_manifest
    import write_pack
    from openpyxl import load_workbook

    deal = json.loads((HERE / "fixtures" / "synthetic_full.json").read_text())
    names = ("wavg_in_place_rent", "cohort_effective_change_pct", "cohort_prior_rent_effective")
    with tempfile.TemporaryDirectory() as tmp:
        built = {
            "manifest workbook": (write_manifest.build(Path(tmp) / "m.xlsx", deal), "Model-Ready Outputs"),
            "pack's manifest sheet": (write_pack.build(deal, Path(tmp) / "p.xlsx"), "Model-ready outputs"),
        }
        for label, (path, sheet) in built.items():
            rows = {r[1]: r for r in load_workbook(path, read_only=True)[sheet].iter_rows(
                values_only=True) if r and len(r) > 3}
            for name in names:
                got, want = rows.get(name), by_name.get(name)
                c.equal(f"{label}: row {name} == schema definition",
                        got[3] if got else None, want.definition if want else "<no field>")
    # The preview lists field names, not definitions, so the row is what it can show.
    html = render_previews.manifest_fragment(deal)
    c.ok("manifest preview names prior_rent_effective's field", "cohort_prior_rent_effective" in html)
    return c.finish()


if __name__ == "__main__":
    sys.exit(guard(main))
