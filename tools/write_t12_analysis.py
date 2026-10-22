"""JSON -> `sample-t12-analysis.xlsx`.

Three tabs: the monthly matrix, the P&L rollup, and the checks-plus-pricing surface. The
split is the layout — data, rollup, and verification each get their own sheet.

Two formula rules carry the weight and are visible in the workbook rather than baked into
a value:

* Gross Scheduled Rent is ``=IF(COUNTIF(categories, label) > 0, SUMIF(...), GPR + GLL)``
  — printed where the statement prints it, derived where it does not. The gate is the
  row's EXISTENCE, never whether its total happens to be nonzero: a printed line that
  nets to zero, or that has a zero month, is still a printed line.
* Total Revenue sums from Gross Scheduled Rent down through Revenue Adjustments, never
  from Gross Potential Rent, which would count loss-to-lease twice.

Usage:
    python tools/write_t12_analysis.py --deal <deal.json> --out <path.xlsx>
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import Workbook

import compute
import schema
from _xl import (
    BOLD,
    FMT_COUNT,
    FMT_CURRENCY,
    FMT_PCT,
    FMT_PCT2,
    FMT_TEXT,
    ILLUSTRATIVE_MARK,
    SUBTLE,
    SUBTOTAL_FILL,
    WRAP,
    Block,
    save_workbook,
    write_audit_trail,
    col,
    freeze,
    header_row,
    label_value,
    section_bar,
    set_widths,
    sheet_header,
    write_cells,
)

SHEET_DATA = "T-12 Data"
SHEET_PNL = "P&L Summary"
SHEET_CHECKS = "Checks & Pricing"

MONTH_FIRST_COL = 3  # column C
REVENUE_HELPER_LABEL = "Monthly revenue (computed)"


def _month_cols(count: int) -> tuple[int, int]:
    return MONTH_FIRST_COL, MONTH_FIRST_COL + count - 1


def _write_data_tab(wb: Workbook, deal: dict) -> dict:
    """The monthly matrix. Returns row/column bookkeeping the other tabs need."""
    ws = wb.active
    ws.title = SHEET_DATA
    t12 = deal["t12"]
    labels = t12["month_labels"]
    order = compute.detect_month_order(labels)
    first_m, last_m = _month_cols(len(labels))
    t3_first, t3_last = compute.recent_span(labels, 3)
    t6_first, t6_last = compute.recent_span(labels, 6)

    row = sheet_header(
        ws,
        "T-12 data",
        deal["meta"]["deal_name"],
        deal["meta"]["as_of_date"],
        note=(
            f"Months run {order.replace('_', '-')}, detected from the labels rather than "
            f"assumed — the trailing-three window depends on it."
        ),
    )
    headers = ["Line item", "Category", *labels, "T-12 total", "T-3 total"]
    row = header_row(ws, row, headers)
    data_first = row

    total_col = last_m + 1
    t3_col = last_m + 2
    formats = [None, None] + [FMT_CURRENCY] * len(labels) + [FMT_CURRENCY, FMT_CURRENCY]

    for line in t12["rows"]:
        values = [line["label"], line.get("category") or ""]
        values.extend(line["monthly"])
        values.append(f"=SUM({col(first_m)}{row}:{col(last_m)}{row})")
        values.append(
            f"=SUM({col(first_m + t3_first)}{row}:{col(first_m + t3_last)}{row})"
        )
        uncategorized = not line.get("category")
        write_cells(
            ws,
            row,
            values,
            formats=formats,
            bold=uncategorized,
            fill=SUBTOTAL_FILL if uncategorized else None,
        )
        row += 1
    data_last = row - 1
    row += 1

    # A monthly revenue row, expressed as the same SUMIF logic the rollup uses, so the
    # trailing-6 trend on the checks tab is a formula rather than a pasted number.
    p = compute.pnl(deal)
    cat_col = "B"
    span = f"${cat_col}${data_first}:${cat_col}${data_last}"

    def sumif(month_col: str, category: str) -> str:
        return (
            f'SUMIF({span},"{category}",'
            f"{month_col}${data_first}:{month_col}${data_last})"
        )

    helper_row = row
    ws.cell(row=helper_row, column=1, value=REVENUE_HELPER_LABEL).font = BOLD
    ws.cell(
        row=helper_row, column=2, value="derived"
    ).font = SUBTLE
    tail = [c for c in p["revenue_span"] if c != "Gross Scheduled Rent"]
    # Existence gate, not a per-month aggregate one. A printed gross-scheduled-rent row
    # with a zero in some month is still a printed row for that month; a nonzero-total
    # test applied per column would substitute the fallback for that month alone, and the
    # T-6 check built on this row would then disagree with compute.revenue_for_window,
    # which treats the printed row as authoritative for every month. Same gate as the
    # rollup and as compute.gsr_is_printed (I6).
    gsr_exists = f'COUNTIF({span},"Gross Scheduled Rent")>0'
    for i in range(len(labels)):
        c = col(first_m + i)
        gsr = (
            f"IF({gsr_exists},"
            f'{sumif(c, "Gross Scheduled Rent")},'
            f'{sumif(c, "Gross Potential Rent")}+{sumif(c, "Gain / (Loss) to Lease")})'
        )
        parts = [gsr] + [sumif(c, cat) for cat in tail]
        cell = ws.cell(row=helper_row, column=first_m + i, value="=" + "+".join(parts))
        cell.number_format = FMT_CURRENCY
        cell.font = BOLD
    for target, label in ((total_col, "T-12"), (t3_col, "T-3")):
        cell = ws.cell(row=helper_row, column=target)
        if label == "T-12":
            cell.value = f"=SUM({col(first_m)}{helper_row}:{col(last_m)}{helper_row})"
        else:
            cell.value = (
                f"=SUM({col(first_m + t3_first)}{helper_row}:"
                f"{col(first_m + t3_last)}{helper_row})"
            )
        cell.number_format = FMT_CURRENCY
        cell.font = BOLD
    ws.cell(
        row=helper_row + 1,
        column=1,
        value=(
            "Revenue by month, built from the same category SUMIFs as the rollup. Gross "
            "scheduled rent falls back to gross potential rent plus gain/(loss) to lease "
            "where the statement did not print it."
        ),
    ).font = SUBTLE

    set_widths(ws, {"A": 34, "B": 24})
    for c in range(first_m, t3_col + 1):
        ws.column_dimensions[col(c)].width = 13
    freeze(ws, "C6")

    return {
        "data_first": data_first,
        "data_last": data_last,
        "first_m": first_m,
        "last_m": last_m,
        "total_col": total_col,
        "t3_col": t3_col,
        "helper_row": helper_row,
        "t6_span": (t6_first, t6_last),
        "labels": labels,
    }


def _write_pnl_tab(wb: Workbook, deal: dict, ref: dict) -> Block:
    ws = wb.create_sheet(SHEET_PNL)
    t12 = deal["t12"]
    # A NUMBER or nothing. Since round 19's ruling 2 `t12.units` may be the MISSING
    # sentinel, and interpolating that into a formula wrote `=IFERROR(B6/MISSING,"")` —
    # a reference to a defined name that does not exist. IFERROR swallowed it, so the
    # column came out blank and the mistake was invisible; a workbook that ever did
    # define that name would have computed a wrong per-unit figure instead.
    units = compute.as_number(t12["units"])
    p = compute.pnl(deal)

    row = sheet_header(
        ws,
        "P&L summary",
        deal["meta"]["deal_name"],
        deal["meta"]["as_of_date"],
        note=(
            "Every category row is a SUMIF over the T-12 Data tab, so the arithmetic is "
            "auditable in the workbook rather than pasted in."
        ),
    )
    row = header_row(
        ws,
        row,
        ["Line item", "T-12", "T-3 annualized", "Per unit (T-12)", "% of revenue"],
    )

    tcol = col(ref["total_col"])
    t3col = col(ref["t3_col"])
    span_cat = f"'{SHEET_DATA}'!$B${ref['data_first']}:$B${ref['data_last']}"
    span_t12 = f"'{SHEET_DATA}'!${tcol}${ref['data_first']}:${tcol}${ref['data_last']}"
    span_t3 = f"'{SHEET_DATA}'!${t3col}${ref['data_first']}:${t3col}${ref['data_last']}"

    block = Block()
    subtotals = {
        "Gross Scheduled Rent",
        "Total Revenue",
        schema.PNL_TOTAL_OPEX_LABEL,
        schema.PNL_NOI_LABEL,
    }

    def sumif(span_values: str, at_row: int) -> str:
        """SUMIF keyed on the label sitting in column A of ``at_row``.

        Keying on the cell rather than a literal string is what makes the rollup survive
        a relabel: change the text in column A and the formula follows it.
        """
        return f"SUMIF({span_cat},$A{at_row},{span_values})"

    def mapped(at_row: int) -> str:
        """Whether any data line carries the category in column A of ``at_row``. A category
        no line was mapped to is left BLANK, never 0 — the same rule
        compute.pnl applies, so the workbook and the preview agree (I6)."""
        return f"COUNTIF({span_cat},$A{at_row})>0"

    for label in schema.pnl_row_order():
        is_subtotal = label in subtotals
        block.mark(label, row)
        a = ws.cell(row=row, column=1, value=label)
        a.font = BOLD if is_subtotal else None
        if is_subtotal:
            for c in range(1, 6):
                ws.cell(row=row, column=c).fill = SUBTOTAL_FILL
                ws.cell(row=row, column=c).font = BOLD
        row += 1

    def put(label: str, b: str, c: str) -> None:
        r = block.row(label)
        cell_b = ws.cell(row=r, column=2, value=b)
        cell_b.number_format = FMT_CURRENCY
        cell_c = ws.cell(row=r, column=3, value=c)
        cell_c.number_format = FMT_CURRENCY
        d = ws.cell(
            row=r, column=4,
            value=(f'=IFERROR(B{r}/{units:g},"")' if units else schema.MISSING),
        )
        d.number_format = FMT_CURRENCY
        e = ws.cell(
            row=r, column=5, value=f'=IFERROR(B{r}/B{block.row("Total Revenue")},"")'
        )
        e.number_format = FMT_PCT

    for cat in schema.T12_REVENUE_CATEGORIES:
        if cat == "Gross Scheduled Rent":
            gpr = block.row("Gross Potential Rent")
            gll = block.row("Gain / (Loss) to Lease")
            r = block.row(cat)
            # Both columns gate on whether a GSR line EXISTS, via COUNTIF -- not on
            # whether its total is nonzero. A printed line whose twelve months net to
            # zero is still a printed line, and testing the aggregate would silently
            # substitute the fallback into the T-3 window. This mirrors
            # compute.gsr_is_printed exactly (I6).
            exists = f"COUNTIF({span_cat},$A{r})>0"
            # The fallback is blank when neither GPR nor GLL was mapped. An absent
            # input cell is "", and "" + a number is #VALUE!, so each is read as 0 explicitly.
            either = f"COUNTIF({span_cat},$A{gpr})+COUNTIF({span_cat},$A{gll})>0"

            def z(ref: str) -> str:
                return f'IF({ref}="",0,{ref})'

            put(
                cat,
                f'=IF({exists},{sumif(span_t12, r)},IF({either},{z(f"B{gpr}")}+{z(f"B{gll}")},""))',
                f'=IF({exists},{sumif(span_t3, r)}*4,IF({either},{z(f"C{gpr}")}+{z(f"C{gll}")},""))',
            )
            continue
        if cat == "Total Revenue":
            gsr = block.row("Gross Scheduled Rent")
            adj = block.row("Revenue Adjustments")
            put(cat, f"=SUM(B{gsr}:B{adj})", f"=SUM(C{gsr}:C{adj})")
            continue
        r = block.row(cat)
        put(cat, f'=IF({mapped(r)},{sumif(span_t12, r)},"")', f'=IF({mapped(r)},{sumif(span_t3, r)}*4,"")')

    for cat in schema.T12_EXPENSE_CATEGORIES:
        r = block.row(cat)
        put(cat, f'=IF({mapped(r)},{sumif(span_t12, r)},"")', f'=IF({mapped(r)},{sumif(span_t3, r)}*4,"")')

    first_exp = block.row(schema.T12_EXPENSE_CATEGORIES[0])
    last_exp = block.row(schema.T12_EXPENSE_CATEGORIES[-1])
    put(
        schema.PNL_TOTAL_OPEX_LABEL,
        f"=SUM(B{first_exp}:B{last_exp})",
        f"=SUM(C{first_exp}:C{last_exp})",
    )
    rev = block.row("Total Revenue")
    opex = block.row(schema.PNL_TOTAL_OPEX_LABEL)
    put(schema.PNL_NOI_LABEL, f"=B{rev}-B{opex}", f"=C{rev}-C{opex}")

    row += 1
    ws.cell(
        row=row,
        column=1,
        value=(
            f"Gross scheduled rent basis: {p['gsr_basis']}. Total revenue sums from gross "
            f"scheduled rent down through revenue adjustments — not from gross potential "
            f"rent, which would double-count the loss-to-lease line."
        ),
    ).font = SUBTLE
    row += 2

    # ---- Top-4-plus-Rest breakouts ---------------------------------------------------
    breakouts = compute.breakouts(deal)
    for cat in schema.BREAKOUT_CATEGORIES:
        data = breakouts[cat]
        row = section_bar(ws, row, f"{cat} — top {data['depth']} plus rest", width=3)
        row = header_row(ws, row, ["Line item", "T-12", "T-3 annualized"])
        for entry in data["top"]:
            write_cells(
                ws, row, [entry["label"], entry["t12"], entry["t3_annualized"]],
                formats=[None, FMT_CURRENCY, FMT_CURRENCY],
            )
            row += 1
        write_cells(
            ws,
            row,
            [
                f"All other {cat.lower()} ({data['rest_count']} lines)",
                data["rest_t12"],
                data["rest_t3_annualized"],
            ],
            formats=[None, FMT_CURRENCY, FMT_CURRENCY],
        )
        row += 1
        cat_row = block.row(cat)
        write_cells(
            ws,
            row,
            [
                f"Total {cat.lower()}",
                f"='{SHEET_PNL}'!B{cat_row}",
                f"='{SHEET_PNL}'!C{cat_row}",
            ],
            formats=[None, FMT_CURRENCY, FMT_CURRENCY],
            bold=True,
            fill=SUBTOTAL_FILL,
        )
        row += 2

    # ---- Trailing annualized T-3 revenue (E2 rule 8) ---------------------------------
    # Six figures the model takes: each month plus the two before it, times four. Values,
    # computed by compute.t3_trailing_series — the same numbers the preview shows (I6).
    series = compute.t3_trailing_series(deal)
    if series:
        row = section_bar(ws, row, "Trailing T-3 annualized revenue", width=3)
        row = header_row(ws, row, ["Month", "T-3 annualized revenue", ""])
        for e in series:
            write_cells(ws, row, [e["month"], e["t3_revenue_annualized"], None],
                        formats=[FMT_TEXT, FMT_CURRENCY, None])
            row += 1
        ws.cell(
            row=row, column=1,
            value=("Each figure is that month plus the two before it, times four — the trend "
                   "the T-3 cap rate sits at the end of. Fewer than eight months of statement "
                   "give fewer figures, never a padded zero."),
        ).font = SUBTLE
        row += 2

    set_widths(ws, {"A": 34, "B": 16, "C": 18, "D": 16, "E": 14})
    freeze(ws, "A6")
    return block


def write_error_check_section(ws, row: int, deal: dict, pnl_block: Block) -> int:
    """The T-12 audit trail: our P&L against the statement's own subtotals, then the
    statement against the rent roll. Returns the next free row."""
    # ---- Error Check -----------------------------------------------------------------
    row = section_bar(ws, row, "Error check", width=6)
    row = header_row(
        ws,
        row,
        ["Check", "Our total", "Source subtotal", "Difference", "Status", "Note"],
    )
    source = deal["t12"].get("source_subtotals") or {}
    for check in compute.error_checks(deal):
        label = check["check"]
        pnl_row = pnl_block.row(label)
        src = source.get(label)
        # Defence in depth: validate() already rejects a non-numeric subtotal, but the
        # cross-tie formulas subtract this cell, so nothing but a number may reach it.
        if isinstance(src, bool) or not isinstance(src, (int, float)):
            src = None
        write_cells(
            ws,
            row,
            [
                label,
                f"='{SHEET_PNL}'!B{pnl_row}",
                src if src is not None else None,
                # An ABSENT category (a blank P&L cell) compares as zero —
                # compute.error_checks' rule — instead of turning the status into #VALUE!.
                f'=IF(C{row}="","",IF(B{row}="",0,B{row})-C{row})',
                f'=IF(C{row}="","NOT CHECKED",IF(ABS(D{row})<=1,"OK","REVIEW"))',
                ""
                if src is not None
                else "The source statement printed no subtotal for this line.",
            ],
            formats=[None, FMT_CURRENCY, FMT_CURRENCY, FMT_CURRENCY, None, None],
        )
        row += 1
    row += 1
    ws.cell(
        row=row,
        column=1,
        value=(
            "A difference over one dollar reads REVIEW rather than quietly passing. A "
            "subtotal the source never printed reads NOT CHECKED — not OK, which would "
            "claim a tie that was never tested."
        ),
    ).font = SUBTLE
    row += 2

    # ---- Statement sanity checks (E2 rules 1, 3, 4, 5) ---------------------------------
    # Values: each guard is OK or REVIEW with the lines it names. Nothing here changes a
    # number — a flip the engine made on its own would be a claim about the statement.
    guards = compute.statement_guards(deal)
    if guards:
        row = section_bar(ws, row, "Statement sanity checks", width=6)
        row = header_row(ws, row, ["Guard", "Status", "Lines named", "Detail", "", ""])
        for g in guards:
            write_cells(ws, row, [g["guard"], g["status"], "; ".join(g["rows"]) or "—", g["detail"], None, None])
            ws.cell(row=row, column=4).alignment = WRAP
            row += 1
        ws.cell(
            row=row, column=1,
            value=("A REVIEW here is a fact about how the statement is printed, flagged for a "
                   "reader to resolve; the P&L above is the statement as categorised, unflipped."),
        ).font = SUBTLE
        row += 2

    # ---- Across the two documents: statement vs rent roll --------------------------------
    # Values, not formulas: the roll is a different workbook, so these are the same
    # numbers compute.cross_document_checks hands the preview (I6), written as figures.
    row = section_bar(ws, row, "Across the two documents: statement vs rent roll", width=6)
    row = header_row(ws, row, ["Tie", "Statement", "Rent roll", "Gap", "Status", "Basis / why not"])
    for check in compute.cross_document_checks(deal):
        write_cells(
            ws, row,
            [check["check"], check["ours"], check["source"], check["pct"], check["status"],
             check.get("reason") or check["basis"]],
            formats=[None, FMT_CURRENCY, FMT_CURRENCY, FMT_PCT, None, None],
        )
        row += 1
    ws.cell(
        row=row, column=1,
        value=(f"Figures, not formulas — the roll is another workbook. A gap inside "
               f"{compute.CROSS_DOC_BAND:.0%} reads OK; wider reads REVIEW; a side that is not "
               "there reads NOT CHECKED with the reason."),
    ).font = SUBTLE
    row += 2

    return row


def write_roll_checks_section(ws, row: int, deal: dict) -> int:
    """The rent-roll audit trail: the normalized roll against the export's own stated
    totals, then the seeded rent readback. Returns the next free row."""
    # ---- Rent-roll tie-outs ----------------------------------------------------------
    # The T-12 checks tie our P&L to the income statement's own subtotals. These tie the
    # normalized rent roll to the ROLL's own stated totals — a different document, a
    # different failure mode. Both belong on one surface: a reader asking "can I trust
    # these numbers?" should not have to find two.
    rrc = compute.rent_roll_checks(deal)
    # ---- The reconciliation verdict (E2 rule 16), as a VALUE at the top ----------------
    row = section_bar(ws, row, "Reconciliation", width=6)
    write_cells(ws, row, ["Verdict", rrc["verdict"], "Units off by", rrc["verdict_delta"]["units"],
                          "Market rent off by", rrc["verdict_delta"]["market_rent"]],
                formats=[None, None, None, FMT_COUNT, None, FMT_CURRENCY], bold=True)
    row += 1
    ws.cell(row=row, column=1, value=rrc["verdict_basis"]).font = SUBTLE
    row += 2
    row = section_bar(ws, row, "Rent-roll tie-outs", width=6)
    row = header_row(
        ws, row, ["Check", "Our calc", "Source", "Difference", "Status", "Note"],
    )
    for chk in rrc["rows"]:
        is_pct = chk.get("format") == "pct"
        fmt = FMT_PCT if is_pct else FMT_COUNT
        write_cells(
            ws, row,
            [chk["item"], chk["ours"], chk["source"], chk["diff"], chk["verdict"],
             chk.get("note") or ""],
            formats=[None, fmt, fmt, fmt, None, None],
        )
        ws.cell(row=row, column=6).alignment = WRAP
        row += 1
    row += 1
    ws.cell(row=row, column=1, value=rrc["summary"]).font = BOLD
    row += 1
    ws.cell(
        row=row, column=1,
        value=(
            "A total the source export never printed reads NOT CHECKED and stays in the "
            "denominator above — excluding it would let a check that could not run "
            "flatter the ratio. Occupancy here is PHYSICAL: "
            + rrc["occupancy_basis"]
            + ". Down, model and employee units count as vacant."
        ),
    ).font = SUBTLE
    row += 2

    # ---- Duplicate rows disclosed (E2 rule 14) --------------------------------------
    dup = compute.duplicate_disclosure(deal)
    row = section_bar(ws, row, f"Duplicate rows — {dup['total']} unit(s) listed twice", width=6)
    if dup["rows"]:
        row = header_row(ws, row, ["Unit", "Kept row", "Kept status", "Kept rent", "Rows not counted", "Why"])
        for d in dup["rows"]:
            not_counted = "; ".join(
                f"row {x['row']}: {x['status'] or 'no status'}, {compute.fmt_currency(x['rent']) if x['rent'] is not None else 'no rent'}"
                for x in d["dropped"]) or "—"
            write_cells(ws, row, [d["unit_id"], d["kept_row"], d["kept_status"] or "—",
                                  d["kept_rent"], not_counted, d["why"]],
                        formats=[FMT_TEXT, FMT_COUNT, None, FMT_CURRENCY, None, None])
            ws.cell(row=row, column=6).alignment = WRAP
            row += 1
        if dup["overflow"]:
            ws.cell(row=row, column=1,
                    value=f"… and {dup['overflow']} more unit(s) not listed here; every one is in the run's JSON.").font = SUBTLE
            row += 1
    else:
        ws.cell(row=row, column=1, value="The export lists no unit twice.").font = SUBTLE
        row += 1
    row += 1

    # ---- Rent readback ---------------------------------------------------------------
    row = section_bar(
        ws, row,
        f"Rent readback — {len(rrc['readback'])} units, sampled reproducibly", width=6,
    )
    row = header_row(
        ws, row,
        ["Unit", "Plan", "Our in-place rent", "Source cell value", "Source locator",
         "Status"],
    )
    for rb in rrc["readback"]:
        write_cells(
            ws, row,
            [rb["unit_id"], rb.get("floor_plan"), rb.get("our_rent"),
             rb.get("source_rent"), rb.get("source_locator") or "—", rb["verdict"]],
            formats=[FMT_TEXT, None, FMT_CURRENCY, FMT_CURRENCY, None, None],
        )
        row += 1
    row += 1
    ws.cell(
        row=row, column=1,
        value=(
            f"{rrc['readback_matched']} of {len(rrc['readback'])} match within "
            f"${rrc['readback_tolerance']:.2f}. Sample seeded from the roll's as-of date "
            f"({rrc['readback_seed']}), so the same roll always checks the same units and "
            f"a reviewer re-running this sees these rows."
        ),
    ).font = BOLD
    row += 1
    ws.cell(row=row, column=1, value=rrc["readback_limit_note"]).font = SUBTLE
    ws.cell(row=row, column=1).alignment = WRAP
    ws.row_dimensions[row].height = 34
    row += 2

    return row


def write_pricing_section(ws, row: int, deal: dict, ref: dict, pnl_block: Block, mark: str) -> int:
    """Pricing and cap rate, the broker pro forma, the revenue trend. Returns the next free row."""
    # ---- Pricing ---------------------------------------------------------------------
    row = section_bar(ws, row, f"Pricing and cap rate{mark}", width=6)
    block = Block()
    pr = compute.pricing(deal)
    row = label_value(
        ws, row, "Price per unit", pr["price_per_unit"], fmt=FMT_CURRENCY,
        note=deal["meta"].get("pricing_note"), block=block,
    )
    row = label_value(
        ws, row, "Unit count", pr["units_total"], fmt=FMT_COUNT, block=block
    )
    ppu = block.row("Price per unit")
    units_row = block.row("Unit count")
    # The formula multiplies two CELLS, and either can hold the MISSING sentinel — text,
    # which Excel answers with #VALUE!. `compute.pricing` has already worked out whether
    # there is a total to state, so the sentinel it returns is written straight in.
    total = pr["total_price"]
    row = label_value(
        ws, row, "Total purchase price",
        f"=B{ppu}*B{units_row}" if isinstance(total, (int, float)) else total,
        fmt=FMT_CURRENCY, block=block,
    )
    price_row = block.row("Total purchase price")
    row += 1

    rev_row = pnl_block.row("Total Revenue")
    opex_row = pnl_block.row(schema.PNL_TOTAL_OPEX_LABEL)
    noi_row = pnl_block.row(schema.PNL_NOI_LABEL)

    ladders = (
        ("T-12", "B"),
        ("T-3 annualized", "C"),
    )
    # Which resolved cap rate belongs to which ladder. `compute.pricing` answers MISSING
    # when there is no price or no count, and the formula below divides by the
    # total-price cell — which then holds text, so IFERROR blanked it.
    cap_rates = {"T-12": pr["t12_cap_rate"], "T-3 annualized": pr["t3_cap_rate"]}
    for name, pnl_col in ladders:
        row = label_value(
            ws, row, f"{name} revenue", f"='{SHEET_PNL}'!{pnl_col}{rev_row}",
            fmt=FMT_CURRENCY, block=block,
        )
        row = label_value(
            ws, row, f"{name} operating expenses",
            f"='{SHEET_PNL}'!{pnl_col}{opex_row}", fmt=FMT_CURRENCY, block=block,
        )
        row = label_value(
            ws, row, f"{name} net operating income",
            f"='{SHEET_PNL}'!{pnl_col}{noi_row}", fmt=FMT_CURRENCY, block=block,
        )
        noi_here = block.row(f"{name} net operating income")
        resolved = cap_rates[name]
        row = label_value(
            ws, row, f"{name} cap rate{mark}",
            f'=IFERROR(B{noi_here}/B{price_row},"")'
            if isinstance(resolved, (int, float)) else resolved,
            fmt=FMT_PCT2, block=block,
        )
        row += 1

    if pr["proforma"]:
        pf = deal["t12"]["broker_proforma"]
        row = label_value(
            ws, row, "Broker pro forma revenue", pf.get("revenue"), fmt=FMT_CURRENCY,
            note="Kept in its own block; never merged into the T-12.", block=block,
        )
        row = label_value(
            ws, row, "Broker pro forma operating expenses", pf.get("opex"),
            fmt=FMT_CURRENCY, block=block,
        )
        row = label_value(
            ws, row, "Broker pro forma net operating income", pf.get("noi"),
            fmt=FMT_CURRENCY, block=block,
        )
        pf_noi = block.row("Broker pro forma net operating income")
        pf_cap = pr["proforma"].get("cap_rate")
        row = label_value(
            ws, row, f"Broker pro forma cap rate{mark}",
            f'=IFERROR(B{pf_noi}/B{price_row},"")'
            if isinstance(pf_cap, (int, float)) else pf_cap,
            fmt=FMT_PCT2, block=block,
        )
        row += 1

    # ---- trailing trend --------------------------------------------------------------
    row = section_bar(ws, row, "Revenue trend", width=6)
    helper = ref["helper_row"]
    t6_first, t6_last = ref["t6_span"]
    fm = ref["first_m"]
    row = label_value(
        ws, row, "T-12 revenue", f"='{SHEET_PNL}'!B{rev_row}", fmt=FMT_CURRENCY,
        block=block,
    )
    row = label_value(
        ws,
        row,
        "T-6 revenue (annualized)",
        f"=SUM('{SHEET_DATA}'!{col(fm + t6_first)}{helper}:"
        f"{col(fm + t6_last)}{helper})*2",
        fmt=FMT_CURRENCY,
        note="Trailing six months times two, off the T-12 Data tab's computed revenue row.",
        block=block,
    )
    row = label_value(
        ws, row, "T-3 revenue (annualized)", f"='{SHEET_PNL}'!C{rev_row}",
        fmt=FMT_CURRENCY, block=block,
    )
    t12r = block.row("T-12 revenue")
    t3r = block.row("T-3 revenue (annualized)")
    row = label_value(
        ws, row, "T-3 versus T-12", f'=IFERROR(B{t3r}/B{t12r}-1,"")', fmt=FMT_PCT,
        note="Positive means the trailing three months are running ahead of the year.",
        block=block,
    )

    return row


def _write_checks_tab(wb: Workbook, deal: dict, ref: dict, pnl_block: Block) -> None:
    ws = wb.create_sheet(SHEET_CHECKS)
    illustrative = deal["meta"].get("illustrative_pricing", False)
    mark = f" {ILLUSTRATIVE_MARK}" if illustrative else ""

    row = sheet_header(
        ws,
        "Checks and pricing",
        deal["meta"]["deal_name"],
        deal["meta"]["as_of_date"],
        note=schema.TOOLKIT_NAME,
    )

    row = write_error_check_section(ws, row, deal, pnl_block)
    row = write_roll_checks_section(ws, row, deal)
    row = write_pricing_section(ws, row, deal, ref, pnl_block, mark)
    set_widths(ws, {"A": 38, "B": 18, "C": 18, "D": 16, "E": 14, "F": 52})
    freeze(ws, "A5")


def build(deal: dict, out_path: Path) -> Path:
    errors = schema.validate(deal)
    if errors:
        raise ValueError(
            "deal failed schema validation:\n  - " + "\n  - ".join(errors)
        )
    wb = Workbook()
    ref = _write_data_tab(wb, deal)
    pnl_block = _write_pnl_tab(wb, deal, ref)
    _write_checks_tab(wb, deal, ref, pnl_block)
    write_audit_trail(wb, deal, "t12_analysis", title="T-12 data, the P&L rollup and pricing")
    return save_workbook(wb, out_path, deal["meta"].get("generated_on"))


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--deal", required=True, type=Path)
    ap.add_argument("--out", required=True, type=Path)
    args = ap.parse_args(argv)
    deal = json.loads(args.deal.read_text(encoding="utf-8"))
    path = build(deal, args.out)
    print(f"wrote {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
