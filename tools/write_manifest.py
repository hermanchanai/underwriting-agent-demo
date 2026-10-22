"""Schema -> `sample-model-ready-outputs.xlsx`.

The fourth deliverable, and the argument of the whole set: the pipeline does not build a
financial model, because every shop already has one. It fills one. This workbook is the
list of values it hands over and the model input each one belongs in.

Rows come from ``schema.MANIFEST_FIELDS`` and nowhere else (I2) — the field list is not
typed twice, so a field added to the schema appears here automatically and a field removed
disappears. The `Example (demo deal)` column is filled from a deal JSON when one is
supplied and reads an em dash otherwise.

Usage:
    python tools/write_manifest.py --out <path.xlsx> [--deal <deal.json>]
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import Workbook

import compute
import schema
from _xl import (
    BOLD,
    TITLE,
    WRAP,
    save_workbook,
    write_audit_trail,
    freeze,
    header_row,
    section_bar,
    set_widths,
    write_cells,
)

SHEET = "Model-Ready Outputs"

PITCH = (
    "These are the values the pipeline extracts from a deal's raw documents and hands to "
    "your model. It does not build a model — every shop already has one, and theirs is "
    "the one their committee trusts. The work worth automating is filling it in "
    "correctly."
)

BOUNDARY = (
    "Comps come only from an operator's own weekly rent scrape. Licensed and subscription "
    "data sources are excluded by design, so a field one of them would have filled reads "
    "MISSING rather than being quietly sourced elsewhere. Every price and cap rate in the "
    "sample is illustrative."
)


def boundary(deal: dict | None) -> str:
    """``BOUNDARY``, plus the memorandum-comparables clause when the deal carries them (the
    same clause the preview's manifest uses); the published sample carries none, so its text
    is ``BOUNDARY`` byte for byte."""
    import render_previews
    if render_previews._has_om_comps(deal):
        head, _, tail = BOUNDARY.partition(".")
        return f"{head}{render_previews.OM_COMPS_BOUNDARY_CLAUSE}.{tail}"
    return BOUNDARY


def _example(deal: dict | None, field: schema.Field) -> object:
    if deal is None or not field.path:
        return "—"
    # What the deal ASSERTS, not merely what it holds — the same resolver the preview
    # fragment uses, so a not_asserted amount reads MISSING on both, never $0.
    value = schema.manifest_value(deal, field)
    if value is None:
        return "—"
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, (int, float)):
        if field.type == "percent":
            return compute.fmt_pct(value)
        if field.type == "currency":
            return compute.fmt_currency(value)
        if schema.is_year_field(field):
            return compute.fmt_year(value)
        return compute.fmt_number(value, places=1 if isinstance(value, float) else 0)
    if isinstance(value, list):
        return f"{len(value)} entries"
    text = str(value)
    return text if len(text) <= 180 else text[:177] + "…"


def write_manifest_sheet(ws, deal: dict | None, fields) -> None:
    """The field manifest onto ``ws`` (titled by the caller)."""
    ws["A1"] = schema.TOOLKIT_NAME
    ws["A1"].font = TITLE
    ws["A2"] = "Model-ready outputs — field manifest"
    ws["A2"].font = BOLD

    row = 4
    for label, text in (
        ("What this is", PITCH),
        ("What it excludes", boundary(deal)),
    ):
        ws.cell(row=row, column=1, value=label).font = BOLD
        cell = ws.cell(row=row, column=2, value=text)
        cell.alignment = WRAP
        ws.row_dimensions[row].height = 46
        row += 1

    ws.cell(row=row, column=1, value="Fields listed").font = BOLD
    ws.cell(row=row, column=2, value=len(fields))
    row += 1
    ws.cell(row=row, column=1, value="Generated").font = BOLD
    ws.cell(
        row=row,
        column=2,
        value=(deal or {}).get("meta", {}).get("generated_on")
        or dt.date.today().isoformat(),
    )
    row += 1
    ws.cell(row=row, column=1, value="Example column").font = BOLD
    if deal is None:
        note = "No deal supplied — the Example column reads em dashes."
    elif deal["meta"].get("synthetic"):
        note = (
            f"From \"{deal['meta']['deal_name']}\", a synthetic deal built to exercise the "
            f"formats. {schema.synthetic_note(deal)}"
        )
    else:
        note = (
            f"From \"{deal['meta']['deal_name']}\". Pricing is illustrative; see the deal "
            f"summary for the boundary note."
        )
    ws.cell(row=row, column=2, value=note).alignment = WRAP
    row += 2

    row = header_row(ws, row, list(schema.MANIFEST_COLUMNS))
    header_at = row - 1

    current_group: str | None = None
    for field in fields:
        if field.group != current_group:
            current_group = field.group
            row = section_bar(ws, row, field.group, width=len(schema.MANIFEST_COLUMNS))
        write_cells(
            ws,
            row,
            [
                field.group,
                field.name,
                field.type,
                field.definition,
                field.source,
                field.fills,
                schema.public_demo_status(field.path),
                (schema.with_source(deal, field.path, str(_example(deal, field)))
                 if schema.public_substitute(field.path) else _example(deal, field)),
            ],
        )
        for col_idx in (4, 6, 7, 8):
            ws.cell(row=row, column=col_idx).alignment = WRAP
        row += 1

    set_widths(
        ws,
        {"A": 22, "B": 30, "C": 11, "D": 62, "E": 32, "F": 38, "G": 34, "H": 30},
    )
    freeze(ws, f"A{header_at + 1}")


def build(out_path: Path, deal: dict | None = None) -> Path:
    if deal is not None:
        errors = schema.validate(deal)
        if errors:
            raise ValueError(
                "deal failed schema validation:\n  - " + "\n  - ".join(errors)
            )

    fields = schema.manifest_fields_for(deal)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET

    write_manifest_sheet(ws, deal, fields)
    if deal is not None:
        write_audit_trail(wb, deal, "manifest", title="The whole trail, every stamped value")
    return save_workbook(wb, out_path, (deal or {}).get("meta", {}).get("generated_on"))


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--out", required=True, type=Path)
    ap.add_argument("--deal", type=Path, default=None)
    args = ap.parse_args(argv)
    deal = (
        json.loads(args.deal.read_text(encoding="utf-8")) if args.deal else None
    )
    path = build(args.out, deal)
    print(f"wrote {path} ({len(schema.manifest_fields_for(deal))} fields)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
