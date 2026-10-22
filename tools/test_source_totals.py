#!/usr/bin/env python3
"""The totals a rent-roll export states about itself, by vendor block shape.

A tie-out against a figure this project computed is not a tie-out, so the reconciliation
gate (E2 rule 16) reads the export's OWN summary block. Two shapes so far:

* RealPage `Occupancy and Rents Summary` — a per-status table whose header names the
  columns and whose `Totals:` row is labelled.
* ResMan `Property Occupancy` — a `Type | Status | Market Rent | Units | Square Footage`
  header, a `Total / Occupied` and a `Total / Vacant` row, then an UNLABELLED totals row
  (numbers only). An earlier triage misnamed this export; its footer says ResMan.

Both are pinned on synthetic workbooks built here, so the reader's contract is the
block's shape and never a real deal's file. The real-file proof is a live run.

Usage: OPENPYXL_LXML=False python tools/test_source_totals.py
"""
from __future__ import annotations

import copy
import datetime as dt
import json
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from _testkit import Checker, guard  # noqa: E402

import openpyxl  # noqa: E402

import compute  # noqa: E402
import source_readers as ex  # noqa: E402

FIXTURE = HERE / "fixtures" / "synthetic_full.json"


def _resman_workbook(path: Path) -> tuple[int, int]:
    """A ResMan-shaped roll: unit rows with charge lines beneath, a per-column totals
    row, a `Total Charges` block, the `Property Occupancy` block, a `Unit Type
    Occupancy` block. Returns (header_row, last_unit_row)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet"
    ws["A1"] = "Sample Community"
    ws["A2"] = "Rent Roll"
    ws["A3"] = "7/31/2026"
    hdr = 5
    for c, label in enumerate(
        ["Unit", "Type", "Sq. Feet", "Residents", "Status", "Market Rent", "Description",
         "Amount", "Move In", "Lease Start", "Lease End"], start=1):
        ws.cell(hdr, c, label)
    r = hdr + 1
    units = [
        ("A-100", "2x1", 900, "Vacant Unit", None, 1890, None, 0, None, None, None, []),
        ("A-101", "2x1", 900, "Resident", "C", 1800, "Tenant Rent", 1750,
         dt.datetime(2024, 1, 15), dt.datetime(2026, 1, 15), dt.datetime(2027, 1, 14),
         [("Renter's Insurance", 15.5)]),
        ("A-102", "1x1", 650, "Resident", "NTV", 1300, "Tenant Rent", 1250,
         dt.datetime(2023, 6, 1), dt.datetime(2025, 6, 1), dt.datetime(2026, 8, 31), []),
        ("A-103", "1x1", 650, "Resident", "C", 1300, "RENT", 1275,
         dt.datetime(2022, 3, 1), dt.datetime(2026, 3, 1), dt.datetime(2027, 2, 28),
         [("Pet Rent - Dog", 35)]),
        ("A-104", "1x1", 650, "Vacant Unit", None, 1350, None, 0, None, None, None, []),
    ]
    last_unit_row = None
    for unit, plan, sqft, who, status, market, desc, amt, mi, ls, le, charges in units:
        for c, v in enumerate([unit, plan, sqft, who, status, market, desc, amt, mi, ls, le], start=1):
            if v is not None:
                ws.cell(r, c, v)
        last_unit_row = r
        r += 1
        total = amt
        for label, value in charges:
            ws.cell(r, 7, label)
            ws.cell(r, 8, value)
            total += value
            r += 1
        ws.cell(r, 7, "Total")
        ws.cell(r, 8, total)
        r += 2
    # The export's per-column totals row (market rent sums in its own column). Everything
    # below mirrors the real file's geometry: the trailer blocks put their labels in A and
    # their numbers in columns NO unit-row role is mapped to (the real export uses D/H/P/V
    # beside unit columns C/E/K/M/W), which is what keeps them out of the unit list.
    ws.cell(r, 6, 7640)
    ws.cell(r, 8, 4325.5)
    r += 1
    ws.cell(r, 1, "Total Charges"); r += 1
    ws.cell(r, 1, "Description"); ws.cell(r, 12, "Total"); r += 1
    ws.cell(r, 1, "Pet Rent"); ws.cell(r, 12, 35); r += 1
    ws.cell(r, 1, "RENT"); ws.cell(r, 12, 1275); r += 1
    ws.cell(r, 1, "Renter's Insurance"); ws.cell(r, 12, 15.5); r += 1
    ws.cell(r, 1, "Tenant Rent"); ws.cell(r, 12, 3000); r += 1
    ws.cell(r, 12, 4325.5); r += 2
    # the block under test: A=Type, L=Status, M=Market Rent, N=Units, O=Square Footage
    ws.cell(r, 1, "Property Occupancy"); r += 1
    ws.cell(r, 1, "Type"); ws.cell(r, 12, "Status"); ws.cell(r, 13, "Market Rent")
    ws.cell(r, 14, "Units"); ws.cell(r, 15, "Square Footage"); r += 2   # one blank row after the header
    ws.cell(r, 1, "Total"); ws.cell(r, 12, "Occupied"); ws.cell(r, 13, 4400); ws.cell(r, 16, 0.5759)
    ws.cell(r, 14, 3); ws.cell(r, 17, 0.6); ws.cell(r, 15, 2200); ws.cell(r, 18, 0.5867); r += 1
    ws.cell(r, 1, "Total"); ws.cell(r, 12, "Vacant"); ws.cell(r, 13, 3240); ws.cell(r, 16, 0.4241)
    ws.cell(r, 14, 2); ws.cell(r, 17, 0.4); ws.cell(r, 15, 1550); ws.cell(r, 18, 0.4133); r += 1
    ws.cell(r, 13, 7640); ws.cell(r, 14, 5); ws.cell(r, 15, 3750); r += 2       # unlabelled totals
    ws.cell(r, 1, "Unit Type Occupancy"); r += 1
    ws.cell(r, 1, "Type"); ws.cell(r, 12, "Status"); ws.cell(r, 13, "Market Rent")
    ws.cell(r, 14, "Units"); ws.cell(r, 15, "Square Footage"); r += 2
    for plan, status, market, n, sq in [("1x1", "Occupied", 2600, 2, 1300), ("1x1", "Vacant", 1350, 1, 650),
                                       ("2x1", "Occupied", 1800, 1, 900), ("2x1", "Vacant", 1890, 1, 900)]:
        ws.cell(r, 1, plan); ws.cell(r, 12, status); ws.cell(r, 13, market); ws.cell(r, 14, n); ws.cell(r, 15, sq); r += 1
    ws.cell(r, 13, 7640); ws.cell(r, 14, 5); ws.cell(r, 15, 3750); r += 2
    ws.cell(r, 1, "* denotes unit is excluded from occupancy calculations\n© ResMan, LLC")
    wb.save(path)
    return hdr, last_unit_row


def _realpage_workbook(path: Path) -> tuple[int, int]:
    """The shape the showcase export has: a labelled `Totals:` row under the units, then
    `Occupancy and Rents Summary` with a `Unit Status | LRO Market Rent | # Units |
    Current Rent` header and a labelled `Totals:` line."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "100"; ws["B1"] = "Sample Property"; ws["H1"] = "One Line Rent Roll"
    ws["A2"] = "Total Sqft:"; ws["B2"] = 3200
    hdr = 3
    for c, label in enumerate(["Bldg", "Unit ID", "Name", "Status", "Unit Status", "Sqft",
                               "Floorplan", "Occupied", "Lease Start", "Lease Rent", "LRO Market"], start=1):
        ws.cell(hdr, c, label)
    rows = [
        ("01", "1001", "Resident", "Current", "Occupied", 1000, "2X2", "9/25/2019", "6/1/2026", 1500, 1450),
        ("01", "1002", "Resident", "Current", "Notice Available", 1000, "2X2", "1/1/2021", "1/1/2026", 1525, 1450),
        ("01", "1003", "Vacant,", None, "Vacant Available", 1200, "3X2", None, None, 0, 1700),
    ]
    r = hdr + 1
    for row in rows:
        for c, v in enumerate(row, start=1):
            if v is not None:
                ws.cell(r, c, v)
        r += 1
    last_unit_row = r - 1
    ws.cell(r, 2, "Total Units:"); ws.cell(r, 3, "3"); ws.cell(r, 10, 3025); ws.cell(r, 11, 4600); r += 1
    ws.cell(r, 1, "Occupancy and Rents Summary for Current Date"); r += 1
    ws.cell(r, 1, "Unit Status"); ws.cell(r, 3, "LRO Market Rent"); ws.cell(r, 4, "# Units"); ws.cell(r, 5, "Current Rent"); r += 1
    for status, market, n, current in [("Notice Available", 1450, 1, 1525), ("Occupied", 1450, 1, 1500),
                                       ("Vacant Available", 1700, 1, 1700)]:
        ws.cell(r, 1, status); ws.cell(r, 3, market); ws.cell(r, 4, n); ws.cell(r, 5, current); r += 1
    ws.cell(r, 1, "Totals:"); ws.cell(r, 3, 4600); ws.cell(r, 4, 3); ws.cell(r, 5, 4725); r += 1
    ws.cell(r, 1, "Period Summary Billing By Transaction Code")
    wb.save(path)
    return hdr, last_unit_row


def main() -> int:
    c = Checker("test_source_totals")
    tmp = Path(tempfile.mkdtemp(prefix="source-totals-"))

    c.section("ResMan — `Property Occupancy` with an unlabelled totals row")
    rp = tmp / "resman.xlsx"
    hdr, last = _resman_workbook(rp)
    ws = openpyxl.load_workbook(rp, data_only=True)["Sheet"]
    st = ex.parse_source_totals(ws, hdr, last)
    c.equal("units_total from the Units column of the totals row", st.get("units_total"), 5)
    c.equal("market_rent from the Market Rent column of the totals row", st.get("market_rent"), 7640)
    c.equal("sqft from the Square Footage column of the totals row", st.get("sqft"), 3750)
    c.equal("buckets keyed by the Status cell of the labelled rows", st.get("buckets"), {"Occupied": 3, "Vacant": 2})
    c.ok("block_row is the `Property Occupancy` title row",
         isinstance(st.get("block_row"), int) and ws.cell(st["block_row"], 1).value == "Property Occupancy",
         str(st.get("block_row")))
    c.equal("current_rent is not stated by this block", st.get("current_rent"), None)
    c.ok("the per-floor-plan `Unit Type Occupancy` block below is not mistaken for the property block",
         st.get("units_total") == 5 and st.get("buckets", {}).get("Occupied") == 3)

    c.section("ResMan — through parse_rent_roll, the way the upload path reads it")
    columns = {"unit": 1, "plan": 2, "sqft": 3, "status": 5, "lease_status": 5,
               "market_rent": 6, "current_rent": 8, "move_in": 9, "lease_start": 10}
    units, notes, stats = ex.parse_rent_roll(
        rp, "2026-07-31", header_row=hdr, columns=columns,
        status_map={"C": "occupied", "NTV": "notice"},
    )
    c.equal("five units read (charge lines and the trailer are not units)", len(units), 5)
    c.equal("stats.source_totals carries the block", (stats.get("source_totals") or {}).get("units_total"), 5)
    stated = ex._roll_stated(stats)
    c.equal("stated.units_total", stated.get("units_total"), 5)
    c.equal("stated.market_rent", stated.get("market_rent"), 7640)
    c.equal("stated.sqft falls back to the block when nothing is stated above the header", stated.get("sqft"), 3750)
    c.equal("stated.occupied_units = the Occupied bucket", stated.get("occupied_units"), 3)
    c.equal("stated.occupied_basis names it", stated.get("occupied_basis"), "Occupied")

    c.section("ResMan — the reconciliation gate can now PASS / FAIL, not only NOT CHECKED")
    deal = json.loads(FIXTURE.read_text(encoding="utf-8"))
    d = copy.deepcopy(deal)
    d["rent_roll"]["units"] = units
    d["rent_roll"]["stated"] = stated
    rrc = compute.rent_roll_checks(d)
    c.equal("units 5 vs 5, market 7,640 vs 7,640 → PASS", rrc.get("verdict"), "PASS")
    c.equal("verdict_delta units 0", (rrc.get("verdict_delta") or {}).get("units"), 0)
    sq = next((r for r in rrc["rows"] if "square" in r["item"].lower()), None)
    c.ok("the square-footage row ties against the block's 3,750", sq is not None and sq["verdict"] == "OK",
         str(sq))
    short = copy.deepcopy(d)
    short["rent_roll"]["units"] = units[:-1]
    c.equal("one unit short → FAIL", compute.rent_roll_checks(short).get("verdict"), "FAIL")

    c.section("ResMan — a BLANK status cell is an absence, not a code")
    # This export's status column holds LEASE codes (C / NTV / R / F) and is empty on a
    # vacant unit, whose Residents cell says "Vacant Unit". The mapper's vocabulary cannot
    # name a blank, so the blank fell through _status_of to Non-revenue: 11 vacant homes
    # on the real file read as non-revenue and occupancy was 88.7% against the export's
    # own 91.1%. A blank says nothing; the unit's own evidence (no rent, no lease, no
    # move-in) says vacant — the same inference a roll with NO status column gets.
    by_id = {u["unit_id"]: u for u in units}
    c.equal("A-100: blank status, no rent, no dates → Vacant", by_id["A-100"]["status"], "Vacant")
    c.equal("A-104: likewise", by_id["A-104"]["status"], "Vacant")
    c.equal("A-101: a mapped code still goes through the vocabulary", by_id["A-101"]["status"], "Occupied")
    c.equal("A-102: NTV → notice → Occupied", by_id["A-102"]["status"], "Occupied")
    vac = next((r for r in rrc["rows"] if r["item"].lower().startswith("occupied units (physical)")), None)
    c.ok("the physical-occupancy tie reads OK, 3 of 3", vac is not None and vac["verdict"] == "OK" and vac["ours"] == 3, str(vac))
    t = compute.roll_totals(d)
    c.equal("roll totals: 2 vacant, 0 non-revenue", (t.get("vacant_units"), t.get("non_revenue_units")), (2, 0))

    c.section("a status code the vocabulary can only call `other` is inferred from evidence")
    # The mapper's vocabulary has six meanings; `other` is "none of them". On the export
    # behind one filed defect the status column holds LEASE codes and the model called `R`
    # (a renewal — a resident, paying) `other`, which _MEANING_TO_STATUS read as
    # Non-revenue: two paying homes left the occupied count. `other` is not a physical
    # state, it is the absence of one, so the unit's own evidence decides — as it does for
    # a blank cell. An explicit vacant / model / down meaning is still the model's call.
    wb2 = openpyxl.Workbook(); ws2 = wb2.active; ws2.title = "Sheet"
    for col, label in enumerate(["Unit", "Type", "Sq. Feet", "Residents", "Status", "Market Rent",
                                 "Description", "Amount", "Move In", "Lease Start"], start=1):
        ws2.cell(1, col, label)
    rows2 = [
        ("B-1", "1x1", 650, "Resident", "R", 1300, "Tenant Rent", 1250, dt.datetime(2022, 5, 1), dt.datetime(2026, 5, 1)),
        ("B-2", "1x1", 650, "Resident", "C", 1300, "Tenant Rent", 1275, dt.datetime(2024, 2, 1), dt.datetime(2026, 2, 1)),
        ("B-3", "1x1", 650, "Model", "M", 1300, None, 0, None, None),
        ("B-4", "1x1", 650, "Resident", "X", 1300, None, 0, None, None),
    ]
    for i, row in enumerate(rows2, start=2):
        for col, v in enumerate(row, start=1):
            if v is not None:
                ws2.cell(i, col, v)
    p2 = tmp / "other.xlsx"; wb2.save(p2)
    u2, n2, _ = ex.parse_rent_roll(
        p2, "2026-07-31", header_row=1, columns=columns,
        status_map={"R": "other", "C": "occupied", "M": "model", "X": "other"},
    )
    s2 = {u["unit_id"]: u["status"] for u in u2}
    c.equal("R (other) with a paying lease → Occupied", s2.get("B-1"), "Occupied")
    c.equal("C → Occupied through the vocabulary", s2.get("B-2"), "Occupied")
    c.equal("an explicit model meaning stays Non-revenue", s2.get("B-3"), "Non-revenue")
    c.equal("other with no rent and no lease → Vacant", s2.get("B-4"), "Vacant")
    c.ok("the notes say how many statuses were inferred, and why",
         any("inferred" in n and "2 unit" in n for n in n2), str(n2))

    c.section
    # A future lease is not occupancy: `F` (future) with rent and a lease starting after
    # the as-of date is a home nobody lives in yet. And a raw status that SAYS model /
    # down / unrentable / employee is non-revenue whatever meaning the vocabulary gave it.
    wb3 = openpyxl.Workbook(); ws3 = wb3.active; ws3.title = "Sheet"
    for col, label in enumerate(["Unit", "Type", "Sq. Feet", "Residents", "Status", "Market Rent",
                                 "Description", "Amount", "Move In", "Lease Start"], start=1):
        ws3.cell(1, col, label)
    rows3 = [
        ("C-1", "1x1", 650, "Resident", "F", 1500, "Tenant Rent", 1500, None, dt.datetime(2026, 12, 1)),
        ("C-2", "1x1", 650, "Resident", "F", 1500, "Tenant Rent", 1500, dt.datetime(2026, 6, 1), dt.datetime(2026, 9, 1)),
        ("C-3", "1x1", 650, "Model Unit", "Model", 1500, "Tenant Rent", 1500, dt.datetime(2020, 1, 1), dt.datetime(2026, 1, 1)),
        ("C-4", "1x1", 650, "Employee", "Employee Unit", 1500, None, 0, None, None),
        ("C-5", "1x1", 650, "Resident", None, 1500, "Tenant Rent", 1400, None, dt.datetime(2026, 8, 15)),
    ]
    for i, row in enumerate(rows3, start=2):
        for col, v in enumerate(row, start=1):
            if v is not None:
                ws3.cell(i, col, v)
    p3 = tmp / "asof.xlsx"; wb3.save(p3)
    u3, n3, _ = ex.parse_rent_roll(
        p3, "2026-07-31", header_row=1, columns=columns,
        status_map={"F": "other", "Model": "other", "Employee Unit": "other"},
    )
    s3 = {u["unit_id"]: u["status"] for u in u3}
    # an earlier review: one future date and one missing date is AMBIGUOUS — a renewal whose
    # start is next month with no move-in column is a resident, not an empty home. With
    # rent it is Unknown (out of the occupied count, and says so); without rent it is Vacant.
    c.equal("F (other), rent, lease starts AFTER the as-of date, no move-in → Unknown, never Occupied", s3.get("C-1"), "Unknown")
    c.equal("F (other), moved in BEFORE the as-of date (a renewal not yet effective) → Occupied", s3.get("C-2"), "Occupied")
    c.equal("a raw status that says Model stays Non-revenue even when the vocabulary said other", s3.get("C-3"), "Non-revenue")
    c.equal("a raw status that says Employee Unit likewise", s3.get("C-4"), "Non-revenue")
    c.equal("a BLANK status with rent and a lease starting after the as-of date → Unknown", s3.get("C-5"), "Unknown")
    c.equal("no rent and only a future lease start → Vacant (pre-leased, empty on the as-of date)",
            ex._inferred_status(None, "2026-09-01", None, "2026-07-31"), "Vacant")
    c.equal("rent, move-in stored after the as-of date, no lease start → Unknown (a lag, not a vacancy)",
            ex._inferred_status(1500, None, "2026-08-01", "2026-07-31"), "Unknown")
    # an earlier review: an explicit move-in DECIDES physical presence; the lease start is
    # only the fallback when the roll has no move-in. A lease that began in July with a
    # move-in keyed for August is not a July resident.
    c.equal("rent, lease start before the as-of date but move-in AFTER it → Unknown, not Occupied",
            ex._inferred_status(1500, "2026-07-01", "2026-08-01", "2026-07-31"), "Unknown")
    c.equal("rent, move-in before the as-of date, lease start after (a renewal) → Occupied",
            ex._inferred_status(1500, "2026-09-01", "2026-06-01", "2026-07-31"), "Occupied")

    c.section
    # Two variants the real export did not show: a totals row that carries the word
    # `Total` in column A, and a property block with NO totals row at all, both followed
    # by the `Unit Type Occupancy` table whose rows must never overwrite the property's.
    def _resman_variant(path, totals_text):
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Sheet"
        ws.cell(1, 1, "Unit"); ws.cell(1, 2, "Type"); ws.cell(1, 5, "Status"); ws.cell(1, 6, "Market Rent"); ws.cell(1, 8, "Amount")
        ws.cell(2, 1, "A-1"); ws.cell(2, 2, "1x1"); ws.cell(2, 5, "C"); ws.cell(2, 6, 1000); ws.cell(2, 8, 950)
        r = 5
        ws.cell(r, 1, "Property Occupancy"); r += 1
        ws.cell(r, 1, "Type"); ws.cell(r, 12, "Status"); ws.cell(r, 13, "Market Rent"); ws.cell(r, 14, "Units"); ws.cell(r, 15, "Square Footage"); r += 2
        ws.cell(r, 1, "Total"); ws.cell(r, 12, "Occupied"); ws.cell(r, 13, 4400); ws.cell(r, 14, 3); ws.cell(r, 15, 2200); r += 1
        ws.cell(r, 1, "Total"); ws.cell(r, 12, "Vacant"); ws.cell(r, 13, 3240); ws.cell(r, 14, 2); ws.cell(r, 15, 1550); r += 1
        if totals_text is not None:
            if totals_text:
                ws.cell(r, 1, totals_text)
            ws.cell(r, 13, 7640); ws.cell(r, 14, 5); ws.cell(r, 15, 3750); r += 1
        r += 1
        ws.cell(r, 1, "Unit Type Occupancy"); r += 1
        ws.cell(r, 1, "Type"); ws.cell(r, 12, "Status"); ws.cell(r, 13, "Market Rent"); ws.cell(r, 14, "Units"); ws.cell(r, 15, "Square Footage"); r += 2
        for plan, status, market, n, sq in [("1x1", "Occupied", 2600, 1, 1300), ("1x1", "Vacant", 1350, 1, 650)]:
            ws.cell(r, 1, plan); ws.cell(r, 12, status); ws.cell(r, 13, market); ws.cell(r, 14, n); ws.cell(r, 15, sq); r += 1
        ws.cell(r, 13, 3950); ws.cell(r, 14, 2); ws.cell(r, 15, 1950)
        wb.save(path)
        return openpyxl.load_workbook(path, data_only=True)["Sheet"]
    ws_a = _resman_variant(tmp / "resman_total_text.xlsx", "Total")
    st_a = ex.parse_source_totals(ws_a, 1, 2)
    c.equal("a totals row that says `Total` in column A is the totals row", (st_a.get("units_total"), st_a.get("market_rent"), st_a.get("sqft")), (5, 7640, 3750))
    c.equal("and the buckets are the property's, not the per-plan table's", st_a.get("buckets"), {"Occupied": 3, "Vacant": 2})
    # an earlier review, three more shapes: (i) the next table's title row also carries its
    # header labels (Status text on the title row) with no totals row and no blank before
    # it; (ii) a blank row BETWEEN the bucket rows, then the totals; (iii) the totals row
    # says `Total` in the Status column.
    def _resman_shape(path, shape):
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Sheet"
        ws.cell(1, 1, "Unit"); ws.cell(1, 2, "Type"); ws.cell(1, 5, "Status"); ws.cell(1, 6, "Market Rent"); ws.cell(1, 8, "Amount")
        ws.cell(2, 1, "A-1"); ws.cell(2, 2, "1x1"); ws.cell(2, 5, "C"); ws.cell(2, 6, 1000); ws.cell(2, 8, 950)
        r = 5
        ws.cell(r, 1, "Property Occupancy"); r += 1
        ws.cell(r, 1, "Type"); ws.cell(r, 12, "Status"); ws.cell(r, 13, "Market Rent"); ws.cell(r, 14, "Units"); ws.cell(r, 15, "Square Footage"); r += 2
        ws.cell(r, 1, "Total"); ws.cell(r, 12, "Occupied"); ws.cell(r, 13, 4400); ws.cell(r, 14, 3); ws.cell(r, 15, 2200); r += 1
        if shape == "blank_between":
            r += 1
        ws.cell(r, 1, "Total"); ws.cell(r, 12, "Vacant"); ws.cell(r, 13, 3240); ws.cell(r, 14, 2); ws.cell(r, 15, 1550); r += 1
        if shape == "total_in_status":
            ws.cell(r, 12, "Total"); ws.cell(r, 13, 7640); ws.cell(r, 14, 5); ws.cell(r, 15, 3750); r += 1
        elif shape == "blank_between":
            ws.cell(r, 13, 7640); ws.cell(r, 14, 5); ws.cell(r, 15, 3750); r += 1
        if shape != "title_header_row":
            r += 1
            ws.cell(r, 1, "Unit Type Occupancy"); r += 1
            ws.cell(r, 1, "Type"); ws.cell(r, 12, "Status"); ws.cell(r, 13, "Market Rent"); ws.cell(r, 14, "Units"); ws.cell(r, 15, "Square Footage"); r += 2
        else:
            ws.cell(r, 1, "Unit Type Occupancy"); ws.cell(r, 12, "Status"); ws.cell(r, 13, "Market Rent"); ws.cell(r, 14, "Units"); ws.cell(r, 15, "Square Footage"); r += 1
        for plan, status, market, n, sq in [("1x1", "Occupied", 2600, 1, 1300), ("1x1", "Vacant", 1350, 1, 650)]:
            ws.cell(r, 1, plan); ws.cell(r, 12, status); ws.cell(r, 13, market); ws.cell(r, 14, n); ws.cell(r, 15, sq); r += 1
        ws.cell(r, 13, 3950); ws.cell(r, 14, 2); ws.cell(r, 15, 1950)
        wb.save(path)
        return openpyxl.load_workbook(path, data_only=True)["Sheet"]
    st_c = ex.parse_source_totals(_resman_shape(tmp / "resman_title_header.xlsx", "title_header_row"), 1, 2)
    c.equal("(i) a title row that also carries `Status` ends the block: the property's buckets stand", st_c.get("buckets"), {"Occupied": 3, "Vacant": 2})
    c.equal("(i) and no totals are taken from the per-plan table", (st_c.get("units_total"), st_c.get("market_rent")), (None, None))
    st_d = ex.parse_source_totals(_resman_shape(tmp / "resman_blank_between.xlsx", "blank_between"), 1, 2)
    c.equal("(ii) a blank row between the bucket rows does not end the block", st_d.get("buckets"), {"Occupied": 3, "Vacant": 2})
    c.equal("(ii) and the totals below are read", (st_d.get("units_total"), st_d.get("market_rent"), st_d.get("sqft")), (5, 7640, 3750))
    st_e = ex.parse_source_totals(_resman_shape(tmp / "resman_total_in_status.xlsx", "total_in_status"), 1, 2)
    c.equal("(iii) `Total` in the Status column is the totals row, not a bucket", st_e.get("buckets"), {"Occupied": 3, "Vacant": 2})
    c.equal("(iii) and its figures are the totals", (st_e.get("units_total"), st_e.get("market_rent"), st_e.get("sqft")), (5, 7640, 3750))
    # an earlier review: (iv) a totals row titled `Total Occupancy` in column A with `Total`
    # in the Status column is the totals row, not the next table; (v) many blank rows
    # between the bucket rows do not exhaust the scan.
    def _resman_shape2(path, shape):
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Sheet"
        ws.cell(1, 1, "Unit"); ws.cell(1, 2, "Type"); ws.cell(1, 5, "Status"); ws.cell(1, 6, "Market Rent"); ws.cell(1, 8, "Amount")
        ws.cell(2, 1, "A-1"); ws.cell(2, 2, "1x1"); ws.cell(2, 5, "C"); ws.cell(2, 6, 1000); ws.cell(2, 8, 950)
        r = 5
        ws.cell(r, 1, "Property Occupancy"); r += 1
        ws.cell(r, 1, "Type"); ws.cell(r, 12, "Status"); ws.cell(r, 13, "Market Rent"); ws.cell(r, 14, "Units"); ws.cell(r, 15, "Square Footage"); r += 2
        ws.cell(r, 1, "Total"); ws.cell(r, 12, "Occupied"); ws.cell(r, 13, 4400); ws.cell(r, 14, 3); ws.cell(r, 15, 2200); r += 1
        if shape == "many_blanks":
            r += 14
        ws.cell(r, 1, "Total"); ws.cell(r, 12, "Vacant"); ws.cell(r, 13, 3240); ws.cell(r, 14, 2); ws.cell(r, 15, 1550); r += 1
        if shape == "total_occupancy_title":
            ws.cell(r, 1, "Total Occupancy"); ws.cell(r, 12, "Total")
        ws.cell(r, 13, 7640); ws.cell(r, 14, 5); ws.cell(r, 15, 3750); r += 2
        ws.cell(r, 1, "Unit Type Occupancy"); r += 1
        ws.cell(r, 1, "Type"); ws.cell(r, 12, "Status"); ws.cell(r, 13, "Market Rent"); ws.cell(r, 14, "Units"); ws.cell(r, 15, "Square Footage"); r += 2
        ws.cell(r, 1, "1x1"); ws.cell(r, 12, "Occupied"); ws.cell(r, 13, 2600); ws.cell(r, 14, 1); ws.cell(r, 15, 1300)
        wb.save(path)
        return openpyxl.load_workbook(path, data_only=True)["Sheet"]
    st_f = ex.parse_source_totals(_resman_shape2(tmp / "resman_total_occupancy.xlsx", "total_occupancy_title"), 1, 2)
    c.equal("(iv) `Total Occupancy` / `Total` is the totals row", (st_f.get("units_total"), st_f.get("market_rent"), st_f.get("sqft")), (5, 7640, 3750))
    c.equal("(iv) buckets intact", st_f.get("buckets"), {"Occupied": 3, "Vacant": 2})
    st_g = ex.parse_source_totals(_resman_shape2(tmp / "resman_many_blanks.xlsx", "many_blanks"), 1, 2)
    c.equal("(v) fourteen blank rows between buckets: both buckets still read", st_g.get("buckets"), {"Occupied": 3, "Vacant": 2})
    c.equal("(v) and the totals", (st_g.get("units_total"), st_g.get("market_rent")), (5, 7640))
    ws_b = _resman_variant(tmp / "resman_no_totals.xlsx", None)
    st_b = ex.parse_source_totals(ws_b, 1, 2)
    c.equal("no totals row at all: buckets are still the property's", st_b.get("buckets"), {"Occupied": 3, "Vacant": 2})
    c.equal("and the totals are honestly None, never the per-plan table's 2", (st_b.get("units_total"), st_b.get("market_rent")), (None, None))

    c.section("RealPage — the labelled `Occupancy and Rents Summary` still reads")
    pp = tmp / "realpage.xlsx"
    hdr2, last2 = _realpage_workbook(pp)
    ws2 = openpyxl.load_workbook(pp, data_only=True)[openpyxl.load_workbook(pp).sheetnames[0]]
    st2 = ex.parse_source_totals(ws2, hdr2, last2)
    c.equal("units_total from the Totals: row", st2.get("units_total"), 3)
    c.equal("market_rent from the LRO Market Rent column", st2.get("market_rent"), 4600)
    c.equal("current_rent from the Current Rent column", st2.get("current_rent"), 4725)
    c.equal("buckets", st2.get("buckets"), {"Notice Available": 1, "Occupied": 1, "Vacant Available": 1})
    c.equal("occupied_units counts the notice bucket as physically occupied",
            ex._roll_stated({"source_totals": st2}).get("occupied_units"), 2)

    c.section("neither block → an honest empty, never an assumed zero")
    bare = tmp / "bare.xlsx"
    wb = openpyxl.Workbook(); ws3 = wb.active
    ws3["A1"] = "Unit"; ws3["B1"] = "Rent"; ws3["A2"] = "101"; ws3["B2"] = 1000; ws3["A4"] = "Total"; ws3["B4"] = 1000
    wb.save(bare)
    st3 = ex.parse_source_totals(openpyxl.load_workbook(bare)[wb.sheetnames[0]], 1, 2)
    c.equal("units_total None", st3.get("units_total"), None)
    c.equal("market_rent None", st3.get("market_rent"), None)
    c.equal("buckets empty", st3.get("buckets"), {})
    return c.finish()


if __name__ == "__main__":
    raise SystemExit(guard(main))
