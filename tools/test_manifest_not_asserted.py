"""A closing-cost amount the pipeline never asserted reads MISSING on the manifest, never $0.

Found live on 2026-09-04 (the browser pass against the deployed API, a rent-roll-only
upload): the Field manifest tab printed ``$0`` for ``transfer_tax_amount`` and
``mortgage_tax_amount`` while the deal summary tab, reading the same block, said the
section was not run. On the upload path ``closing_costs`` is the absent shape —
``amount_status: "not_asserted"``, ``amount: 0``, ``applies: False`` — and both manifest
consumers (the preview fragment and the workbook writer) resolved the raw ``amount``
through ``schema.resolve_path`` without asking what the status said. A zero a reader can
copy into a model is a claim; ``zero_by_law`` (the published sets, where the statute says
$0) stays ``$0``, and every other status that is not an assertion reads MISSING.

Usage: python tools/test_manifest_not_asserted.py
"""

from __future__ import annotations

import copy
import json
import re
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from _testkit import Checker, guard  # noqa: E402

import render_previews  # noqa: E402
import schema  # noqa: E402
import write_manifest  # noqa: E402

# The COMMITTED fixture (md5-identical to outputs/fixture/deal_outputs.json), so the test
# runs from a clean checkout — outputs/ is gitignored.
FIXTURE = Path(__file__).resolve().parent / "fixtures" / "synthetic_full.json"
AMOUNT_FIELDS = ("transfer_tax_amount", "mortgage_tax_amount")


def _fragment_examples(deal: dict) -> dict[str, str]:
    html = render_previews.manifest_fragment(deal)
    text = re.sub(r"<style.*?</style>", "", html, flags=re.S)
    cells = [re.sub(r"\s+", " ", re.sub(r"<[^>]+>", "", c)).strip() for c in re.findall(r"<td[^>]*>(.*?)</td>", text, flags=re.S)]
    out: dict[str, str] = {}
    for i, cell in enumerate(cells):
        # Name, Value, Source, Fills in the model (2026-09-14) → the value is the next cell
        if cell in AMOUNT_FIELDS and i + 1 < len(cells):
            out[cell] = cells[i + 1]
    return out


AMOUNT_PATHS = ("closing_costs.transfer_tax.amount", "closing_costs.mortgage_tax.amount")


def _workbook_examples(deal: dict) -> dict[str, str]:
    from openpyxl import load_workbook

    with tempfile.TemporaryDirectory() as tmp:
        path = write_manifest.build(Path(tmp) / "manifest.xlsx", deal)
        ws = load_workbook(path)[write_manifest.SHEET]
        out: dict[str, str] = {}
        for row in ws.iter_rows(values_only=True):
            if row and row[1] in AMOUNT_FIELDS:
                out[row[1]] = row[7]   # column H: Example (demo deal), after Public demo
        return out


def _audit_trail_values(deal: dict) -> dict[str, object]:
    """The Audit trail sheet's "Value as published" for the two closing-cost amount paths."""
    from openpyxl import load_workbook

    with tempfile.TemporaryDirectory() as tmp:
        path = write_manifest.build(Path(tmp) / "manifest.xlsx", deal)
        wb = load_workbook(path)
        sheet = next(n for n in wb.sheetnames if "audit" in n.lower())
        out: dict[str, object] = {}
        for row in wb[sheet].iter_rows(values_only=True):
            if row and row[0] in AMOUNT_PATHS:
                out[row[0]] = row[1]
        return out


def main() -> int:
    c = Checker("test_manifest_not_asserted")
    if not FIXTURE.is_file():
        c.ok("committed fixture present", False, str(FIXTURE))
        return c.finish()
    published = json.loads(FIXTURE.read_text(encoding="utf-8"))

    c.section("zero_by_law — the published set — still reads $0")
    for name, val in _fragment_examples(published).items():
        c.equal(f"fragment: {name}", val, "$0")
    for name, val in _workbook_examples(published).items():
        # The Example column is the formatted string, as it always has been.
        c.equal(f"workbook: {name} reads $0", val, "$0")

    c.section("not_asserted — the upload path's absent shape — reads MISSING, never $0")
    absent = copy.deepcopy(published)
    for key in ("transfer_tax", "mortgage_tax"):
        absent["closing_costs"][key].update(
            {"applies": False, "amount_status": "not_asserted", "amount": 0,
             "rate_label": "Not looked up", "citation": "Not run in the upload demo.", "as_of": ""}
        )
    absent["closing_costs"]["price"] = None
    frag = _fragment_examples(absent)
    c.equal("both amount fields found on the fragment", sorted(frag), sorted(AMOUNT_FIELDS))
    for name in AMOUNT_FIELDS:
        c.equal(f"fragment: {name} reads MISSING", frag.get(name), schema.MISSING)
        c.ok(f"fragment: {name} is not $0", frag.get(name) != "$0", repr(frag.get(name)))
    book = _workbook_examples(absent)
    for name in AMOUNT_FIELDS:
        c.equal(f"workbook: {name} reads MISSING", book.get(name), schema.MISSING)

    c.section("the Audit trail sheet of the same workbook says MISSING too")
    trail_ok = _audit_trail_values(published)
    c.equal("zero_by_law: both amount paths are stamped on the trail", sorted(trail_ok), sorted(AMOUNT_PATHS))
    for path, val in trail_ok.items():
        c.ok(f"trail: {path} publishes a zero", str(val).strip("$") in ("0", "0.0"), repr(val))
    trail = _audit_trail_values(absent)
    for path in AMOUNT_PATHS:
        c.equal(f"trail: {path} reads MISSING when not asserted", trail.get(path), schema.MISSING)

    c.section("the resolver is one function both consumers call")
    field = next(f for f in schema.MANIFEST_FIELDS if f.name == "transfer_tax_amount")
    c.ok("schema.manifest_value exists", hasattr(schema, "manifest_value"))
    if hasattr(schema, "manifest_value"):
        c.equal("not_asserted -> MISSING", schema.manifest_value(absent, field), schema.MISSING)
        c.equal("zero_by_law -> 0", schema.manifest_value(published, field), 0.0)
    c.ok("schema.asserted_value exists (by path, for the audit trail)", hasattr(schema, "asserted_value"))
    if hasattr(schema, "asserted_value"):
        c.equal("by path: not_asserted -> MISSING", schema.asserted_value(absent, AMOUNT_PATHS[0]), schema.MISSING)
        c.equal("by path: a non-amount path is untouched", schema.asserted_value(published, "meta.deal_name"), published["meta"]["deal_name"])
    return c.finish()


if __name__ == "__main__":
    raise SystemExit(guard(main))
