#!/usr/bin/env python3
"""The Checks tab after 2026-09-10: NOT CHECKED says why, the untied group subtotals
collapse into one line, and the statement is tied ACROSS to the rent roll.

Usage: python tools/test_checks.py
"""
from __future__ import annotations

import copy
import json
import re
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from _testkit import Checker, guard  # noqa: E402

import compute  # noqa: E402
import render_previews  # noqa: E402
import schema  # noqa: E402
import write_t12_analysis  # noqa: E402

FIXTURE = Path(__file__).resolve().parent / "fixtures" / "synthetic_full.json"


def _text(html: str) -> str:
    html = re.sub(r"<style.*?</style>", "", html, flags=re.S)
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", html))


def main() -> int:
    c = Checker("test_checks")
    deal = json.loads(FIXTURE.read_text(encoding="utf-8"))

    c.section("NOT CHECKED carries its reason; the headline ties are flagged")
    checks = compute.error_checks(deal)
    c.equal("three headline ties", [x["check"] for x in checks if x["headline"]],
            ["Total Revenue", schema.PNL_TOTAL_OPEX_LABEL, schema.PNL_NOI_LABEL])
    nc = [x for x in checks if x["status"] == "NOT CHECKED"]
    c.ok("a NOT CHECKED row names the missing subtotal",
         all("printed no" in (x.get("reason") or "") and x["check"] in x["reason"] for x in nc), str(nc))
    c.ok("an OK row carries no reason", all(x.get("reason") is None for x in checks if x["status"] == "OK"))

    c.section("two or more untied group subtotals collapse into one line on the tab")
    two = copy.deepcopy(deal)
    for cat in ("Other Income", "Utilities"):
        two["t12"]["source_subtotals"].pop(cat, None)
    untied = [x for x in compute.error_checks(two) if not x["headline"] and x["status"] == "NOT CHECKED"]
    c.ok("the fixture now has >= 2 untied group ties", len(untied) >= 2, str(len(untied)))
    text = _text(render_previews.t12_checks_fragment(two))
    c.ok("one collapsed line names the groups", "Group subtotals (" in text and "Other Income" in text)
    c.ok("…and says why", "prints no subtotal for these" in text)
    c.ok("the headline ties are still their own rows", "Total Revenue" in text and "Net Operating Income" in text)
    one = copy.deepcopy(deal)
    one["t12"]["source_subtotals"].pop("Utility Reimbursements", None)
    text1 = _text(render_previews.t12_checks_fragment(one))
    c.ok("a single untied group is left as its own row (no collapse)", "Group subtotals (" not in text1)
    c.ok("with the reason beside it", "printed no Utility Reimbursements subtotal" in text1)

    c.section("across the two documents: statement vs rent roll")
    x = {r["check"]: r for r in compute.cross_document_checks(deal)}
    gpr = x["Gross Potential Rent vs the roll's market rent x 12"]
    rev = x["Total Revenue vs the roll's in-place rent x 12"]
    p = compute.pnl(deal)
    units = deal["rent_roll"]["units"]
    c.close("GPR tie: statement side is the P&L's GPR", gpr["ours"], p["rows"]["Gross Potential Rent"]["t12"], 0.01)
    c.close("GPR tie: roll side is market rent x 12 over every unit", gpr["source"],
            sum(u["market_rent"] for u in units) * 12, 0.01)
    c.ok("GPR tie: inside the band reads OK", gpr["status"] == "OK" and abs(gpr["pct"]) <= compute.CROSS_DOC_BAND,
         f"{gpr['status']} {gpr['pct']}")
    occ = [u for u in units if u["status"].lower() == "occupied" and isinstance(u.get("in_place_rent"), (int, float))]
    c.close("revenue tie: roll side is occupied in-place rent x 12", rev["source"],
            sum(u["in_place_rent"] for u in occ) * 12, 0.01)
    c.close("revenue tie: statement side is total revenue", rev["ours"], p["total_revenue"], 0.01)
    wide = copy.deepcopy(deal)
    for u in wide["rent_roll"]["units"]:
        u["market_rent"] = round(u["market_rent"] * 1.25, 2)
    gpr_w = {r["check"]: r for r in compute.cross_document_checks(wide)}["Gross Potential Rent vs the roll's market rent x 12"]
    c.equal("a 25% gap reads REVIEW", gpr_w["status"], "REVIEW")
    nomarket = copy.deepcopy(deal)
    for u in nomarket["rent_roll"]["units"]:
        u["market_rent"] = None
    gpr_n = {r["check"]: r for r in compute.cross_document_checks(nomarket)}["Gross Potential Rent vs the roll's market rent x 12"]
    c.equal("no market-rent column -> NOT CHECKED", gpr_n["status"], "NOT CHECKED")
    c.ok("…with the reason", "market rent on 0 of" in (gpr_n.get("reason") or ""), str(gpr_n.get("reason")))
    not12 = copy.deepcopy(deal)
    not12["t12"] = {"month_labels": [], "units": None, "rows": [], "source_subtotals": {},
                    "broker_proforma": None, "absent_reason": "No trailing-twelve income statement was uploaded."}
    rows_n = compute.cross_document_checks(not12)
    c.ok("no statement -> both ties NOT CHECKED with the reason",
         all(r["status"] == "NOT CHECKED" and "no trailing-twelve" in r["reason"] for r in rows_n), str(rows_n))

    c.section
    nogpr = copy.deepcopy(deal)
    for r in nogpr["t12"]["rows"]:
        if r.get("category") == "Gross Potential Rent":
            r["category"] = "Gross Scheduled Rent"
    gpr_ng = {r["check"]: r for r in compute.cross_document_checks(nogpr)}["Gross Potential Rent vs the roll's market rent x 12"]
    c.equal("a statement with no GPR line -> NOT CHECKED, not a $0 REVIEW", gpr_ng["status"], "NOT CHECKED")
    c.ok("…with the reason naming the missing line", "Gross Potential Rent" in (gpr_ng.get("reason") or ""), str(gpr_ng.get("reason")))
    zero = copy.deepcopy(deal)
    for r in zero["t12"]["rows"]:
        if r.get("category") == "Gross Potential Rent":
            r["monthly"] = [0.0 for _ in (r.get("monthly") or [])]
    gpr_z = {r["check"]: r for r in compute.cross_document_checks(zero)}["Gross Potential Rent vs the roll's market rent x 12"]
    c.ok("a GPR line that nets to zero is still compared (a stated zero is a claim)",
         gpr_z["status"] == "REVIEW" and gpr_z["ours"] == 0, str(gpr_z))
    occ_all = [u for u in deal["rent_roll"]["units"] if str(u.get("status", "")).lower() == "occupied"]
    sparse = copy.deepcopy(deal)
    kept = False
    for u in sparse["rent_roll"]["units"]:
        if str(u.get("status", "")).lower() == "occupied":
            if not kept:
                kept = True
                continue
            u["in_place_rent"] = None
    rev_s = {r["check"]: r for r in compute.cross_document_checks(sparse)}["Total Revenue vs the roll's in-place rent x 12"]
    c.equal("in-place rent on 1 of N occupied units -> NOT CHECKED", rev_s["status"], "NOT CHECKED")
    c.ok("…with the coverage in the reason", f"1 of {len(occ_all)} occupied" in (rev_s.get("reason") or ""), str(rev_s.get("reason")))
    one_gap = copy.deepcopy(deal)
    next(u for u in one_gap["rent_roll"]["units"] if str(u.get("status", "")).lower() == "occupied")["in_place_rent"] = None
    rev_g = {r["check"]: r for r in compute.cross_document_checks(one_gap)}["Total Revenue vs the roll's in-place rent x 12"]
    c.ok("one missing rent among many still ties, and the basis states the coverage",
         rev_g["status"] != "NOT CHECKED" and f"{len(occ_all) - 1} of {len(occ_all)} occupied" in rev_g["basis"], str(rev_g))

    c.section
    occ_units = [u for u in deal["rent_roll"]["units"] if str(u.get("status", "")).lower() == "occupied"]
    two = copy.deepcopy(deal)
    two["rent_roll"]["units"] = copy.deepcopy(occ_units[:2])
    two["rent_roll"]["units"][1]["in_place_rent"] = None
    rev_2 = {r["check"]: r for r in compute.cross_document_checks(two)}["Total Revenue vs the roll's in-place rent x 12"]
    c.equal("in-place rent on 1 of 2 occupied units -> NOT CHECKED (int(0.9*2) floored to 1)", rev_2["status"], "NOT CHECKED")
    nine = copy.deepcopy(deal)
    nine["rent_roll"]["units"] = copy.deepcopy(deal["rent_roll"]["units"][:9])
    nine["rent_roll"]["units"][0]["market_rent"] = None
    gpr_9 = {r["check"]: r for r in compute.cross_document_checks(nine)}["Gross Potential Rent vs the roll's market rent x 12"]
    c.equal("market rent on 8 of 9 units (88.9%) -> NOT CHECKED", gpr_9["status"], "NOT CHECKED")
    ten = copy.deepcopy(deal)
    ten["rent_roll"]["units"] = copy.deepcopy(deal["rent_roll"]["units"][:10])
    ten["rent_roll"]["units"][0]["market_rent"] = None
    gpr_10 = {r["check"]: r for r in compute.cross_document_checks(ten)}["Gross Potential Rent vs the roll's market rent x 12"]
    c.ok("market rent on 9 of 10 units (exactly 90%) still ties", gpr_10["status"] != "NOT CHECKED", gpr_10["status"])

    c.section("the tab and the workbook show the same ties (I6)")
    text = _text(render_previews.t12_checks_fragment(deal))
    c.ok("the tab has the cross-document table", "Across the two documents" in text and "Gross Potential Rent vs the roll" in text)
    with tempfile.TemporaryDirectory() as tmp:
        from openpyxl import load_workbook
        path = write_t12_analysis.build(deal, Path(tmp) / "t12.xlsx")
        ws = load_workbook(path)["Checks & Pricing"]
        cells = [str(ws.cell(row=r, column=1).value or "") for r in range(1, ws.max_row + 1)]
        c.ok("the workbook has the cross-document block", any("Across the two documents" in v for v in cells))
        row = next((r for r, v in enumerate(cells, start=1) if v.startswith("Gross Potential Rent vs the roll")), None)
        c.ok("…with the GPR tie row", row is not None)
        if row:
            c.close("…as the same figure the tab shows (a value, not a formula)",
                    float(ws.cell(row=row, column=2).value), gpr["ours"], 0.01)
            c.equal("…and the same status", ws.cell(row=row, column=5).value, gpr["status"])
    return c.finish()


if __name__ == "__main__":
    raise SystemExit(guard(main))
