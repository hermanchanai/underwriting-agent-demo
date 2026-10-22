#!/usr/bin/env python3
"""The underwriting pack: every output in ONE workbook, tabs in the order the page shows them.

One download carrying every output, one tab each. The tabs, in order:

    Deal summary · Demand · Rent comps · Property tax · Supply ·
    T-12 Raw data · T-12 Analysis · T-12 Audit trail ·
    RR Raw data · RR Analysis · RR Audit trail ·
    Model-ready outputs · Audit trail

Nothing here writes a cell of its own. Each tab is the existing writer's sheet function,
called into one workbook with the sheet constants pointed at the pack's tab names for the
duration of the call, so every cross-sheet formula (the P&L's SUMIFs over the data tab,
the audit trail's cell references) resolves against the pack's own tabs. The four
single-purpose workbooks are still built beside it — they are the writers' own contract.

    python tools/write_pack.py --deal outputs/<set>/deal_outputs.json --out <path>.xlsx
"""
from __future__ import annotations

import argparse
import contextlib
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import Workbook  # noqa: E402

import schema  # noqa: E402
import write_deal_summary as ds  # noqa: E402
import write_manifest as mf  # noqa: E402
import write_rent_roll as rr  # noqa: E402
import write_t12_analysis as t12  # noqa: E402
from _xl import ILLUSTRATIVE_MARK, save_workbook, set_widths, sheet_header, write_audit_trail  # noqa: E402

TABS = (
    "Deal summary", "Demand", "Rent comps", "Property tax", "Supply",
    "T-12 Raw data", "T-12 Analysis", "T-12 Audit trail",
    "RR Raw data", "RR Analysis", "RR Audit trail",
    "Model-ready outputs", "Audit trail",
)


@contextlib.contextmanager
def _renamed(module, **names):
    """Point a writer's sheet-name constants at the pack's tab names for one call."""
    saved = {k: getattr(module, k) for k in names}
    try:
        for k, v in names.items():
            setattr(module, k, v)
        yield
    finally:
        for k, v in saved.items():
            setattr(module, k, v)


def _fresh_active(wb: Workbook):
    """A blank sheet made active, for writers that take ``wb.active`` and title it."""
    ws = wb.create_sheet()
    wb.active = wb.index(ws)
    return ws


def build(deal: dict, out_path: Path) -> Path:
    errors = schema.validate(deal)
    if errors:
        raise ValueError("deal failed schema validation:\n  - " + "\n  - ".join(errors))
    wb = Workbook()
    illustrative = deal["meta"].get("illustrative_pricing", False)
    mark = f" {ILLUSTRATIVE_MARK}" if illustrative else ""

    with _renamed(ds, SHEET_SUMMARY="Deal summary", SHEET_DEMAND="Demand", SHEET_COMPS="Rent comps",
                  SHEET_TAX="Property tax", SHEET_SUPPLY="Supply"):
        ds._write_summary_tab(wb, deal)
        ds._write_demand_tab(wb, deal)
        ds._write_comps_tab(wb, deal)
        ds._write_tax_tab(wb, deal)
        ds._write_supply_tab(wb, deal)

    if not schema.section_absent(deal, "t12"):
        with _renamed(t12, SHEET_DATA="T-12 Raw data", SHEET_PNL="T-12 Analysis", SHEET_CHECKS="T-12 Audit trail"):
            _fresh_active(wb)
            ref = t12._write_data_tab(wb, deal)
            pnl_block = t12._write_pnl_tab(wb, deal, ref)
            pnl_ws = wb["T-12 Analysis"]
            row = pnl_ws.max_row + 2
            t12.write_pricing_section(pnl_ws, row, deal, ref, pnl_block, mark)
            chk = wb.create_sheet("T-12 Audit trail")
            row = sheet_header(chk, "T-12 audit trail", deal["meta"]["deal_name"],
                               deal["meta"]["as_of_date"], note=schema.TOOLKIT_NAME)
            t12.write_error_check_section(chk, row, deal, pnl_block)
            set_widths(chk, {"A": 38, "B": 18, "C": 18, "D": 16, "E": 14, "F": 52})

    if not schema.section_absent(deal, "rent_roll"):
        with _renamed(rr, SHEET_ROLL="RR Raw data", SHEET_MIX="RR Analysis"):
            _fresh_active(wb)
            ref_rr = rr._write_roll_tab(wb, deal)
            rr._write_mix_tab(wb, deal, ref_rr)
        rchk = wb.create_sheet("RR Audit trail")
        row = sheet_header(rchk, "Rent roll audit trail", deal["meta"]["deal_name"],
                           deal["meta"]["as_of_date"], note=schema.TOOLKIT_NAME)
        t12.write_roll_checks_section(rchk, row, deal)
        set_widths(rchk, {"A": 38, "B": 18, "C": 18, "D": 16, "E": 14, "F": 52})

    man = wb.create_sheet("Model-ready outputs")
    mf.write_manifest_sheet(man, deal, schema.manifest_fields_for(deal))

    write_audit_trail(wb, deal, "manifest", title="The whole trail, every stamped value")
    # openpyxl de-duplicates titles case-insensitively against the sheet's OWN name, so a
    # direct "Audit Trail" -> "Audit trail" lands on "Audit trail1"; go through a temp name.
    trail = wb["Audit Trail"]
    trail.title = "_audit_tmp"
    trail.title = "Audit trail"

    present = [n for n in wb.sheetnames]
    expected = [n for n in TABS if n in present]
    stray = [n for n in present if n not in TABS]
    if stray:
        raise RuntimeError(f"pack: unexpected sheet(s) {stray}")
    wb._sheets = [wb[n] for n in expected]  # the page's order, exactly
    wb.active = 0
    return save_workbook(wb, out_path, deal["meta"].get("generated_on"))


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--deal", required=True, type=Path)
    ap.add_argument("--out", required=True, type=Path)
    args = ap.parse_args(argv)
    deal = json.loads(args.deal.read_text(encoding="utf-8"))
    path = build(deal, args.out)
    print(f"wrote {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
