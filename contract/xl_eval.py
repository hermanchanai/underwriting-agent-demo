"""A small, strict formula evaluator for the workbooks this repository writes.

The writers emit formulas as strings and leave the arithmetic to Excel, so a workbook saved
by ``openpyxl`` carries no computed values. The chain contract still has to read the number
a person opening the file would see. This module computes it, for the formula vocabulary the
writers actually use, and refuses everything else.

Two failure kinds, never confused:

* :class:`XlErr` is a spreadsheet error VALUE (``#DIV/0!``, ``#VALUE!``, ``#N/A``). It
  propagates through arithmetic, and ``IFERROR`` / ``ISNUMBER`` consume it, as in Excel.
* :class:`Unsupported` is an EXCEPTION: this evaluator cannot compute the expression
  (unknown function, wildcard criterion, syntax it does not parse). ``IFERROR`` does not
  catch it, so an unsupported formula fails the contract by name instead of turning into
  a plausible fallback value.

Supported: numbers, strings, TRUE/FALSE, ``+ - * / ^ &``, ``= <> < <= > >=``, unary minus,
A1 references and ranges (``$`` anchors, quoted or bare sheet names), and the functions
SUM SUMIF COUNTIF AVERAGEIF AVERAGEIFS SUMPRODUCT IF IFERROR ISNUMBER AND OR NOT ABS MAX MIN.
"""

from __future__ import annotations

import re
from typing import Any

from openpyxl.utils import column_index_from_string

__all__ = ["XlErr", "Unsupported", "Evaluator"]


class Unsupported(Exception):
    """The evaluator cannot compute this expression. Never a spreadsheet value."""


class XlErr:
    __slots__ = ("code",)

    def __init__(self, code: str):
        self.code = code

    def __repr__(self) -> str:
        return self.code

    def __eq__(self, other: object) -> bool:
        return isinstance(other, XlErr) and other.code == self.code

    def __hash__(self) -> int:
        return hash(self.code)


DIV0 = XlErr("#DIV/0!")
VALUE = XlErr("#VALUE!")

# --------------------------------------------------------------------------------------
# Tokenizer
# --------------------------------------------------------------------------------------

_SHEET = r"(?:'(?:[^']|'')+'|[A-Za-z_][A-Za-z0-9_.]*)!"
_CELL = r"\$?[A-Z]{1,3}\$?\d+"
_TOKEN = re.compile(
    rf"""
    (?P<ws>\s+)
  | (?P<str>"(?:[^"]|"")*")
  | (?P<num>\d+(?:\.\d+)?(?:[eE][+-]?\d+)?)
  | (?P<ref>(?:{_SHEET})?{_CELL}(?::{_CELL})?)
  | (?P<func>[A-Z][A-Z0-9.]*)\(
  | (?P<bool>TRUE|FALSE)\b
  | (?P<op><>|<=|>=|[-+*/^&=<>(),])
    """,
    re.VERBOSE,
)


def tokenize(text: str) -> list[tuple[str, str]]:
    out: list[tuple[str, str]] = []
    pos = 0
    while pos < len(text):
        m = _TOKEN.match(text, pos)
        if not m:
            raise Unsupported(f"cannot tokenize at {text[pos:pos + 20]!r}")
        kind = m.lastgroup
        if kind == "func":
            out.append(("func", m.group("func")))
            out.append(("op", "("))
        elif kind != "ws":
            out.append((kind, m.group(kind)))
        pos = m.end()
    return out


# --------------------------------------------------------------------------------------
# Parser: precedence, low to high: comparison, &, + -, * /, ^, unary
# --------------------------------------------------------------------------------------


class _Parser:
    def __init__(self, tokens: list[tuple[str, str]]):
        self.t = tokens
        self.i = 0

    def peek(self) -> tuple[str, str] | None:
        return self.t[self.i] if self.i < len(self.t) else None

    def take(self) -> tuple[str, str]:
        tok = self.t[self.i]
        self.i += 1
        return tok

    def expect(self, value: str) -> None:
        tok = self.peek()
        if tok is None or tok[1] != value:
            raise Unsupported(f"expected {value!r}, got {tok!r}")
        self.i += 1

    def parse(self):
        node = self.comparison()
        if self.peek() is not None:
            raise Unsupported(f"trailing tokens from {self.peek()!r}")
        return node

    def _binary(self, ops: tuple[str, ...], nxt):
        node = nxt()
        while (tok := self.peek()) is not None and tok[0] == "op" and tok[1] in ops:
            self.take()
            node = ("bin", tok[1], node, nxt())
        return node

    def comparison(self):
        return self._binary(("=", "<>", "<", "<=", ">", ">="), self.concat)

    def concat(self):
        return self._binary(("&",), self.additive)

    def additive(self):
        return self._binary(("+", "-"), self.multiplicative)

    def multiplicative(self):
        return self._binary(("*", "/"), self.power)

    def power(self):
        return self._binary(("^",), self.unary)

    def unary(self):
        tok = self.peek()
        if tok is not None and tok[0] == "op" and tok[1] in ("-", "+"):
            self.take()
            operand = self.unary()
            return ("neg", operand) if tok[1] == "-" else operand
        return self.primary()

    def primary(self):
        tok = self.peek()
        if tok is None:
            raise Unsupported("unexpected end of formula")
        kind, value = self.take()
        if kind == "num":
            return ("lit", float(value))
        if kind == "str":
            return ("lit", value[1:-1].replace('""', '"'))
        if kind == "bool":
            return ("lit", value == "TRUE")
        if kind == "ref":
            return ("ref", value)
        if kind == "func":
            self.expect("(")
            args = []
            if not (self.peek() and self.peek()[1] == ")"):
                while True:
                    args.append(self.comparison())
                    if self.peek() and self.peek()[1] == ",":
                        self.take()
                        continue
                    break
            self.expect(")")
            return ("call", value, args)
        if kind == "op" and value == "(":
            node = self.comparison()
            self.expect(")")
            return node
        raise Unsupported(f"unexpected token {value!r}")


# --------------------------------------------------------------------------------------
# Values
# --------------------------------------------------------------------------------------


def _is_num(v: Any) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _to_num(v: Any):
    """Scalar coercion for arithmetic: blank is 0, TRUE is 1, numeric text converts."""
    if isinstance(v, XlErr):
        return v
    if v is None:
        return 0.0
    if isinstance(v, bool):
        return float(v)
    if _is_num(v):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v.replace(",", ""))
        except ValueError:
            return VALUE
    return VALUE


def _flatten(v: Any) -> list:
    if isinstance(v, list):
        return [x for row in v for x in (row if isinstance(row, list) else [row])]
    return [v]


def _compare(op: str, a: Any, b: Any) -> bool:
    if a is None and isinstance(b, str):
        a = ""
    if b is None and isinstance(a, str):
        b = ""
    if a is None:
        a = 0.0
    if b is None:
        b = 0.0
    rank = lambda x: 2 if isinstance(x, bool) else 1 if isinstance(x, str) else 0  # noqa: E731
    if rank(a) != rank(b):
        ra, rb = rank(a), rank(b)
        return {"=": False, "<>": True, "<": ra < rb, "<=": ra < rb,
                ">": ra > rb, ">=": ra > rb}[op]
    if isinstance(a, str):
        a, b = a.lower(), b.lower()
    return {"=": a == b, "<>": a != b, "<": a < b, "<=": a <= b, ">": a > b, ">=": a >= b}[op]


_CRIT_OP = re.compile(r"^(<>|<=|>=|=|<|>)")


def _criterion(crit: Any):
    """A SUMIF/COUNTIF criterion as a predicate over one cell value."""
    if isinstance(crit, XlErr):
        raise Unsupported(f"criterion is an error value {crit!r}")
    op = "="
    target: Any = crit
    if isinstance(crit, str):
        m = _CRIT_OP.match(crit)
        if m:
            op, target = m.group(1), crit[m.end():]
        if any(ch in str(target) for ch in "*?~"):
            raise Unsupported(f"wildcard criterion {crit!r}")
        if isinstance(target, str):
            try:
                target = float(target)
            except ValueError:
                pass
    if target is None:
        target = ""

    def match(v: Any) -> bool:
        if isinstance(v, XlErr):
            return False
        if _is_num(target):
            if not _is_num(v):
                return op == "<>"
            return _compare(op, float(v), float(target))
        if v is None:
            v = ""
        if not isinstance(v, str):
            return op == "<>"
        return _compare(op, v, target)

    return match


# --------------------------------------------------------------------------------------
# Evaluator
# --------------------------------------------------------------------------------------

_REF = re.compile(rf"^(?:(?P<sheet>{_SHEET[:-1]})!)?(?P<a>{_CELL})(?::(?P<b>{_CELL}))?$")
_CELL_PARTS = re.compile(r"^\$?([A-Z]{1,3})\$?(\d+)$")


def _rc(cell: str) -> tuple[int, int]:
    m = _CELL_PARTS.match(cell)
    return int(m.group(2)), column_index_from_string(m.group(1))


class Evaluator:
    """Lazy, memoised evaluation of an ``openpyxl`` workbook loaded WITH formulas."""

    def __init__(self, workbook):
        self.wb = workbook
        self._memo: dict[tuple[str, int, int], Any] = {}
        self._active: set[tuple[str, int, int]] = set()

    # -- public ---------------------------------------------------------------------

    def cell(self, sheet: str, row: int, col: int) -> Any:
        key = (sheet, row, col)
        if key in self._memo:
            return self._memo[key]
        if key in self._active:
            raise Unsupported(f"circular reference at {sheet}!R{row}C{col}")
        raw = self.wb[sheet].cell(row=row, column=col).value
        if isinstance(raw, str) and raw.startswith("="):
            self._active.add(key)
            try:
                value = self.formula(raw[1:], sheet)
            finally:
                self._active.discard(key)
        else:
            value = raw
        self._memo[key] = value
        return value

    def formula(self, text: str, sheet: str) -> Any:
        node = _Parser(tokenize(text)).parse()
        value = self._eval(node, sheet)
        if isinstance(value, list):
            raise Unsupported("a formula evaluated to a range, not a value")
        return value

    # -- internals --------------------------------------------------------------------

    def _ref(self, text: str, sheet: str) -> Any:
        m = _REF.match(text)
        if not m:
            raise Unsupported(f"bad reference {text!r}")
        target = m.group("sheet") or sheet
        if target.startswith("'"):
            target = target[1:-1].replace("''", "'")
        r1, c1 = _rc(m.group("a"))
        if not m.group("b"):
            return self.cell(target, r1, c1)
        r2, c2 = _rc(m.group("b"))
        return [[self.cell(target, r, c) for c in range(c1, c2 + 1)] for r in range(r1, r2 + 1)]

    def _eval(self, node, sheet: str) -> Any:
        kind = node[0]
        if kind == "lit":
            return node[1]
        if kind == "ref":
            return self._ref(node[1], sheet)
        if kind == "neg":
            v = _to_num(self._scalar(node[1], sheet))
            return v if isinstance(v, XlErr) else -v
        if kind == "bin":
            return self._binop(node[1], self._scalar(node[2], sheet), self._scalar(node[3], sheet))
        if kind == "call":
            return self._call(node[1], node[2], sheet)
        raise Unsupported(f"unknown node {kind}")

    def _scalar(self, node, sheet: str) -> Any:
        v = self._eval(node, sheet)
        if isinstance(v, list):
            raise Unsupported("a range was used where a single value is needed")
        return v

    def _binop(self, op: str, a: Any, b: Any) -> Any:
        if isinstance(a, XlErr):
            return a
        if isinstance(b, XlErr):
            return b
        if op == "&":
            text = lambda v: "" if v is None else ("TRUE" if v is True else "FALSE" if v is False else  # noqa: E731
                                                  (f"{v:g}" if _is_num(v) else str(v)))
            return text(a) + text(b)
        if op in ("=", "<>", "<", "<=", ">", ">="):
            return _compare(op, a, b)
        x, y = _to_num(a), _to_num(b)
        if isinstance(x, XlErr):
            return x
        if isinstance(y, XlErr):
            return y
        if op == "+":
            return x + y
        if op == "-":
            return x - y
        if op == "*":
            return x * y
        if op == "/":
            return DIV0 if y == 0 else x / y
        if op == "^":
            return x ** y
        raise Unsupported(f"operator {op!r}")

    def _call(self, name: str, args: list, sheet: str) -> Any:
        ev = lambda n: self._eval(n, sheet)  # noqa: E731
        if name == "IF":
            if len(args) not in (2, 3):
                raise Unsupported("IF takes 2 or 3 arguments")
            cond = self._scalar(args[0], sheet)
            if isinstance(cond, XlErr):
                return cond
            truthy = bool(_to_num(cond)) if not isinstance(cond, str) else bool(cond)
            if truthy:
                return self._scalar(args[1], sheet)
            return self._scalar(args[2], sheet) if len(args) == 3 else False
        if name == "IFERROR":
            v = self._scalar(args[0], sheet)
            return self._scalar(args[1], sheet) if isinstance(v, XlErr) else v
        if name == "ISNUMBER":
            return _is_num(self._scalar(args[0], sheet))
        if name in ("AND", "OR"):
            vals = [x for a in args for x in _flatten(ev(a)) if x is not None]
            for v in vals:
                if isinstance(v, XlErr):
                    return v
            bools = [bool(_to_num(v)) for v in vals]
            return all(bools) if name == "AND" else any(bools)
        if name == "NOT":
            v = self._scalar(args[0], sheet)
            return v if isinstance(v, XlErr) else not bool(_to_num(v))
        if name == "ABS":
            v = _to_num(self._scalar(args[0], sheet))
            return v if isinstance(v, XlErr) else abs(v)
        if name in ("SUM", "MAX", "MIN"):
            nums: list[float] = []
            for a in args:
                v = ev(a)
                if isinstance(v, list):
                    for x in _flatten(v):
                        if isinstance(x, XlErr):
                            return x
                        if _is_num(x):
                            nums.append(float(x))
                else:
                    x = _to_num(v)
                    if isinstance(x, XlErr):
                        return x
                    nums.append(x)
            if name == "SUM":
                return sum(nums)
            if not nums:
                return 0.0
            return max(nums) if name == "MAX" else min(nums)
        if name in ("SUMIF", "COUNTIF", "AVERAGEIF"):
            rng = _flatten(ev(args[0]))
            pred = _criterion(self._scalar(args[1], sheet))
            if name == "COUNTIF":
                return float(sum(1 for v in rng if pred(v)))
            vals = _flatten(ev(args[2])) if len(args) > 2 else rng
            if len(vals) != len(rng):
                raise Unsupported(f"{name}: ranges of different sizes")
            picked = [v for v, k in zip(vals, rng) if pred(k) and _is_num(v)]
            if name == "SUMIF":
                return float(sum(picked))
            return DIV0 if not picked else sum(picked) / len(picked)
        if name == "AVERAGEIFS":
            vals = _flatten(ev(args[0]))
            keep = [True] * len(vals)
            for r_node, c_node in zip(args[1::2], args[2::2]):
                rng = _flatten(ev(r_node))
                if len(rng) != len(vals):
                    raise Unsupported("AVERAGEIFS: ranges of different sizes")
                pred = _criterion(self._scalar(c_node, sheet))
                keep = [k and pred(v) for k, v in zip(keep, rng)]
            picked = [v for v, k in zip(vals, keep) if k and _is_num(v)]
            return DIV0 if not picked else sum(picked) / len(picked)
        if name == "SUMPRODUCT":
            arrays = [_flatten(ev(a)) for a in args]
            if len({len(a) for a in arrays}) != 1:
                raise Unsupported("SUMPRODUCT: arrays of different sizes")
            total = 0.0
            for items in zip(*arrays):
                prod = 1.0
                for x in items:
                    if isinstance(x, XlErr):
                        return x
                    prod *= float(x) if _is_num(x) else 0.0
                total += prod
            return total
        raise Unsupported(f"function {name} is not supported")
