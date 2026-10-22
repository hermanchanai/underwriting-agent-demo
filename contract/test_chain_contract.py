#!/usr/bin/env python3
"""Chain contract: source totals -> saved workbooks -> every reader.

    python contract/test_chain_contract.py

One end-to-end run of the real pipeline (``server/pipeline.py``, model in mock mode) on the
bundled synthetic deal, then every surface a person or a downstream model reads is checked
against the SOURCE's own printed numbers:

* **oracle** - read straight out of the raw T-12 and rent roll by this file's own small
  readers: the statement's printed Total Revenue / Total Operating Expenses / Net
  Operating Income rows (summed across their twelve month cells), and the roll's unit rows,
  occupied rents and printed square-footage total. Never the pipeline's output, and never
  its ``source_subtotals`` (that is the pipeline's own read of the same rows).
* **producer** - the pipeline writes the interchange JSON, five workbooks and the previews.
* **readers** - the interchange JSON re-read from disk; each workbook loaded from disk with
  its formula cells EVALUATED (``contract/xl_eval.py``; the writers leave the arithmetic to
  Excel); the HTML preview fragments.

Then the red leg: one month of one revenue line is changed in a copy of the source while
its printed total is left alone. The chain has to notice: the workbook's own error check
must read REVIEW, and the contract's comparison against the printed total must fail.
A contract nobody has seen fail is not a contract.
"""

from __future__ import annotations

import html
import json
import os
import re
import shutil
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("OPENPYXL_LXML", "False")

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
for _p in (str(REPO), str(REPO / "tools")):
    if _p not in sys.path:
        sys.path.insert(0, _p)

import openpyxl  # noqa: E402

from _testkit import Checker, guard  # noqa: E402
from contract.xl_eval import Evaluator  # noqa: E402

SAMPLE_INPUTS = REPO / "sample" / "inputs"
T12_NAME = "Wildflower Crossing T-12.xlsx"
ROLL_NAME = "Wildflower Crossing Rent Roll.xlsx"
TOTAL_LABELS = ("Total Revenue", "Total Operating Expenses", "Net Operating Income")
TOLERANCE = 1.0  # dollars; the same tolerance the workbook's own error check uses

# ---------------------------------------------------------------------------------------
# Oracles: deliberately naive, independent of tools/source_readers.py
# ---------------------------------------------------------------------------------------


def oracle_t12(path: Path) -> dict[str, float]:
    """The statement's printed total rows, each summed across its numeric month cells."""
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    out: dict[str, float] = {}
    for row in ws.iter_rows(values_only=True):
        label = next((v.strip() for v in row[:3] if isinstance(v, str) and v.strip() in TOTAL_LABELS), None)
        if label:
            out[label] = sum(v for v in row[2:] if isinstance(v, (int, float)) and not isinstance(v, bool))
    return out


def oracle_roll(path: Path) -> dict:
    """Unit rows below the 'Unit ID' header, the occupied rents, and the printed sqft total."""
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    printed_sqft = None
    for row in rows:
        for v in row:
            m = re.match(r"\s*Total Sqft:\s*([\d,]+)", str(v or ""))
            if m:
                printed_sqft = float(m.group(1).replace(",", ""))
    head = next(i for i, r in enumerate(rows) if r and r[0] == "Unit ID")
    cols = {name: i for i, name in enumerate(rows[head]) if name}
    units = [r for r in rows[head + 1:] if r and r[0] not in (None, "")]
    occupied = [r for r in units if str(r[cols["Unit Status"]]).strip().lower() == "occupied"]
    rent = sum(float(r[cols["Lease Rent"]]) for r in occupied)
    return {
        "units": len(units),
        "occupied": len(occupied),
        "sqft": float(sum(r[cols["Sqft"]] for r in units)),
        "printed_sqft": printed_sqft,
        "in_place_avg": rent / len(occupied),
    }


# ---------------------------------------------------------------------------------------
# Readers of the produced set
# ---------------------------------------------------------------------------------------


def row_by_label(ws, label: str, col: int = 1) -> int:
    """The row whose column ``col`` reads ``label``. Located by label, never by position."""
    for row in ws.iter_rows():
        v = row[col - 1].value
        if isinstance(v, str) and v.strip() == label:
            return row[0].row
    raise LookupError(f"{ws.title}: no row labelled {label!r}")


def evaluated(path: Path) -> tuple[openpyxl.Workbook, Evaluator]:
    wb = openpyxl.load_workbook(path)
    return wb, Evaluator(wb)


def money(x: float) -> str:
    return f"${x:,.0f}"


def run_pipeline(inputs: Path, out: Path) -> dict:
    from server import pipeline

    return pipeline.run([inputs], out, price_per_unit=215000, generated_on="2026-10-21")


# ---------------------------------------------------------------------------------------


def main() -> int:
    c = Checker("chain contract")
    tmp = Path(tempfile.mkdtemp(prefix="chain_contract_"))
    try:
        return _main(c, tmp)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _main(c: Checker, tmp: Path) -> int:
    c.section("oracle: the source's own printed numbers")
    t12 = oracle_t12(SAMPLE_INPUTS / T12_NAME)
    roll = oracle_roll(SAMPLE_INPUTS / ROLL_NAME)
    c.equal("the statement prints all three total rows", sorted(t12), sorted(TOTAL_LABELS))
    c.ok("and its NOI row is its own revenue less its own expenses",
         abs(t12["Total Revenue"] - t12["Total Operating Expenses"] - t12["Net Operating Income"]) <= TOLERANCE,
         str(t12))
    c.equal("the roll's printed sqft total equals the sum of its unit rows",
            roll["printed_sqft"], roll["sqft"])
    print(f"    oracle: revenue {money(t12['Total Revenue'])}, opex {money(t12['Total Operating Expenses'])}, "
          f"NOI {money(t12['Net Operating Income'])}; {roll['units']} units, {roll['occupied']} occupied")

    c.section("producer: the pipeline, end to end, mock model")
    out = tmp / "run"
    result = run_pipeline(SAMPLE_INPUTS, out)
    c.equal("it spent nothing", float(result["ledger"].get("dollars") or 0.0), 0.0)
    for name in ("deal-outputs.json", "t12-analysis.xlsx", "rent-roll-analysis.xlsx",
                 "deal-summary.xlsx", "underwriting-pack.xlsx", "model-ready-outputs.xlsx"):
        c.ok(f"wrote {name}", (out / name).is_file())

    c.section("reader 1: the interchange JSON, re-read from disk")
    import compute

    deal = json.loads((out / "deal-outputs.json").read_text(encoding="utf-8"))
    pnl = compute.pnl(deal)
    for label, key in zip(TOTAL_LABELS, ("total_revenue", "total_opex", "noi")):
        c.ok(f"{label} ties to the printed row", abs(pnl[key] - t12[label]) <= TOLERANCE,
             f"json {pnl[key]:,.2f} vs source {t12[label]:,.2f}")
    totals = compute.roll_totals(deal)
    c.equal("unit count", totals["total_units"], roll["units"])
    c.equal("occupied units", totals["occupied_units"], roll["occupied"])
    c.ok("weighted in-place rent", abs(totals["wavg_in_place_rent"] - roll["in_place_avg"]) <= 0.01,
         f"{totals['wavg_in_place_rent']} vs {roll['in_place_avg']}")
    c.equal("summary unit count", deal["summary"]["units_total"], roll["units"])

    c.section("reader 2: the T-12 analysis workbook, formulas evaluated")
    wb, ev = evaluated(out / "t12-analysis.xlsx")
    pnl_ws, checks_ws = wb["P&L Summary"], wb["Checks & Pricing"]
    for label in TOTAL_LABELS:
        got = ev.cell(pnl_ws.title, row_by_label(pnl_ws, label), 2)
        c.ok(f"P&L Summary: {label}", abs(got - t12[label]) <= TOLERANCE,
             f"workbook {got:,.2f} vs source {t12[label]:,.2f}")
        status = ev.cell(checks_ws.title, row_by_label(checks_ws, label), 5)
        c.equal(f"Error check: {label} status", status, "OK")

    c.section("reader 3: the combined underwriting pack, formulas evaluated")
    wb, ev = evaluated(out / "underwriting-pack.xlsx")
    ws = wb["T-12 Analysis"]
    for label in TOTAL_LABELS:
        got = ev.cell(ws.title, row_by_label(ws, label), 2)
        c.ok(f"T-12 Analysis: {label}", abs(got - t12[label]) <= TOLERANCE,
             f"workbook {got:,.2f} vs source {t12[label]:,.2f}")

    c.section("reader 4: the rent-roll analysis workbook, formulas evaluated")
    wb, ev = evaluated(out / "rent-roll-analysis.xlsx")
    mix = wb["Unit Mix Summary"]
    c.equal("Unit Mix: total units", ev.cell(mix.title, row_by_label(mix, "Total units"), 2), roll["units"])
    c.equal("Unit Mix: TOTAL row unit count",
            ev.cell(mix.title, row_by_label(mix, "TOTAL", col=2), 4), roll["units"])
    occupied = ev.cell(mix.title, row_by_label(mix, "Occupied units"), 2)
    c.equal("Unit Mix: occupied units", float(occupied), float(roll["occupied"]))
    rr = wb["Rent Roll"]
    head = row_by_label(rr, "Unit ID")
    listed = sum(1 for r in rr.iter_rows(min_row=head + 1) if r[0].value not in (None, ""))
    c.equal("Rent Roll tab lists every unit once", listed, roll["units"])

    c.section("reader 5: the deal summary workbook")
    wb, ev = evaluated(out / "deal-summary.xlsx")
    ws = wb["Deal Summary"]
    c.equal("Unit count", ev.cell(ws.title, row_by_label(ws, "Unit count"), 2), roll["units"])
    noi = ev.cell(ws.title, row_by_label(ws, "T-12 net operating income"), 2)
    c.ok("T-12 net operating income", abs(noi - t12["Net Operating Income"]) <= TOLERANCE, f"{noi:,.2f}")
    cap = ev.cell(ws.title, row_by_label(ws, "T-12 cap rate"), 2)
    c.ok("T-12 cap rate is NOI over the illustrative price",
         abs(cap - t12["Net Operating Income"] / (215000 * roll["units"])) < 1e-9, f"{cap}")

    c.section("reader 6: the HTML previews")
    pnl_html = html.unescape((out / "previews" / "t12-pnl.html").read_text(encoding="utf-8"))
    for label in TOTAL_LABELS:
        m = re.search(re.escape(label) + r"</td><td[^>]*>(\$[\d,]+)<", pnl_html)
        c.equal(f"t12-pnl: {label}", m.group(1) if m else None, money(t12[label]))
    mix_html = (out / "previews" / "unit-mix.html").read_text(encoding="utf-8")
    m = re.search(r"<dt>Total units</dt><dd>(\d+)<", mix_html)
    c.equal("unit-mix: total units", int(m.group(1)) if m else None, roll["units"])

    c.section("red leg: a source cell changed under an unchanged printed total")
    bent = tmp / "bent"
    shutil.copytree(SAMPLE_INPUTS, bent)
    wb = openpyxl.load_workbook(bent / T12_NAME)
    ws = wb.worksheets[0]
    target = next(r for r in ws.iter_rows() if r[1].value == "Other Income")
    target[2].value = float(target[2].value) + 5000.0
    wb.save(bent / T12_NAME)
    bent_out = tmp / "bent-run"
    run_pipeline(bent, bent_out)
    bent_oracle = oracle_t12(bent / T12_NAME)
    c.equal("the printed total did not move", bent_oracle["Total Revenue"], t12["Total Revenue"])
    wb, ev = evaluated(bent_out / "t12-analysis.xlsx")
    pnl_ws, checks_ws = wb["P&L Summary"], wb["Checks & Pricing"]
    got = ev.cell(pnl_ws.title, row_by_label(pnl_ws, "Total Revenue"), 2)
    c.ok("the workbook's revenue no longer ties to the printed row (the contract would fail)",
         abs(got - t12["Total Revenue"]) > TOLERANCE, f"{got:,.2f} vs {t12['Total Revenue']:,.2f}")
    c.equal("and the workbook's own error check says REVIEW",
            ev.cell(checks_ws.title, row_by_label(checks_ws, "Total Revenue"), 5), "REVIEW")

    return c.finish()


if __name__ == "__main__":
    raise SystemExit(guard(main))
